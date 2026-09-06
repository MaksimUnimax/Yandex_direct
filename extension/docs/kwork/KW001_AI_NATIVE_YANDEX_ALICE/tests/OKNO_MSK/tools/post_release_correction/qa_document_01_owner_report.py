from pathlib import Path
import argparse, re, json, hashlib
from docx import Document
from pypdf import PdfReader

parser=argparse.ArgumentParser(description="Independent recipient QA for OKNO_MSK document 01")
parser.add_argument("source_md", type=Path)
parser.add_argument("docx", type=Path)
parser.add_argument("pdf", type=Path)
parser.add_argument("output_json", type=Path)
args=parser.parse_args()

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
s=args.source_md.read_text(encoding="utf-8")
lines=s.splitlines()

def table_rows_after(heading):
    start=next(i for i,l in enumerate(lines) if l.startswith("## ") and heading in l)
    rows=[]
    for l in lines[start+1:]:
        if l.startswith("## "): break
        if l.startswith("|"):
            if re.match(r"^\|\s*:?-+",l): continue
            if "№ |" in l: continue
            rows.append(l)
    return rows

sec5_start=next(i for i,l in enumerate(lines) if l.startswith("## 5. "))
sec6_start=next(i for i,l in enumerate(lines) if l.startswith("## 6. "))
sec5_subheads=[l for l in lines[sec5_start+1:sec6_start] if l.startswith("### ")]
app1=table_rows_after("Приложение 1.")
app2=table_rows_after("Приложение 2.")
forbidden=["Stage","Step","QF","S18-","CMS","READY","HOLD","SEARCH_REQUIRED","DE_RISK","NO_CHANGE","INSUFFICIENT","proxy","прокси","пользовательская Алиса","Дополнительная проверка восьми важных тем","связанные страницы","изменения запрещены","спорные семейства","частично готовая работа","не создавать новые страницы","GitHub","checkpoint","пересборка отчёта"]
forbidden_hits={t:[i+1 for i,l in enumerate(lines) if t.lower() in l.lower()] for t in forbidden}
forbidden_hits={k:v for k,v in forbidden_hits.items() if v}
brands=["REHAU","Provedal","KBE","Accado","Vorne","Futurus","fapim","rehau thermo"]

d=Document(str(args.docx))
docx_text="\n".join([p.text for p in d.paragraphs] + ["\t".join(c.text for c in row.cells) for t in d.tables for row in t.rows])
reader=PdfReader(str(args.pdf))
pdf_text="\n".join((p.extract_text() or "") for p in reader.pages)
pdf_norm=re.sub(r"\s+", " ", pdf_text)
checks={
 "ordinary_labels_8":s.count("Что показала обычная выдача:")==8,
 "generative_labels_8":s.count("Что показал генеративный поиск:")==8,
 "site_meaning_labels_8":s.count("Что это означает для сайта:")==8,
 "seven_recommendations_have_why":s.count("Почему это важно:")==7,
 "official_generative_yandex_visible":"официальный генеративный поиск яндекса" in s.lower(),
 "alice_technology_visible":"технологий Алисы" in s,
 "demand_groups_21":len(sec5_subheads)==21,
 "ordinary_yandex_rows_75":len(app1)==75,
 "research_results_rows_34":len(app2)==34,
 "forbidden_internal_terms_absent":not forbidden_hits,
 "brands_and_source_spellings_preserved":all(b in s for b in brands),
 "docx_core_content_present":"Что показали обычный и генеративный поиск Яндекса" in docx_text and "холодное остекление балкона" in docx_text,
 "pdf_core_content_present":"Что показали обычный и генеративный поиск Яндекса" in pdf_norm and "Сводная карта результатов исследования" in pdf_norm,
 "docx_forbidden_terms_absent":not any(t.lower() in docx_text.lower() for t in forbidden),
 "pdf_forbidden_terms_absent":not any(t.lower() in pdf_text.lower() for t in forbidden),
}
result={"status":"PASS" if all(checks.values()) else "FAIL","checks":checks,"counts":{"generative_comparison_cases":8,"priority_recommendations":7,"company_fact_requests":2,"demand_groups":len(sec5_subheads),"ordinary_yandex_observations":len(app1),"research_result_rows":len(app2),"pdf_pages":len(reader.pages)},"forbidden_hits":forbidden_hits,"preserved_brand_spellings":brands,"sha256":{"md":sha(args.source_md),"docx":sha(args.docx),"pdf":sha(args.pdf)},"provider_calls_during_rebuild":0}
args.output_json.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(result,ensure_ascii=False,indent=2))
raise SystemExit(0 if result["status"]=="PASS" else 2)
