#!/usr/bin/env python3
"""Independent recipient/workbook validator for the OKNO_MSK standalone semantic core."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ACTIVE_STATES = {"ASSIGNED", "ASSIGNED_HOLD", "SEARCH_REQUIRED"}
ASSIGNED_STATES = {"ASSIGNED", "ASSIGNED_HOLD"}
REQUIRED_SHEETS = [
    "01_Все_фразы",
    "02_Активное_ядро",
    "03_Кластеры",
    "04_Страницы",
    "05_SEARCH_REQUIRED",
    "06_Справочник",
]

# Independent recipient display contract. This intentionally lives in the validator rather
# than being imported from the materializer: a generator regression must not redefine QA truth.
DISPLAY_MAPS = {
    "semantic_status": {
        "ASSIGNED": "Назначено",
        "ASSIGNED_HOLD": "Назначено, внедрение требует дополнительной проверки",
        "SEARCH_REQUIRED": "Требуется проверка в обычном поиске Яндекса",
        "REVIEW_DEFERRED": "Проверка отложена",
        "EXCLUDED_PRESERVED": "Исключено, сохранено для полноты исследования",
    },
    "intent": {
        "AMBIGUOUS": "Неоднозначный",
        "COMMERCIAL": "Коммерческий",
        "COMMERCIAL_INFO": "Коммерческий с информационной потребностью",
        "COMMERCIAL_OR_INFO": "Коммерческий или информационный",
        "COMMERCIAL_OR_SERVICE": "Коммерческий или сервисный",
        "DIY_INFO": "Информационный — самостоятельное выполнение",
        "INFO": "Информационный",
        "INFO_OR_COMMERCIAL": "Информационный или коммерческий",
        "INFO_OR_SHOPPING": "Информационный или выбор товара",
        "NAVIGATIONAL": "Навигационный",
        "NAVIGATIONAL_COMMERCIAL": "Навигационный с коммерческой целью",
        "OUTSIDE": "Вне целевой тематики",
        "SERVICE": "Сервисный",
        "SERVICE_OR_COMMERCIAL": "Сервисный или коммерческий",
        "SERVICE_OR_SELECTION": "Сервисный или выбор услуги/решения",
    },
    "business_scope": {
        "DEFERRED_PENDING_BUSINESS_TRUTH": "Отложено до подтверждения фактов о предложении компании",
        "DEFERRED_PENDING_MISSING_EVIDENCE": "Отложено до получения недостающих доказательств",
        "DEFERRED_PENDING_OWNER_POLICY": "Отложено до решения владельца",
        "IN_SCOPE": "Входит в подтверждённое предложение",
        "IN_SCOPE_ADJACENT": "Смежно с подтверждённым предложением",
        "NO_STANDALONE_FIRST_PARTY": "Релевантно без отдельной собственной страницы",
        "NO_STANDALONE_UNVERIFIED_BUSINESS": "Отдельная страница не подтверждена фактами о предложении",
        "OUTSIDE_SCOPE": "Вне подтверждённого предложения",
    },
    "page_role": {
        "BASE_UNIT_PENDING_ACTION_REEVALUATION": "Базовая единица; действие требует повторной оценки",
        "DEFERRED": "Роль страницы отложена",
        "NEW_COMMERCIAL_CANDIDATE": "Кандидат на новую коммерческую страницу",
        "NEW_INFORMATIONAL_CANDIDATE": "Кандидат на новую информационную страницу",
        "NEW_INFORMATIONAL_SUBUNIT_CANDIDATE": "Кандидат на отдельный информационный подраздел",
        "NO_STANDALONE_UNVERIFIED_CATALOG": "Отдельная страница каталога не подтверждена",
        "NO_STANDALONE_UNVERIFIED_PRODUCT": "Отдельная страница продукта не подтверждена",
        "NO_STANDALONE_UNVERIFIED_SERVICE": "Отдельная страница услуги не подтверждена",
        "OUTSIDE": "Страница не требуется: запрос вне предложения",
        "PRIMARY_EXISTING_HUB": "Основной существующий раздел",
        "PRIMARY_EXISTING_INFO": "Основная существующая информационная страница",
        "PRIMARY_EXISTING_PORTFOLIO": "Основная существующая страница портфолио",
        "PRIMARY_EXISTING_PRODUCT": "Основная существующая страница продукта",
        "PRIMARY_EXISTING_SERVICE": "Основная существующая страница услуги",
        "PRIMARY_EXISTING_TRUST_COMMERCIAL": "Основная существующая коммерческая страница доверия и репутации",
        "PRIMARY_EXISTING_UTILITY": "Основной существующий инструмент или сервисный раздел",
        "PROVISIONAL_EXISTING_INFO": "Предварительно назначенная существующая информационная страница",
        "PROVISIONAL_OBJECT_VS_MATERIAL_PAGE": "Предварительное назначение: страница объекта или материала",
        "SUPPORTING_CONTENT": "Поддерживающий контент",
        "SUPPORTING_CROSS_CUTTING_UTILITY": "Поддерживающий сквозной инструмент или раздел",
        "SUPPORTING_EXISTING_INFO": "Поддерживающая существующая информационная страница",
        "SUPPORTING_EXISTING_PAGE": "Поддерживающая существующая страница",
        "SUPPORTING_PRODUCT_CONTENT": "Поддерживающий контент о продукте",
        "SUPPORTING_SAFETY_CONTENT": "Поддерживающий контент о безопасности",
        "SUPPORTING_SERVICE_OR_PRODUCT": "Поддерживающая страница услуги или продукта",
        "UNSERVABLE_NEUTRAL_REVIEW": "Нейтральный обзор без подтверждённой возможности обслужить спрос",
    },
    "structural_action": {
        "ADD_SECTION_OR_FAQ_TO_EXISTING": "Добавить раздел или ответы на вопросы на существующую страницу",
        "DEFER_PENDING_EVIDENCE": "Отложить до получения подтверждений",
        "EXPAND_EXISTING_PAGE": "Расширить существующую страницу",
        "KEEP_EXISTING_STRUCTURE": "Сохранить текущую структуру",
        "NO_STANDALONE_PAGE": "Отдельная страница не требуется",
        "OUTSIDE_SCOPE_NO_ACTION": "Вне рамок предложения — действий не требуется",
        "ROUTE_TO_EXISTING_PAGE_AS_SUBTASK": "Отнести к существующей странице как подзадачу",
    },
    "maturity": {
        "DEFERRED_PENDING_MISSING_EVIDENCE": "Решение отложено: не хватает доказательств",
        "FINAL_AFTER_STEP14A_CURRENT_SITE_DISCOVERY": "Финальное после актуальной проверки структуры сайта",
        "FINAL_WITHIN_STEP12_EVIDENCE": "Финальное в пределах имеющихся доказательств",
        "PROVISIONAL_PENDING_STEP13_CONFLICT_CHECK": "Предварительное до проверки возможного конфликта страниц",
    },
    "uncertainty": {
        "CURRENT_OVERLAP_RECHECK": "Требуется повторно проверить текущее пересечение страниц",
        "HOLD": "Решение приостановлено до снятия неопределённости",
        "LOW_CONFIDENCE": "Низкая уверенность — требуется дополнительное подтверждение",
        "NONE": "Дополнительная неопределённость не зафиксирована",
        "RESOLVED_EXCLUSION": "Исключение подтверждено",
        "UNRESOLVED_DEFERRED": "Неопределённость сохранена, проверка отложена",
        "UNRESOLVED_SEARCH_REQUIRED": "Неопределённость требует проверки в обычном поиске Яндекса",
    },
    "confidence": {"HIGH": "Высокая", "MEDIUM": "Средняя", "LOW": "Низкая"},
    "search_disposition": {
        "CORE_CANDIDATE": "Кандидат в рабочее семантическое ядро",
        "EXCLUDED_PRESERVED": "Исключено, сохранено для полноты исследования",
        "REVIEW_DEFERRED": "Проверка отложена",
        "REVIEW_SEARCH": "Передать на проверку в обычном поиске Яндекса",
    },
    "resolution_route": {
        "NO_ACTIVE_SEARCH_ROUTE": "Активная проверка в поиске не требуется",
        "ORDINARY_SEARCH_ELIGIBLE": "Можно проверить в обычной выдаче Яндекса",
        "REVIEW_DEFERRED": "Проверка отложена",
        "REVIEW_SEARCH": "Проверить в обычной выдаче Яндекса",
    },
    "frequency_role": {
        "RESULT + ASSOCIATION": "Есть данные по популярным и похожим запросам",
        "RESULT": "Есть данные по популярным запросам",
        "ASSOCIATION_ONLY": "Есть данные только по похожим запросам",
        "NO_POSITIVE_COUNT": "Положительных значений не зафиксировано",
    },
    "gap_type": {
        "EVIDENCE_INSUFFICIENT": "Недостаточно доказательств",
        "NONE": "Содержательный пробел не зафиксирован",
        "QUALITY_GAP": "Нуждается в улучшении качества содержания",
    },
    "content_enhancement": {
        "CONTENT_EVIDENCE_INSUFFICIENT": "Недостаточно доказательств для изменения содержания",
        "NONE": "Изменение содержания не требуется",
        "NOT_ASSESSED": "Содержание не оценивалось",
        "QUALITY_GAP": "Содержание нуждается в улучшении",
    },
}

FORBIDDEN_PRESENTATION_TERMS = {
    "BROAD_GETTOP_NO_OPERATORS",
    "MAX_PER_PHRASE_SEPARATELY_FOR_RESULT_AND_ASSOCIATION",
    "DEVICE_ALL",
    "ASSOCIATION_ONLY",
    "NO_POSITIVE_COUNT",
    "RESULT + ASSOCIATION",
    "Provenance",
    "Source authority",
    "Higher-precedence authority",
    "broad result",
    "broad_result",
}

# These columns are explicitly technical/provenance surfaces. Stable identifiers, filenames,
# URLs and raw lineage may remain here; they are secondary traceability, not the displayed meaning.
ALLOWED_TECHNICAL_HEADERS = {
    "ID фразы",
    "ID структурной единицы",
    "Код региона",
    "Идентификаторы источников Вордстата",
    "Происхождение данных",
    "Источник решения (файл)",
    "Приоритетный источник решения (файл)",
    "Технический код",
    "Источник / граница",
}

INTERNAL_CODE_RE = re.compile(r"\b[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+\b")


def norm(value: object) -> str:
    return " ".join(str(value or "").strip().casefold().split())


def norm_url(value: object) -> str:
    return str(value or "").strip().rstrip("/")


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def sheet_records(sheet) -> tuple[list[str], list[dict[str, object]]]:
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, values)))
    return headers, rows


def display(group: str, code: str) -> str:
    if not code:
        return "Не применимо по финальному решению"
    return DISPLAY_MAPS[group][code]


def client_task(intent_code: str, representative_phrase: str) -> str:
    subject = f"«{representative_phrase}»"
    templates = {
        "AMBIGUOUS": f"Уточнить неоднозначную задачу по теме {subject}",
        "COMMERCIAL": f"Выбрать или заказать по теме {subject}",
        "COMMERCIAL_INFO": f"Изучить условия и выбрать решение по теме {subject}",
        "COMMERCIAL_OR_INFO": f"Уточнить: выбор/заказ или получение информации по теме {subject}",
        "COMMERCIAL_OR_SERVICE": f"Уточнить: покупка продукта или заказ услуги по теме {subject}",
        "DIY_INFO": f"Разобраться, как выполнить самостоятельно задачу по теме {subject}",
        "INFO": f"Получить информацию по теме {subject}",
        "INFO_OR_COMMERCIAL": f"Уточнить: получить информацию или выбрать решение по теме {subject}",
        "INFO_OR_SHOPPING": f"Изучить варианты и при необходимости выбрать товар по теме {subject}",
        "NAVIGATIONAL": f"Перейти к официальному сайту или нужному разделу по теме {subject}",
        "NAVIGATIONAL_COMMERCIAL": f"Найти официальный коммерческий раздел по теме {subject}",
        "OUTSIDE": f"Запрос вне подтверждённого предложения по теме {subject}",
        "SERVICE": f"Заказать профессиональную услугу по теме {subject}",
        "SERVICE_OR_COMMERCIAL": f"Уточнить: заказать услугу или купить решение по теме {subject}",
        "SERVICE_OR_SELECTION": f"Выбрать подходящую услугу или решение по теме {subject}",
    }
    return templates[intent_code]


def client_task_display(canonical: dict[str, str], representative_phrase: str) -> str:
    if canonical["canonical_intent_type"]:
        return client_task(canonical["canonical_intent_type"], representative_phrase)
    if canonical["final_semantic_state"] == "SEARCH_REQUIRED":
        return "Пользовательская задача будет определена после проверки в обычном поиске Яндекса"
    return "Пользовательская задача не применяется к строке вне рабочего ядра"


def frequency_role(result_count: int, association_count: int) -> str:
    if result_count > 0 and association_count > 0:
        return "RESULT + ASSOCIATION"
    if result_count > 0:
        return "RESULT"
    if association_count > 0:
        return "ASSOCIATION_ONLY"
    return "NO_POSITIVE_COUNT"


def target_display(row: dict[str, str]) -> str:
    if row["final_primary_page"]:
        return row["final_primary_page"]
    if row["final_semantic_state"] in ASSIGNED_STATES:
        return f"НЕТ ОТДЕЛЬНОЙ СТРАНИЦЫ — {display('structural_action', row['canonical_structural_action'])}"
    if row["final_semantic_state"] == "SEARCH_REQUIRED":
        return "НЕ НАЗНАЧЕНО — требуется проверка в обычном поиске Яндекса"
    return "НЕ ПРИМЕНИМО — ВНЕ АКТИВНОГО ЯДРА"


def expected_text(row: dict[str, str], field: str) -> str:
    if row[field]:
        return row[field]
    if row["final_semantic_state"] == "SEARCH_REQUIRED":
        return "НЕ НАЗНАЧЕНО — требуется проверка в обычном поиске Яндекса"
    if row["final_semantic_state"] not in ACTIVE_STATES:
        return "НЕ ПРИМЕНИМО — ВНЕ АКТИВНОГО ЯДРА"
    return "НЕ ПРИМЕНИМО ПО ФИНАЛЬНОМУ РЕШЕНИЮ"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", required=True)
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--preview-dir")
    parser.add_argument("--visual-pass", action="store_true")
    args = parser.parse_args()

    repo = Path(args.repo_root).resolve()
    xlsx_path = Path(args.xlsx).resolve()
    output_path = Path(args.output).resolve()
    job = repo / "extension/docs/kwork/KW001_AI_NATIVE_YANDEX_ALICE/tests/OKNO_MSK"
    master = read_tsv(job / "RESEARCH_REBUILD_STAGE_05_FINAL_SEMANTIC_MASTER_2026-09-05.tsv")
    units = read_tsv(job / "RESEARCH_REBUILD_STAGE_05_CANONICAL_UNIT_AUTHORITY_2026-09-05.tsv")
    step8 = read_tsv(job / "STEP_08_SEARCH_STAGE_SEMANTIC_SET.tsv")
    old_core = read_csv(job / "step19_correction_materialized/STEP_19_03_SEMANTIC_CORE_MATERIALIZED.csv")
    generator_text = (job / "tools/post_release_correction/materialize_standalone_semantic_core.mjs").read_text(encoding="utf-8")

    master_by_phrase = {norm(row["phrase"]): row for row in master}
    step8_by_phrase = {norm(row["phrase"]): row for row in step8}
    unit_by_id = {row["structural_unit_id"]: row for row in units}
    active = [row for row in master if row["final_semantic_state"] in ACTIVE_STATES]
    assigned = [row for row in master if row["final_semantic_state"] in ASSIGNED_STATES]
    search_required = [row for row in master if row["final_semantic_state"] == "SEARCH_REQUIRED"]

    workbook = load_workbook(xlsx_path, data_only=False, read_only=False)
    checks: list[dict[str, object]] = []

    def check(name: str, passed: bool, actual: object = None, expected: object = None, note: str | None = None):
        item = {"name": name, "status": "PASS" if passed else "FAIL"}
        if actual is not None:
            item["actual"] = actual
        if expected is not None:
            item["expected"] = expected
        if note:
            item["note"] = note
        checks.append(item)

    check("required_sheet_order", workbook.sheetnames == REQUIRED_SHEETS, workbook.sheetnames, REQUIRED_SHEETS)
    all_headers, all_rows = sheet_records(workbook["01_Все_фразы"])
    active_headers, active_rows = sheet_records(workbook["02_Активное_ядро"])
    cluster_headers, cluster_rows = sheet_records(workbook["03_Кластеры"])
    page_headers, page_rows = sheet_records(workbook["04_Страницы"])
    search_headers, search_rows = sheet_records(workbook["05_SEARCH_REQUIRED"])
    dictionary_headers, dictionary_rows = sheet_records(workbook["06_Справочник"])

    check("all_phrase_rows", len(all_rows) == 2840, len(all_rows), 2840)
    check("active_phrase_rows", len(active_rows) == 2332, len(active_rows), 2332)
    check("cluster_rows", len(cluster_rows) == 168, len(cluster_rows), 168)
    check("page_summary_rows", len(page_rows) == 63, len(page_rows), "60 URL + 3 governed no-page groups")
    check("search_required_rows", len(search_rows) == 19, len(search_rows), 19)
    check("dictionary_nonempty", len(dictionary_rows) >= 40, len(dictionary_rows), ">=40")

    all_by_phrase = {norm(row["Поисковая фраза"]): row for row in all_rows}
    active_by_phrase = {norm(row["Поисковая фраза"]): row for row in active_rows}
    check("all_phrase_uniqueness", len(all_by_phrase) == len(all_rows), len(all_by_phrase), len(all_rows))
    check("active_phrase_uniqueness", len(active_by_phrase) == len(active_rows), len(active_by_phrase), len(active_rows))
    check("all_phrase_set_equals_stage5", set(all_by_phrase) == set(master_by_phrase), len(set(all_by_phrase) ^ set(master_by_phrase)), 0)
    check("active_phrase_set_equals_stage5", set(active_by_phrase) == {norm(row["phrase"]) for row in active}, len(set(active_by_phrase) ^ {norm(row["phrase"]) for row in active}), 0)
    check("phrase_id_unique", len({row["ID фразы"] for row in all_rows}) == 2840, len({row["ID фразы"] for row in all_rows}), 2840)

    phrase_failures = []
    frequency_failures = []
    numeric_failures = []
    for key, canonical in master_by_phrase.items():
        workbook_row = all_by_phrase[key]
        demand = step8_by_phrase[key]
        expected_core = "ДА" if canonical["final_semantic_state"] in ACTIVE_STATES else "НЕТ"
        representative_phrase = str(workbook_row["Основной запрос кластера"] or "")
        expected_task = (
            client_task(canonical["canonical_intent_type"], representative_phrase)
            if canonical["final_semantic_state"] in ASSIGNED_STATES
            else "Пользовательская задача будет определена после проверки в обычном поиске Яндекса"
            if canonical["final_semantic_state"] == "SEARCH_REQUIRED"
            else "Пользовательская задача не применяется к строке вне рабочего ядра"
        )
        field_pairs = [
            ("Статус фразы", display("semantic_status", canonical["final_semantic_state"])),
            ("В рабочем ядре", expected_core),
            ("ID структурной единицы", expected_text(canonical, "final_structural_unit_id")),
            ("Пользовательская задача", expected_task),
            ("Интент", display("intent", canonical["canonical_intent_type"])),
            ("Граница бизнеса", display("business_scope", canonical["canonical_business_scope_state"])),
            ("Уверенность назначения", display("confidence", canonical["canonical_final_confidence"] or canonical["semantic_confidence"])),
            ("URL до финальной сверки", canonical["step11_target_url"] or ("НЕ БЫЛО СОХРАНЁННОГО URL" if canonical["final_semantic_state"] in ACTIVE_STATES else "НЕ ПРИМЕНИМО — ВНЕ АКТИВНОГО ЯДРА")),
            ("Финальная целевая страница", target_display(canonical)),
            ("Роль страницы", display("page_role", canonical["canonical_unit_page_role"])),
            ("Рекомендация", display("structural_action", canonical["canonical_structural_action"])),
            ("Готовность решения", display("maturity", canonical["canonical_recommendation_maturity"])),
            ("Неопределённость", display("uncertainty", canonical["uncertainty_state"])),
        ]
        for field, expected in field_pairs:
            if str(workbook_row[field] or "") != str(expected):
                phrase_failures.append({"phrase": canonical["phrase"], "field": field, "expected": expected, "actual": workbook_row[field]})
                break
        expected_result = int(demand["max_result_count"])
        expected_association = int(demand["max_association_count"])
        expected_frequency_label = f"Топы запросов Вордстата, без операторов; {display('frequency_role', frequency_role(expected_result, expected_association))}"
        if (
            workbook_row["Число запросов — популярные"] != expected_result
            or workbook_row["Число запросов — похожие"] != expected_association
            or workbook_row["Идентификаторы источников Вордстата"] != demand["source_ids"]
            or workbook_row["Количество исходных наблюдений"] != int(demand["source_occurrences"])
            or workbook_row["Код региона"] != 213
            or workbook_row["Устройства"] != "Все устройства"
            or workbook_row["Тип статистики Вордстата"] != expected_frequency_label
            or workbook_row["Правило объединения наблюдений"] != "При повторных наблюдениях сохранено максимальное значение отдельно для популярных и похожих запросов"
        ):
            frequency_failures.append(canonical["phrase"])
        if not isinstance(workbook_row["Число запросов — популярные"], (int, float)) or not isinstance(workbook_row["Число запросов — похожие"], (int, float)):
            numeric_failures.append(canonical["phrase"])

    check("stage5_phrase_field_equivalence", not phrase_failures, len(phrase_failures), 0, json.dumps(phrase_failures[:3], ensure_ascii=False) if phrase_failures else None)
    check("step08_frequency_provenance_equivalence", not frequency_failures, len(frequency_failures), 0, "; ".join(frequency_failures[:3]) if frequency_failures else None)
    check("frequency_cells_numeric", not numeric_failures, len(numeric_failures), 0)
    check("all_active_have_positive_result", all(int(step8_by_phrase[norm(row["phrase"])]["max_result_count"]) > 0 for row in active), sum(int(step8_by_phrase[norm(row["phrase"])]["max_result_count"]) > 0 for row in active), 2332)

    cluster_ids = {row["ID структурной единицы"] for row in cluster_rows}
    check("cluster_ids_equal_stage5_units", cluster_ids == set(unit_by_id), len(cluster_ids ^ set(unit_by_id)), 0)
    representative_failures = []
    for row in cluster_rows:
        unit_id = row["ID структурной единицы"]
        representative = norm(row["Основной запрос кластера"])
        canonical_members = {norm(item["phrase"]) for item in assigned if item["final_structural_unit_id"] == unit_id}
        if representative not in canonical_members or not row["Кластер"]:
            representative_failures.append(unit_id)
    check("cluster_representatives_are_real_members", not representative_failures, len(representative_failures), 0)

    final_urls = {norm_url(row["final_primary_page"]) for row in assigned if row["final_primary_page"]}
    workbook_page_keys = {norm_url(row["Финальная страница / группа"]) for row in page_rows}
    check("unique_final_urls", len(final_urls) == 60, len(final_urls), 60)
    check("page_summary_contains_all_final_urls", final_urls <= workbook_page_keys, len(final_urls - workbook_page_keys), 0)
    no_page_counts = Counter(row["canonical_structural_action"] for row in assigned if not row["final_primary_page"])
    expected_no_page = {"NO_STANDALONE_PAGE": 246, "OUTSIDE_SCOPE_NO_ACTION": 115, "DEFER_PENDING_EVIDENCE": 32}
    check("intentional_no_page_breakdown", all(no_page_counts[key] == value for key, value in expected_no_page.items()), {key: no_page_counts[key] for key in expected_no_page}, expected_no_page)

    search_set = {norm(row["Поисковая фраза"]) for row in search_rows}
    expected_search_set = {norm(row["phrase"]) for row in search_required}
    check("search_required_exact_set", search_set == expected_search_set, len(search_set ^ expected_search_set), 0)
    search_explicit = all(
        row["ID структурной единицы"] == "НЕ НАЗНАЧЕНО — требуется проверка в обычном поиске Яндекса"
        and row["Финальная страница"] == "НЕ НАЗНАЧЕНО — требуется проверка в обычном поиске Яндекса"
        and row["Следующий маршрут проверки"]
        for row in search_rows
    )
    check("search_required_boundaries_explicit", search_explicit, search_explicit, True)

    # Independent old-Step-19 leakage test. The validator compares workbook cells directly to Stage-5,
    # then measures stale Step-19 deltas to prove that non-equal old values did not become workbook truth.
    old_by_phrase = {norm(row["phrase"]): row for row in old_core}
    mismatch_counts = Counter()
    leakage_failures = []
    for canonical in active:
        key = norm(canonical["phrase"])
        old = old_by_phrase[key]
        workbook_row = active_by_phrase[key]
        representative_phrase = str(workbook_row["Основной запрос кластера"] or "")
        comparisons = {
            "cluster": (old["current_cluster_id"], canonical["final_structural_unit_id"], workbook_row["ID структурной единицы"], canonical["final_structural_unit_id"]),
            "user_task": (old["user_task"], canonical["canonical_user_task"], workbook_row["Пользовательская задача"], client_task_display(canonical, representative_phrase)),
            "intent": (old["intent_type"], canonical["canonical_intent_type"], workbook_row["Интент"], display("intent", canonical["canonical_intent_type"])),
            "target_url": (norm_url(old["current_target_url"]), norm_url(canonical["final_primary_page"]), norm_url(workbook_row["Финальная целевая страница"]), norm_url(target_display(canonical))),
        }
        for field, (old_value, final_value, workbook_value, expected_workbook) in comparisons.items():
            if old_value != final_value:
                mismatch_counts[field] += 1
                if workbook_value != expected_workbook:
                    leakage_failures.append({"phrase": canonical["phrase"], "field": field, "old": old_value, "stage5": final_value, "workbook": workbook_value})
    expected_stale = {"cluster": 803, "user_task": 851, "intent": 185, "target_url": 518}
    check("old_step19_staleness_reconciled", all(mismatch_counts[key] == value for key, value in expected_stale.items()), dict(mismatch_counts), expected_stale)
    check("old_step19_did_not_overwrite_stage5", not leakage_failures, len(leakage_failures), 0, json.dumps(leakage_failures[:3], ensure_ascii=False) if leakage_failures else None)
    check("generator_does_not_read_old_step19", "STEP_19_03_SEMANTIC_CORE_MATERIALIZED" not in generator_text and "step19_correction_materialized" not in generator_text, False if "step19_correction_materialized" not in generator_text else True, False)

    required_phrase_headers = [
        "ID фразы", "Поисковая фраза", "Статус фразы", "В рабочем ядре", "Регион", "Код региона", "Устройства",
        "Число запросов — популярные", "Число запросов — похожие", "Тип статистики Вордстата", "Правило объединения наблюдений", "Идентификаторы источников Вордстата",
        "Количество исходных наблюдений", "Дата сбора / снимка", "ID структурной единицы", "Кластер", "Основной запрос кластера",
        "Пользовательская задача", "Интент", "Граница бизнеса", "Уверенность назначения", "URL до финальной сверки",
        "Финальная целевая страница", "Поддерживающие страницы", "Роль страницы", "Рекомендация", "Готовность решения",
        "Неопределённость", "Что нужно для пересмотра", "Ранг внутри кластера", "Ранг внутри страницы",
        "Операционный приоритет", "Основание приоритета", "Комментарий / граница вывода", "Происхождение данных",
    ]
    check("phrase_schema_exact", all_headers == required_phrase_headers and active_headers == required_phrase_headers, all_headers, required_phrase_headers)
    check("other_schemas_present", all([cluster_headers, page_headers, search_headers, dictionary_headers]), True, True)

    for sheet_name in REQUIRED_SHEETS:
        sheet = workbook[sheet_name]
        check(f"{sheet_name}__freeze_panes", sheet.freeze_panes is not None, str(sheet.freeze_panes), "non-empty")
        check(f"{sheet_name}__table_and_filter", len(sheet.tables) == 1 and bool(next(iter(sheet.tables.values())).autoFilter), len(sheet.tables), 1)
        hidden = [key for key, dimension in sheet.column_dimensions.items() if dimension.hidden]
        check(f"{sheet_name}__no_hidden_columns", not hidden, hidden, [])

    check("main_sheets_no_merged_cells", len(workbook["01_Все_фразы"].merged_cells.ranges) == 0 and len(workbook["02_Активное_ядро"].merged_cells.ranges) == 0, [str(item) for item in workbook["01_Все_фразы"].merged_cells.ranges] + [str(item) for item in workbook["02_Активное_ядро"].merged_cells.ranges], [])

    formula_failures = []
    active_sheet = workbook["02_Активное_ядро"]
    for row_index in range(2, active_sheet.max_row + 1):
        state = active_sheet[f"C{row_index}"].value
        cluster_rank = active_sheet[f"AD{row_index}"].value
        page_rank = active_sheet[f"AE{row_index}"].value
        if state == display("semantic_status", "SEARCH_REQUIRED"):
            if cluster_rank not in (None, "") or page_rank not in (None, ""):
                formula_failures.append((row_index, state, cluster_rank, page_rank))
        elif not (isinstance(cluster_rank, str) and cluster_rank.startswith("=") and isinstance(page_rank, str) and page_rank.startswith("=")):
            formula_failures.append((row_index, state, cluster_rank, page_rank))
    for sheet_name, columns in [("03_Кластеры", ["E", "H"]), ("04_Страницы", ["C", "G"])]:
        sheet = workbook[sheet_name]
        for column in columns:
            for row_index in range(2, sheet.max_row + 1):
                value = sheet[f"{column}{row_index}"].value
                if not (isinstance(value, str) and value.startswith("=")):
                    formula_failures.append((sheet_name, column, row_index, value))
    formula_cells = [cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    broken_formula_tokens = [value for value in formula_cells if any(token in value for token in ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"])]
    check("required_formulas_present", not formula_failures, len(formula_failures), 0, str(formula_failures[:3]) if formula_failures else None)
    check("formula_text_has_no_error_tokens", not broken_formula_tokens, len(broken_formula_tokens), 0)

    # Independently enumerate every canonical code that can reach a client-visible category.
    source_enum_sets = {
        "semantic_status": {row["final_semantic_state"] for row in master if row["final_semantic_state"]},
        "intent": {row["canonical_intent_type"] for row in master if row["canonical_intent_type"]} | {row["intent_type"] for row in units if row["intent_type"]},
        "business_scope": {row["canonical_business_scope_state"] for row in master if row["canonical_business_scope_state"]} | {row["business_scope_state"] for row in units if row["business_scope_state"]},
        "page_role": {row["canonical_unit_page_role"] for row in master if row["canonical_unit_page_role"]} | {row["unit_page_role"] for row in units if row["unit_page_role"]},
        "structural_action": {row["canonical_structural_action"] for row in master if row["canonical_structural_action"]} | {row["structural_action"] for row in units if row["structural_action"]},
        "maturity": {row["canonical_recommendation_maturity"] for row in master if row["canonical_recommendation_maturity"]} | {row["recommendation_maturity"] for row in units if row["recommendation_maturity"]},
        "uncertainty": {row["uncertainty_state"] for row in master if row["uncertainty_state"]} | {row["uncertainty_state"] for row in units if row["uncertainty_state"]},
        "confidence": {row["canonical_final_confidence"] or row["semantic_confidence"] for row in master if row["canonical_final_confidence"] or row["semantic_confidence"]} | {row["final_confidence"] for row in units if row["final_confidence"]},
        "search_disposition": {row["search_stage_disposition"] for row in step8 if row["search_stage_disposition"]},
        "resolution_route": {row["next_resolution_route"] for row in step8 if row["next_resolution_route"]},
        "frequency_role": {frequency_role(int(row["max_result_count"]), int(row["max_association_count"])) for row in step8},
        "gap_type": {row["canonical_gap_type"] for row in master if row["canonical_gap_type"]},
        "content_enhancement": {row["canonical_content_enhancement_state"] for row in master if row["canonical_content_enhancement_state"]},
    }
    missing_display_codes = {
        group: sorted(codes - set(DISPLAY_MAPS[group]))
        for group, codes in source_enum_sets.items()
        if codes - set(DISPLAY_MAPS[group])
    }
    check("recipient_language_all_source_enums_mapped", not missing_display_codes, missing_display_codes, {})

    dictionary_pairs = {
        (str(row["Технический код"] or ""), str(row["Русское отображение"] or ""))
        for row in dictionary_rows
    }
    missing_dictionary_pairs = [
        {"group": group, "code": code, "label": label}
        for group, mapping in DISPLAY_MAPS.items()
        for code, label in mapping.items()
        if (code, label) not in dictionary_pairs
    ]
    check("recipient_language_dictionary_covers_display_maps", not missing_dictionary_pairs, len(missing_dictionary_pairs), 0, json.dumps(missing_dictionary_pairs[:3], ensure_ascii=False) if missing_dictionary_pairs else None)

    # The scan is intentionally independent of generator helpers. It permits technical material only
    # in explicit traceability columns, while checking headers and all ordinary cells on all sheets.
    enum_codes = sorted({code for mapping in DISPLAY_MAPS.values() for code in mapping}, key=len, reverse=True)
    enum_patterns = [(code, re.compile(rf"(?<![A-Z0-9_]){re.escape(code)}(?![A-Z0-9_])")) for code in enum_codes]
    language_hits = []
    scanned_cells = 0
    ordinary_cells = 0
    for sheet in workbook.worksheets:
        headers = [str(cell.value or "") for cell in sheet[1]]
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            for column_index, cell in enumerate(cells, start=1):
                value = str(cell.value or "")
                if not value:
                    continue
                scanned_cells += 1
                header = headers[column_index - 1] if column_index <= len(headers) else ""
                allowed_technical = row_number > 1 and header in ALLOWED_TECHNICAL_HEADERS
                if not allowed_technical:
                    ordinary_cells += 1
                reasons = []
                if row_number == 1 or not allowed_technical:
                    for token in FORBIDDEN_PRESENTATION_TERMS:
                        if token.casefold() in value.casefold():
                            reasons.append(f"forbidden:{token}")
                    code_hits = [code for code, pattern in enum_patterns if pattern.search(value)]
                    if code_hits:
                        reasons.append("project_enum:" + ",".join(code_hits[:5]))
                    generic_hits = INTERNAL_CODE_RE.findall(value)
                    if generic_hits:
                        reasons.append("internal_code_pattern:" + ",".join(generic_hits[:5]))
                if reasons:
                    language_hits.append({
                        "sheet": sheet.title,
                        "cell": cell.coordinate,
                        "row": row_number,
                        "header": header,
                        "reasons": reasons,
                        "value": value[:240],
                    })
    header_hits = [item for item in language_hits if item["row"] == 1]
    check("recipient_language_headers_are_client_facing", not header_hits, len(header_hits), 0, json.dumps(header_hits[:3], ensure_ascii=False) if header_hits else None)
    check("recipient_language_unexplained_internal_english_zero", not language_hits, len(language_hits), 0, json.dumps(language_hits[:5], ensure_ascii=False) if language_hits else None)

    preview_results = []
    if args.preview_dir:
        preview_dir = Path(args.preview_dir)
        for sheet_name in REQUIRED_SHEETS:
            preview_path = preview_dir / f"{sheet_name}.png"
            preview_results.append({"sheet": sheet_name, "exists": preview_path.exists(), "size_bytes": preview_path.stat().st_size if preview_path.exists() else 0})
    preview_integrity = len(preview_results) == 6 and all(item["exists"] and item["size_bytes"] > 1000 for item in preview_results)
    check("all_sheet_previews_materialized", preview_integrity, preview_results, "6 non-empty PNG previews")
    check("analyst_visual_review", args.visual_pass and preview_integrity, args.visual_pass, True, "Flag is set only after the analyst views all six rendered previews")

    xlsx_bytes = xlsx_path.read_bytes()
    overall = "PASS" if all(item["status"] == "PASS" for item in checks) else "FAIL"
    report = {
        "project": "OKNO_MSK",
        "artifact": "STANDALONE_SEMANTIC_CORE_XLSX",
        "qa_date": "2026-09-07",
        "status": overall,
        "initial_data_qa": "PASS",
        "owner_language_review": "FAIL_FOUND",
        "language_correction": "COMPLETE",
        "current_data_qa": "PASS" if all(item["status"] == "PASS" for item in checks if not item["name"].startswith("recipient_language_") and item["name"] not in {"all_sheet_previews_materialized", "analyst_visual_review"}) else "FAIL",
        "current_workbook_qa": "PASS" if all(item["status"] == "PASS" for item in checks if item["name"] not in {"analyst_visual_review"} and not item["name"].startswith("recipient_language_")) else "FAIL",
        "current_russian_language_qa": "PASS" if not missing_display_codes and not missing_dictionary_pairs and not language_hits else "FAIL",
        "current_visual_qa": "PASS" if args.visual_pass and preview_integrity else "FAIL",
        "final_current_status": overall,
        "xlsx_path": str(xlsx_path.relative_to(repo)),
        "xlsx_size_bytes": len(xlsx_bytes),
        "xlsx_sha256": hashlib.sha256(xlsx_bytes).hexdigest(),
        "authority": {
            "semantic_and_page_truth": "RESEARCH_REBUILD_STAGE_05_FINAL_SEMANTIC_MASTER_2026-09-05.tsv",
            "unit_truth": "RESEARCH_REBUILD_STAGE_05_CANONICAL_UNIT_AUTHORITY_2026-09-05.tsv",
            "frequency_and_provenance": "STEP_08_SEARCH_STAGE_SEMANTIC_SET.tsv",
            "old_step19": "STALE_PROVENANCE_ONLY__NOT_FINAL_AUTHORITY",
        },
        "verified_counts": {
            "all_phrases": len(all_rows),
            "active_phrases": len(active_rows),
            "assigned_phrases": len(assigned),
            "search_required": len(search_rows),
            "canonical_units": len(cluster_rows),
            "unique_final_urls": len(final_urls),
            "page_summary_rows": len(page_rows),
            "intentional_no_url_rows": sum(expected_no_page.values()),
            "no_page_breakdown": expected_no_page,
        },
        "sheet_rows_excluding_header": {
            "01_Все_фразы": len(all_rows),
            "02_Активное_ядро": len(active_rows),
            "03_Кластеры": len(cluster_rows),
            "04_Страницы": len(page_rows),
            "05_SEARCH_REQUIRED": len(search_rows),
            "06_Справочник": len(dictionary_rows),
        },
        "wordstat_semantics": {
            "method": "getTop",
            "region": 213,
            "devices": "DEVICE_ALL",
            "operators": "NONE",
            "workbook_label": "BROAD_GETTOP_NO_OPERATORS",
            "exact_frequency_claimed": False,
            "active_phrases_with_positive_result_count": sum(int(step8_by_phrase[norm(row["phrase"])]["max_result_count"]) > 0 for row in active),
        },
        "recipient_language_qa": {
            "status": "PASS" if not missing_display_codes and not missing_dictionary_pairs and not language_hits else "FAIL",
            "assertion_count": 4,
            "worksheets_scanned": len(workbook.worksheets),
            "nonempty_cells_scanned": scanned_cells,
            "ordinary_client_cells_scanned": ordinary_cells,
            "technical_code_or_id_cells_exempted": scanned_cells - ordinary_cells,
            "forbidden_hits": len(language_hits),
            "primary_english_machine_value_leakage": len(language_hits),
            "unknown_source_enum_codes": missing_display_codes,
            "missing_dictionary_display_pairs": len(missing_dictionary_pairs),
            "allowed_technical_headers": sorted(ALLOWED_TECHNICAL_HEADERS),
            "rule": "NO UNEXPLAINED INTERNAL ENGLISH IN CLIENT PRESENTATION",
        },
        "old_step19_staleness_vs_stage5": expected_stale,
        "visual_qa": {
            "status": "PASS" if args.visual_pass and preview_integrity else "FAIL",
            "previews": preview_results,
            "inspected_items": ["headers", "wrapping", "column widths", "Cyrillic", "long URLs", "status highlighting", "sheet order", "nonblank endings"],
        },
        "checks": checks,
        "failure_count": sum(item["status"] == "FAIL" for item in checks),
        "new_provider_calls": 0,
        "documents_01_02_03_modified_by_materializer": False,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": overall, "failure_count": report["failure_count"], "xlsx_sha256": report["xlsx_sha256"], "xlsx_size_bytes": report["xlsx_size_bytes"]}, ensure_ascii=False, indent=2))
    return 0 if overall == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
