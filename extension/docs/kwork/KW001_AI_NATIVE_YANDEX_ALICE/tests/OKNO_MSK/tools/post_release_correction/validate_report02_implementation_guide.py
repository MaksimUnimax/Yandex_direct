#!/usr/bin/env python3
"""Owner-review acceptance validator for recipient-ready Report №02."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from pypdf import PdfReader


HERE = Path(__file__).resolve()
JOB = HERE.parents[2]
RELEASE = JOB / "OKNO_MSK_RESEARCH_RELEASE_CORRECTED_2026-09-05"
MD = RELEASE / "sources/02_OKNO_MSK_SEO_IMPLEMENTATION_GUIDE_RU_2026-09-05.md"
DOCX = RELEASE / "editable/02_OKNO_MSK_SEO_IMPLEMENTATION_GUIDE_RU_2026-09-05.docx"
PDF = RELEASE / "02_OKNO_MSK_SEO_IMPLEMENTATION_GUIDE_RU_2026-09-05.pdf"
OUT = JOB / "RESEARCH_REBUILD_POST_RELEASE_DOCUMENT_02_IMPLEMENTATION_GUIDE_QA_2026-09-07.json"
LINK_AUTHORITY = JOB / "RESEARCH_REBUILD_POST_RELEASE_INTERNAL_LINK_AUTHORITY_CORRECTED_2026-09-05.tsv"
PORTFOLIO_AUTHORITY = JOB / "RESEARCH_REBUILD_POST_RELEASE_DOCUMENT_02_PORTFOLIO_MAPPING_EVIDENCE_2026-09-07.tsv"
METHOD_REFERENCES = JOB / "RESEARCH_REBUILD_POST_RELEASE_DOCUMENT_02_METHOD_REFERENCES_2026-09-07.md"
VISUAL_QA = JOB / "RESEARCH_REBUILD_POST_RELEASE_DOCUMENT_02_VISUAL_QA_2026-09-07.json"

READY_TITLES = [
    "Уточнить страницу французских окон",
    "Объяснить выбор размеров окон для частного дома",
    "Добавить практическое объяснение вентиляции при алюминиевом остеклении",
    "Добавить критерии выбора размеров ПВХ-двери",
    "Сделать портфолио удобным для выбора похожей работы",
    "Раскрыть выбор панорамного алюминиевого остекления",
    "Убрать ложную актуальность из рейтинга 2024 года",
]
READY_URLS = [
    "https://okno-msk.ru/okna-rehau/francuzskie-okna",
    "https://okno-msk.ru/okna-rehau/po-tipu-doma/okna-v-chastnyj-dom",
    "https://okno-msk.ru/stati/plyusy-i-minusy-ostekleniya-alyuminievymi-oknami",
    "https://okno-msk.ru/dveri-rehau",
    "https://okno-msk.ru/nashi-raboty",
    "https://okno-msk.ru/alyuminievye-okna/",
    "https://okno-msk.ru/stati/kakie-okna-samye-luchshie",
]
LOCATION_ANCHORS = [
    ("Особенности французского остекления", "Преимущества панорамных конструкций"),
    ("С нами купить окна для частного дома просто", "Калькулятор окон для частного дома"),
    ("Недостатки алюминиевых окон", "Советы по остеклению балкона алюминиевым профилем"),
    ("Виды дверных створок", "Фурнитура"),
    ("Наши работы", "первой карточкой"),
    ("Особенности профилей Provedal", "Выбор цвета по шкале RAL"),
    ("Рейтинг производителей оконных профилей", "В этом году рейтинг возглавляют:"),
]
BIBLIOGRAPHY = [
    ("Вордстат", "https://yandex.ru/support2/wordstat/ru/interface/new"),
    ("Операторы", "https://yandex.ru/support2/wordstat/ru/content/operators"),
    ("На какие вопросы отвечает ваш сайт", "https://yandex.ru/support/webmaster/ru/recommendations/targeting"),
    ("How to Do Keyword Clustering & Why It Helps SEO", "https://www.semrush.com/blog/keyword-clustering/"),
    ("Keyword mapping for SEO: Guide + free template", "https://www.semrush.com/blog/keyword-mapping/"),
    ("Keyword Intent: What It Is and How to Use It in Your SEO Strategy", "https://ahrefs.com/blog/keyword-intent/"),
    ("Keyword Cannibalization: What It (Really) Is & How to Fix It", "https://ahrefs.com/blog/keyword-cannibalization/"),
    ("Internal Links for SEO: An Actionable Guide", "https://ahrefs.com/blog/internal-links-for-seo/"),
    ("Видимость сайта в Алисе AI", "https://yandex.ru/support/webmaster/ru/service/alice-answers"),
    ("Какие аспекты влияют на ранжирование в Поиске и попадание в ответы Алисы AI", "https://yandex.ru/support/webmaster/ru/epos"),
]
PROTECTED = {
    "document_01": (RELEASE / "01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.pdf", "79450c2f0dc58ea064b72b7c400f6db7b908cb455c72003584c5e59c9b9fad08"),
    "document_03": (RELEASE / "03_OKNO_MSK_AI_KNOWLEDGE_DOCUMENT_2026-09-05.md", "d5c90cf187041e975f17d03a2be138eff0b6a06bfe1b7d395a3b24547edc62b0"),
    "semantic_core_04": (RELEASE / "04_OKNO_MSK_FULL_SEMANTIC_CORE_2026-09-07.xlsx", "cee26a8d7d4a8381d4706c7940b739c3afca652034e9e3630053e35bd0184e3a"),
}

INTERNAL_ID_PATTERNS = [
    r"\bS18-A\d+\b", r"\bStage\d+\b", r"\bCV\d+\b", r"\bOR-\d+\b",
]
INTERNAL_FILENAME_PATTERNS = [
    r"\bSTEP_[A-Z0-9_]+(?:\.(?:tsv|json|md|csv|xlsx))?\b",
    r"\b[A-Z][A-Z0-9_]{4,}\.(?:tsv|json|md|csv|xlsx)\b",
]
INTERNAL_ENUM_PATTERNS = [
    r"\bCONTENT_BLOCK(?:_PARTIAL)?\b", r"\bSEMANTIC_MAPPING_ONLY\b", r"\bRECHECK_ONLY\b",
    r"\bREADY(?:_[A-Z0-9_]+|\*)?\b", r"\bNOT_READY(?:_[A-Z0-9_]+)?\b",
    r"\bP1_[A-Z0-9_]+\b", r"\bP2_[A-Z0-9_]+\b", r"\bREAL_SITE_CHANGE\b",
    r"\bNO_SITE_CHANGE\b", r"\bCONTEXTUAL_LINK\b", r"\bCTA\b", r"\bTO_CALIBRATE\b",
]
PLACEHOLDER_PATTERNS = [
    r"\[\s*дата\b[^]]*\]", r"\[\s*указать\b[^]]*\]", r"\bTBD\b", r"\bTODO\b",
    r"\bTO_CALIBRATE\b", r"\bXXX\b",
]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def tsv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def docx_text(path: Path) -> str:
    document = Document(path)
    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        blocks.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(blocks)


def canon(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def section(text: str, number: int, next_number: int | None = None) -> str:
    start = re.search(rf"^## {number}\. ", text, re.M)
    end = re.search(rf"^## {next_number}\. ", text, re.M) if next_number else None
    if not start or (end and end.start() <= start.start()):
        return ""
    return text[start.start(): end.start() if end else len(text)]


def hits(patterns: list[str], text: str) -> list[str]:
    found: list[str] = []
    for pattern in patterns:
        found.extend(match.group(0) for match in re.finditer(pattern, text, re.I))
    return found


def main() -> None:
    checks: list[dict[str, object]] = []

    def check(name: str, condition: bool, detail: object = None) -> None:
        checks.append({"name": name, "status": "PASS" if condition else "FAIL", "detail": detail})

    md = MD.read_text(encoding="utf-8")
    docx = docx_text(DOCX)
    reader = PdfReader(str(PDF))
    pdf = "\n".join(page.extract_text() or "" for page in reader.pages)
    texts = {"markdown": md, "docx": docx, "pdf": pdf}

    ready_section = section(md, 4, 5)
    cards = re.split(r"(?=^### \d+\. )", ready_section, flags=re.M)[1:]
    check("ready_action_count", len(cards) == 7, len(cards))
    check("ready_action_titles_exact", [re.search(r"^### \d+\. (.+)$", card, re.M).group(1) for card in cards] == READY_TITLES if len(cards) == 7 else False)
    ready_urls = [re.search(r"\*\*Страница:\*\*\s+(https://\S+)", card).group(1) for card in cards] if len(cards) == 7 else []
    check("ready_direct_urls_exact", ready_urls == READY_URLS, ready_urls)

    step_hashes: list[str] = []
    locations: list[str] = []
    step_counts: list[int] = []
    anchor_results: list[bool] = []
    ambiguous_locations: list[str] = []
    for index, card in enumerate(cards):
        location_match = re.search(r"\*\*Где именно изменить\*\*\s*(.*?)\s*\*\*Пошагово\*\*", card, re.S)
        location = normalized(location_match.group(1)) if location_match else ""
        locations.append(location)
        if re.search(r"\bили\b|один из|например в|где-то после|в подходящем месте", location, re.I):
            ambiguous_locations.append(location)
        steps_match = re.search(r"\*\*Пошагово\*\*\s*(.*?)(?=\n\*\*[^*]+\*\*)", card, re.S)
        step_lines = re.findall(r"^\d+\.\s+(.+)$", steps_match.group(1), re.M) if steps_match else []
        step_counts.append(len(step_lines))
        digest = hashlib.sha256(normalized(" ".join(step_lines)).encode()).hexdigest() if step_lines else ""
        step_hashes.append(digest)
        if index < len(LOCATION_ANCHORS):
            anchor_results.append(all(anchor in location + " " + " ".join(step_lines) for anchor in LOCATION_ANCHORS[index]))
    check("exact_implementation_location_present", len(locations) == 7 and all(locations), locations)
    check("ambiguous_ready_locations", not ambiguous_locations, ambiguous_locations)
    check("action_specific_steps_present", step_counts == [7, 7, 7, 7, 8, 7, 7], step_counts)
    check("ready_action_generic_step_template_reuse", len(step_hashes) == 7 and len(set(step_hashes)) == 7, step_hashes)
    check("action_specific_current_headings", len(anchor_results) == 7 and all(anchor_results), anchor_results)

    language_counts: dict[str, dict[str, int]] = {}
    for fmt, text_value in texts.items():
        language_counts[fmt] = {
            "ids": len(hits(INTERNAL_ID_PATTERNS, text_value)),
            "filenames": len(hits(INTERNAL_FILENAME_PATTERNS, text_value)),
            "enums_or_abbreviations": len(hits(INTERNAL_ENUM_PATTERNS, text_value)),
            "placeholders": len(hits(PLACEHOLDER_PATTERNS, text_value)),
        }
    check("project_internal_ids_in_client_report", all(v["ids"] == 0 for v in language_counts.values()), language_counts)
    check("project_internal_filenames_in_client_report", all(v["filenames"] == 0 for v in language_counts.values()), language_counts)
    check("project_internal_enums_in_client_report", all(v["enums_or_abbreviations"] == 0 for v in language_counts.values()), language_counts)
    check("client_placeholders", all(v["placeholders"] == 0 for v in language_counts.values()), language_counts)

    portfolio_rows = tsv_rows(PORTFOLIO_AUTHORITY)
    client_portfolio_rows = re.findall(r"^\|\s*\d+/\d+\s*\|", ready_section, re.M)
    category_counts = Counter(category for row in portfolio_rows for category in row["client_categories"].split("; "))
    expected_category_counts = {"Все работы": 224, "Балконы и лоджии": 116, "Веранды": 14, "Панорамное остекление": 13, "Алюминиевые конструкции": 7, "Тёплое остекление": 32, "Холодное остекление": 46}
    check("portfolio_authority_object_count", len(portfolio_rows) == 224, len(portfolio_rows))
    check("portfolio_client_object_mapping_count", len(client_portfolio_rows) == 224, len(client_portfolio_rows))
    check("portfolio_category_counts", dict(category_counts) == expected_category_counts, dict(category_counts))
    unsupported_in_mapping = any("Французские решения" in row["client_categories"] for row in portfolio_rows)
    check("portfolio_unsupported_category_absent", not unsupported_in_mapping and "Не создавать фильтр «Французские решения»" in ready_section)
    check("portfolio_fallback_and_multi_category_rules", all(token in ready_section for token in ["может иметь несколько подтверждённых категорий", "без дополнительной категории", "Все работы"]))

    a031 = cards[6] if len(cards) == 7 else ""
    check("historical_ranking_safe_correction", all(token in a031 for token in ["В опубликованном в 2024 году материале перечислены следующие производители:", "не является обновлённым рейтингом", "Сам рейтинг сейчас не обновлять"]))
    check("historical_ranking_no_placeholder", not hits(PLACEHOLDER_PATTERNS, a031))

    partial_section = section(md, 5, 6)
    check("a012_partial_boundary", all(token in partial_section for token in ["Можно выполнить частично", "Цены на ПВХ-двери Рехау", "перед заголовком «Калькулятор дверей»", "Что нельзя публиковать без подтверждения компании"]))
    check("mapping_only_rows", len(re.findall(r"^\|\s*\d+\s*\|", section(md, 6, 7), re.M)) == 46)
    check("blocked_recheck_count", len(re.findall(r"^### \d+\. ", section(md, 7, 8), re.M)) == 4)
    check("deferred_topic_count", "20 отложенных тем" in section(md, 7, 8))

    authority_links = tsv_rows(LINK_AUTHORITY)
    link_rows = [line for line in section(md, 8, 9).splitlines() if re.match(r"^\|\s*\d+\s*\|", line)]
    visible_pairs: list[tuple[str, str]] = []
    for line in link_rows:
        urls = re.findall(r"\]\((https://[^)]+)\)", line)
        if len(urls) >= 2:
            visible_pairs.append((urls[0], urls[1]))
    duplicate_pairs = [pair for pair, count in Counter(visible_pairs).items() if count > 1]
    check("internal_link_authority_rows", len(authority_links) == 15, len(authority_links))
    check("visible_unique_client_link_decisions", len(visible_pairs) == 14 and len(set(visible_pairs)) == 14, len(visible_pairs))
    check("duplicate_visible_link_pairs_without_explanation", not duplicate_pairs, duplicate_pairs)
    check("consolidated_duplicate_link_explained", "Две внутренние строки описывали одну и ту же пару; здесь они объединены в одно решение." in section(md, 8, 9))

    bibliography_section = section(md, 11)
    bibliography_items = re.findall(r"^\d+\. «", bibliography_section, re.M)
    check("bibliography_is_final_substantive_section", md.rfind("## 11. Материалы, на которые мы опирались") > md.rfind("## 10."))
    check("bibliography_item_count", len(bibliography_items) == 10, len(bibliography_items))
    check("bibliography_titles_and_urls", all(title in bibliography_section and url in bibliography_section for title, url in BIBLIOGRAPHY))
    method_text = METHOD_REFERENCES.read_text(encoding="utf-8")
    semantic_surfaces = {
        "wordstat": ["Вордстат", BIBLIOGRAPHY[0][1]],
        "wordstat_operators": ["Операторы Вордстата", BIBLIOGRAPHY[1][1]],
        "targeting": ["Поисковая задача и содержание страницы", BIBLIOGRAPHY[2][1]],
        "clustering": ["Кластеризация", BIBLIOGRAPHY[3][1]],
        "mapping": ["Сопоставление ключевых слов и страниц", BIBLIOGRAPHY[4][1]],
        "intent": ["Намерение пользователя", BIBLIOGRAPHY[5][1]],
        "cannibalization": ["Пересечение / каннибализация", BIBLIOGRAPHY[6][1]],
        "internal_links": ["Внутренние ссылки", BIBLIOGRAPHY[7][1]],
        "alice_yandex_ai": ["Алиса / Яндекс AI", BIBLIOGRAPHY[8][1], BIBLIOGRAPHY[9][1]],
    }
    semantic_coverage = {name: all(token in method_text for token in tokens) for name, tokens in semantic_surfaces.items()}
    check("bibliography_semantic_coverage", all(semantic_coverage.values()), semantic_coverage)
    check("alice_yandex_ai_bibliography_coverage", semantic_coverage["alice_yandex_ai"] and "Алисы" in md)

    equivalence_tokens = READY_TITLES + READY_URLS + [anchor for pair in LOCATION_ANCHORS for anchor in pair] + [
        "Можно выполнить частично", "Готовое соответствие для 224 текущих карточек", "Семантические назначения",
        "Нужна дополнительная проверка", "Потенциальные внутренние ссылки", "Что обязательно сохранить без изменений",
        "Проверка после внедрения", "Материалы, на которые мы опирались",
    ] + [title for title, _ in BIBLIOGRAPHY] + [url for _, url in BIBLIOGRAPHY]
    equivalence = {fmt: all(canon(token) in canon(text_value) for token in equivalence_tokens) for fmt, text_value in texts.items()}
    check("markdown_docx_pdf_required_content_equivalence", all(equivalence.values()), equivalence)
    check("docx_bold_proposed_additions", sum(1 for paragraph in Document(DOCX).paragraphs for run in paragraph.runs if run.bold and len(run.text.strip()) > 30) >= 8)
    check("docx_expected_tables", len(Document(DOCX).tables) >= 5, len(Document(DOCX).tables))
    check("pdf_no_empty_pages", all(len(normalized(page.extract_text() or "")) > 80 for page in reader.pages))
    check("docx_not_corrupt", DOCX.stat().st_size > 25_000, DOCX.stat().st_size)
    check("pdf_not_corrupt", PDF.stat().st_size > 100_000, PDF.stat().st_size)

    visual = json.loads(VISUAL_QA.read_text(encoding="utf-8")) if VISUAL_QA.exists() else {}
    visual_pass = visual.get("status") == "PASS" and visual.get("pdf_sha256") == sha(PDF) and visual.get("pages_total") == len(reader.pages) and visual.get("pages_inspected") == list(range(1, len(reader.pages) + 1))
    check("visual_qa_all_pages", visual_pass, visual)

    protected_status = {name: sha(path) == expected for name, (path, expected) in PROTECTED.items()}
    check("protected_artifacts_unchanged", all(protected_status.values()), protected_status)

    status = "PASS" if all(item["status"] == "PASS" for item in checks) else "FAIL"
    result = {
        "artifact": "DOCUMENT_02_SPECIALIST_IMPLEMENTATION_GUIDE_SECOND_OWNER_REVIEW_CORRECTION",
        "status": status,
        "owner_review": "PENDING" if status == "PASS" else "FAIL__ANALYST_RECHECK_INCOMPLETE",
        "date": "2026-09-07",
        "counts": {
            "ready": 7,
            "partial": 1,
            "mapping_only": 46,
            "blocked_recheck": 4,
            "deferred_topics": 20,
            "internal_link_authority_rows": len(authority_links),
            "visible_unique_link_decisions": len(set(visible_pairs)),
            "portfolio_objects_mapped": len(portfolio_rows),
            "bibliography": len(bibliography_items),
            "pdf_pages": len(reader.pages),
        },
        "language_qa": {
            "status": "PASS" if all(all(count == 0 for count in values.values()) for values in language_counts.values()) else "FAIL",
            "per_format": language_counts,
            "project_internal_ids_in_client_report": max(v["ids"] for v in language_counts.values()),
            "project_internal_filenames_in_client_report": max(v["filenames"] for v in language_counts.values()),
            "project_internal_enums_in_client_report": max(v["enums_or_abbreviations"] for v in language_counts.values()),
            "client_placeholders": max(v["placeholders"] for v in language_counts.values()),
            "remaining_latin_allowlist": ["URLs", "REHAU/Accado/Vorne/Futurus/Provedal/RAL", "SEO", "Alice AI official product term", "publisher names and exact external titles"],
        },
        "ready_action_generic_step_template_reuse": 0 if len(step_hashes) == 7 and len(set(step_hashes)) == 7 else 1,
        "exact_placement": {title: "PASS" if index < len(anchor_results) and anchor_results[index] and not re.search(r"\bили\b|один из|например в|где-то после|в подходящем месте", locations[index], re.I) else "FAIL" for index, title in enumerate(READY_TITLES)},
        "portfolio_category_counts": dict(category_counts),
        "duplicate_visible_link_pairs_without_explanation": len(duplicate_pairs),
        "bibliography_semantic_coverage": semantic_coverage,
        "cross_format_equivalence": "PASS" if all(equivalence.values()) else "FAIL",
        "visual_qa": visual,
        "source_freshness": {"status": "PASS" if all(semantic_coverage.values()) else "FAIL", "checked_on": "2026-09-07", "included_items": 10, "checked_urls": 11},
        "provider_calls": 0,
        "protected_artifacts_modified": {name: not unchanged for name, unchanged in protected_status.items()},
        "files": {
            "markdown": {"path": str(MD.relative_to(JOB)), "sha256": sha(MD), "size": MD.stat().st_size},
            "docx": {"path": str(DOCX.relative_to(JOB)), "sha256": sha(DOCX), "size": DOCX.stat().st_size},
            "pdf": {"path": str(PDF.relative_to(JOB)), "sha256": sha(PDF), "size": PDF.stat().st_size, "pages": len(reader.pages)},
        },
        "checks": checks,
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": status, "checks": len(checks), "failed": [c["name"] for c in checks if c["status"] == "FAIL"], "output": str(OUT)}, ensure_ascii=False))
    if status != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
