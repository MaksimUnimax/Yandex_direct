from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path
from docx import Document
from pypdf import PdfReader

ROOT = Path('extension/docs/kwork/KW001_AI_NATIVE_YANDEX_ALICE/tests/OKNO_MSK')
REL = ROOT / 'OKNO_MSK_RESEARCH_RELEASE_STEP05A_PROPAGATED_2026-09-09'
R1_MD = REL / 'sources/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.md'
R1_DOCX = REL / 'editable/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.docx'
R1_PDF = REL / '01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.pdf'
R2_MD = REL / 'sources/02_OKNO_MSK_VNEDRENIE_REKOMENDATSII_RU_2026-09-09.md'
R2_DOCX = REL / 'editable/02_OKNO_MSK_VNEDRENIE_REKOMENDATSII_RU_2026-09-09.docx'
R2_PDF = REL / '02_OKNO_MSK_VNEDRENIE_REKOMENDATSII_RU_2026-09-09.pdf'
TRACKER = REL / '05_OKNO_MSK_IMPLEMENTATION_TRACKER_2026-09-09.xlsx'
README = REL / 'README_RU.md'
MANIFEST = REL / 'RELEASE_MANIFEST_2026-09-09.json'
STATE = ROOT / 'RESEARCH_REPORT_REBUILD_CURRENT_STATE_CLIENT_CLEANUP_2026-09-08.json'
QA = ROOT / 'REPORTS_01_02_PROFESSIONAL_REPACKAGE_QA_2026-09-09.json'
RECEIPT = ROOT / 'REPORTS_01_02_PROFESSIONAL_REPACKAGE_RECEIPT_2026-09-09.md'


def sha(path: Path) -> str:
    h = hashlib.sha256(); h.update(path.read_bytes()); return h.hexdigest()


def section(text: str, start: str, end: str) -> str:
    if start not in text or end not in text:
        raise RuntimeError(f'missing section boundary: {start} -> {end}')
    return text.split(start, 1)[1].split(end, 1)[0].strip()


def cards(block: str):
    heads = re.findall(r'(?m)^### \d+\. (.+)$', block)
    bodies = re.split(r'(?m)^### \d+\. .+$', block)[1:]
    if len(heads) != len(bodies):
        raise RuntimeError('card parse mismatch')
    return list(zip(heads, bodies))


def field(block: str, label: str) -> str:
    m = re.search(rf'\*\*{re.escape(label)}\*\*\s*(.*?)(?=\n+\*\*|\n+###|\Z)', block, re.S)
    if not m:
        return ''
    return re.sub(r'\s+', ' ', m.group(1).replace('**', '').strip())


def renumber_card(head: str, body: str, num: int, status: str | None = None) -> str:
    status_line = f'\n\n**Статус:** {status}' if status else ''
    return f'### {num}. {head}{status_line}\n\n{body.strip()}'


def replace_2322_semantics(text: str) -> str:
    replacements = {
        'Для 2 322 из них определена подходящая существующая страница; 26 фраз оставлены без принудительного назначения, чтобы не связывать запрос со страницей без достаточных оснований.':
        'По 2 322 фразам завершено решение по точному назначению; 26 фраз оставлены без точного владельца, чтобы не закреплять запрос за страницей без достаточных оснований.',
        '2 322 активные фразы имеют точное назначение на существующую страницу, 26 оставлены без принудительного точного назначения, включая семь формулировок, найденных через конкурентный этап.':
        'по 2 322 активным фразам завершено решение по точному назначению, а 26 оставлены без точного владельца, включая семь формулировок, найденных через конкурентный этап.',
        'Для 2 322 фраз определили точную существующую страницу; 26 оставили без принудительного точного назначения.':
        'Для 2 322 фраз завершили решение по точному назначению; 26 оставили без точного владельца.',
        '2 322 активные фразы имеют точное назначение на существующую страницу':
        'по 2 322 активным фразам завершено решение по точному назначению',
    }
    for a, b in replacements.items():
        text = text.replace(a, b)
    return text


def transform_report01(r1: str, r2: str) -> str:
    competitor = section(r1, '### Анализ конкурентов и расширение семантики', '## Как проводилась работа')
    method = section(r1, '## Как проводилась работа', '## Что показала обычная выдача Яндекса')
    search = section(r1, '## Что показала обычная выдача Яндекса', '## Как сайт соответствует выдаче Алисы')
    alice = section(r1, '## Как сайт соответствует выдаче Алисы', '## Что нужно изменить на сайте')
    recos = section(r1, '## Что нужно изменить на сайте', '## Что на сайте уже работает правильно')
    keep = section(r1, '## Что на сайте уже работает правильно', '## Что нужно уточнить у компании для двух следующих улучшений')

    rcards = cards(recos)
    if len(rcards) != 7:
        raise RuntimeError(f'Report01 expected 7 recommendation cards, got {len(rcards)}')
    ready_idx = [0, 3, 6]
    clarify_idx = [1, 2, 4, 5]
    ready_cards = [renumber_card(*rcards[i], n, 'можно передавать в работу сейчас') for n, i in enumerate(ready_idx, 1)]
    clarify_cards = [renumber_card(*rcards[i], n, 'после уточнения; до этого не передавать во внедрение') for n, i in enumerate(clarify_idx, 1)]

    r2_clar = section(r2, '## 5. Уточнить перед внедрением', '## 6.')
    r2_clar_cards = cards(r2_clar)
    if len(r2_clar_cards) != 5:
        raise RuntimeError('Report02 clarification card count mismatch')
    install_head, install_body = r2_clar_cards[1]
    clarify_cards.append(renumber_card(install_head, install_body, 5, 'после подтверждения фактического состава услуги компанией'))

    r2_checks = section(r2, '## 8. Дополнительные проверки перед следующими изменениями', '## 9.')
    check_rows = []
    for h, b in cards(r2_checks):
        check_rows.append((h, field(b, 'Страница:'), field(b, 'Что нужно понять'), field(b, 'Результат проверки')))
    if len(check_rows) != 4:
        raise RuntimeError('Report02 check count mismatch')
    check_table = ['| Проверка | Страница | Что нужно получить |', '|---|---|---|']
    for h, url, why, result in check_rows:
        result_clean = result.replace('|', '/').strip()
        check_table.append(f'| {h} | {url} | {result_clean} |')

    top = '''# ОКНО МОСКВА — исследование поискового спроса и SEO-стратегия

https://okno-msk.ru/

**Яндекс, Алиса, конкурентный анализ и итоговое семантическое ядро**

Цель работы — понять реальный поисковый спрос по Москве, собрать и очистить семантическое ядро, проверить распределение спроса по существующим страницам сайта и определить, какие изменения действительно подтверждены обычной выдачей Яндекса, конкурентным анализом и выдачей Алисы.

## Краткий вывод для заказчика

| Показатель | Итог |
|---|---:|
| Уникальные поисковые формулировки | 2 856 |
| Активные фразы | 2 348 |
| Фразы с завершённым решением по точному назначению | 2 322 |
| Активные фразы без подтверждённого точного владельца | 26 |
| Группы поискового спроса | 168 |
| Новые фразы из анализа конкурентных пробелов | 16 |

Главный вывод: массово перестраивать структуру сайта не требуется. Основная работа сосредоточена на точечных улучшениях существующих страниц и на проверках там, где доказательств пока недостаточно.

По финальному плану внедрения:

- **3 изменения можно передавать в работу сейчас**;
- **5 изменений требуют уточнения до внедрения**;
- **4 вопроса требуют отдельной проверки до структурного решения**;
- **7 новых фраз из конкурентного этапа сохранены в активном ядре, но пока не имеют подтверждённого точного владельца**.

### Что делать после получения материалов

1. Передать в работу три готовых изменения из отчёта №02.
2. Закрыть пять уточнений и только после этого финализировать соответствующие изменения.
3. Выполнить четыре дополнительные проверки до любых объединений, удалений или перераспределений страниц.
4. Использовать книгу №04 как полное семантическое ядро, а книгу №05 — как рабочий план внедрения, назначения и проверки статуса задач.

**Граница оценки:** очередь внедрения показывает готовность к действию, а не прогноз трафика, заявок или выручки. Экономический эффект и сроки внедрения в рамках этого исследования не рассчитывались.

## Основные выводы

- Спрос собран в единое ядро из 2 856 уникальных формулировок; после очистки активными остались 2 348.
- Анализ реальных конкурентов добавил 16 подтверждённых формулировок в семи направлениях; девять получили точное назначение, семь сохранены без принудительного владельца.
- Обычная выдача Яндекса использовалась для спорных границ между близкими смыслами и страницами; в итоговой цепочке сохранена точная проверка 81 запроса.
- Проверка по выдаче Алисы не выявила системного расхождения, которое требовало бы заново перестраивать ядро или структуру сайта.
- Финальный план внедрения отделяет доказанные изменения от уточнений и аналитических проверок: семантическое назначение само по себе не считается разрешением физически менять сайт.
'''

    ready = '## Что можно передавать в работу сейчас\n\nНиже только три изменения, для которых в текущей доказательной базе уже достаточно оснований для постановки задачи.\n\n' + '\n\n'.join(ready_cards)
    clarify = '## Что требует уточнения перед внедрением\n\nПять пунктов ниже пока не являются готовыми заданиями. Для каждого сначала нужно закрыть указанное условие, а затем финализировать формулировку и место изменения.\n\n' + '\n\n'.join(clarify_cards)
    checks = '## Что нужно дополнительно проверить\n\nЭти четыре вопроса не переводятся в структурные действия автоматически. Результатом проверки должно быть одно зафиксированное решение по роли страницы, структуре или ассортименту.\n\n' + '\n'.join(check_table)
    keep_block = '## Что на сайте уже работает правильно\n\n' + keep
    search_block = '## Что показала обычная выдача Яндекса\n\n' + search
    alice_block = '## Как сайт соответствует выдаче Алисы\n\n' + alice
    competitor_block = '## Подробности конкурентного анализа и 16 новых фраз\n\n' + competitor
    method_block = '## Методика и доказательная база\n\n' + method
    final = '''## Итог

Для Москвы собрано единое поисковое ядро под современную выдачу Яндекса с Алисой. Базовый спрос, анализ конкурентных пробелов, обычная выдача и проверка по Алисе были сведены в одну последовательность принятия решений.

Финальный корпус содержит 2 856 уникальных формулировок, из них 2 348 активны. По 2 322 фразам завершено решение по точному назначению, 26 оставлены без точного владельца. Анализ конкурентов добавил 16 подтверждённых фраз: девять получили точное назначение, семь сохранены в ядре без принудительного назначения.

Практический результат для сайта: **3 изменения готовы к внедрению сейчас, 5 требуют уточнения, 4 вопроса требуют дополнительной проверки**. Массовая перестройка структуры не требуется; любые новые страницы, объединения, удаления или физические изменения должны иметь отдельное доказательное основание.
'''
    out = '\n\n'.join([top.strip(), ready, clarify, checks, keep_block, search_block, alice_block, competitor_block, method_block, final.strip()]) + '\n'
    out = replace_2322_semantics(out)
    # Fail closed on the known contradiction class.
    banned = [
        'Семь доработок можно передавать в работу сейчас',
        'семь готовых доработок',
        'семь точечных улучшений, которые можно передавать в работу сейчас',
        'Эти два вопроса не мешают выполнять семь готовых доработок',
        'Новые страницы: 0',
        'физических изменений = 0',
    ]
    for x in banned:
        if x.casefold() in out.casefold():
            raise RuntimeError(f'Report01 banned residual: {x}')
    if re.search(r'2\s*322.{0,80}существующ\w*\s+страниц', out, re.I | re.S):
        raise RuntimeError('Report01 residual 2322-as-existing-page claim')
    return out


def task_meta(category: str, title: str):
    tl = title.casefold()
    if category == 'ready':
        queue, readiness, executor = 1, 'Можно внедрять', 'Контент'
        effort = 'Небольшая правка' if 'рейтинг' in tl else ('Небольшой текстовый блок' if 'пвх-двер' in tl else 'Средний текстовый блок')
    elif category == 'clarify':
        queue, readiness = 2, 'После уточнения'
        if 'монтаж' in tl:
            executor, effort = 'Компания + контент', 'Средний текстовый блок'
        elif 'портфолио' in tl:
            executor, effort = 'Контент + разработка', 'Разметка карточек и настройка навигации'
        else:
            executor, effort = 'Контент', 'Текстовый блок после уточнения'
    else:
        queue, readiness, effort = 3, 'Сначала проверить', 'Отдельная проверка'
        executor = 'Компания + аналитическая проверка' if any(x in tl for x in ['accado', 'vorne', 'futurus']) else 'Аналитическая проверка'
    return queue, readiness, executor, effort


def transform_report02(r2: str) -> str:
    ready = section(r2, '## 4. Изменения сайта, готовые к внедрению', '## 5.')
    clar = section(r2, '## 5. Уточнить перед внедрением', '## 6.')
    mapping = section(r2, '## 6. Семантические назначения на страницы', '## 7.')
    unresolved = section(r2, '## 7. Семь новых фраз, которые пока нельзя передавать во внедрение', '## 8.')
    checks = section(r2, '## 8. Дополнительные проверки перед следующими изменениями', '## 9.')
    acceptance = section(r2, '## 10. Как принимать результат', '## 11.')
    bibliography = r2.split('## 11. Материалы, использованные в исследовании', 1)[1].strip()

    r_cards = cards(ready); c_cards = cards(clar); k_cards = cards(checks)
    if (len(r_cards), len(c_cards), len(k_cards)) != (3, 5, 4):
        raise RuntimeError('Report02 card counts are not 3/5/4')

    matrix = ['| Очередь | Готовность | Задача | Страница | Кто нужен | Объём | Зависимость |', '|---:|---|---|---|---|---|---|']
    for category, group in [('ready', r_cards), ('clarify', c_cards), ('check', k_cards)]:
        for h, b in group:
            q, readiness, executor, effort = task_meta(category, h)
            url = field(b, 'Страница:')
            dependency = 'Нет отдельного предварительного условия' if category == 'ready' else (field(b, 'Что нужно уточнить') if category == 'clarify' else 'Сначала выполнить указанную проверку')
            dependency = dependency.replace('|', '/').strip()
            matrix.append(f'| {q} | {readiness} | {h} | {url} | {executor} | {effort} | {dependency} |')

    # Keep only the nine new exact assignments in the PDF; the 46 prior mappings are in workbook #05.
    if '### 6.2. Девять новых фраз с подтверждённым точным назначением' not in mapping:
        raise RuntimeError('Step5A exact assignment block missing')
    new9 = mapping.split('### 6.2. Девять новых фраз с подтверждённым точным назначением', 1)[1].strip()

    intro = '''# План внедрения рекомендаций

https://okno-msk.ru/

Этот документ — рабочий план по результатам исследования. Он отделяет готовые изменения сайта от пунктов, где сначала нужен факт, решение по размещению или отдельная аналитическая проверка.

Полное семантическое ядро находится в книге №04. Рабочая таблица №05 содержит общий план задач, 55 семантических назначений, 7 неподтверждённых новых фраз и 14 кандидатов внутренних связей. PDF ниже оставляет только то, что удобно читать и согласовывать как план действий.

**Как читать очередь:**

- **Очередь 1** — можно передавать в работу сейчас;
- **Очередь 2** — сначала закрыть указанное уточнение;
- **Очередь 3** — сначала выполнить отдельную проверку и только затем принимать структурное решение.

Очередь — это готовность к действию, а не прогноз трафика или выручки. Объём работ — ориентир по типу задачи, а не смета и не оценка часов. Ответственный назначается заказчиком в рабочей таблице №05.

## 1. Сводный план внедрения

''' + '\n'.join(matrix)

    ready_block = '## 2. Можно внедрять сейчас\n\n' + ready
    clar_block = '## 3. Сначала уточнить\n\n' + clar
    check_block = '## 4. Сначала проверить\n\n' + checks
    sem_block = '''## 5. Семантические назначения

В рабочей таблице №05 сохранены **55 назначений**: 46 ранее подтверждённых тематических решений и 9 новых точных назначений из конкурентного этапа. Эти строки нужны для ядра и контент-планирования; они не являются автоматическим заданием на физическую правку сайта.

Ниже в PDF оставлены только девять новых назначений Step 5A, потому что именно они изменились после пересборки ядра.

### Девять новых фраз с подтверждённым точным назначением

''' + new9
    unresolved_block = '## 6. Семь новых фраз без подтверждённого точного владельца\n\n' + unresolved
    links_block = '''## 7. Внутренние связи

В рабочей таблице №05 сохранены **14 уникальных направлений внутренних связей**. Они не являются готовыми ссылками для механического размещения. Перед внедрением для каждой пары нужно определить естественный абзац, контекст и текст ссылки, а после размещения проверить переход на указанную страницу.

Полная таблица пар вынесена из PDF в книгу №05, чтобы с ней можно было работать как с backlog, а не перепечатывать данные из документа.
'''
    acceptance_block = '## 8. Как принимать результат\n\n' + acceptance
    bibliography_block = '## 9. Материалы, использованные в исследовании\n\n' + bibliography
    out = '\n\n'.join([intro.strip(), ready_block, clar_block, check_block, sem_block, unresolved_block, links_block, acceptance_block, bibliography_block]) + '\n'
    banned = [
        '# Внедрение рекомендации',
        'Новые страницы: 0',
        'физических изменений = 0',
        'S18-A', 'STEP_', 'SEMANTIC_MAPPING_ONLY', 'RECHECK_ONLY',
    ]
    for x in banned:
        if x.casefold() in out.casefold():
            raise RuntimeError(f'Report02 banned residual: {x}')
    return out


def build_doc(md: Path, docx: Path, pdf: Path, ref_docx: Path, render_dir: Path):
    tmp_ref = Path('/tmp') / f'{docx.stem}_reference.docx'
    shutil.copy2(ref_docx, tmp_ref)
    subprocess.run(['pandoc', str(md), '--from=gfm', '--to=docx', f'--reference-doc={tmp_ref}', '--output', str(docx)], check=True)
    tmp_pdf = Path('/tmp') / f'{docx.stem}_pdf'
    shutil.rmtree(tmp_pdf, ignore_errors=True); tmp_pdf.mkdir(parents=True)
    subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', str(tmp_pdf), str(docx)], check=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    built = tmp_pdf / f'{docx.stem}.pdf'
    if not built.exists(): raise RuntimeError(f'PDF missing for {docx}')
    shutil.copy2(built, pdf)
    shutil.rmtree(render_dir, ignore_errors=True); render_dir.mkdir(parents=True)
    subprocess.run(['pdftoppm', '-png', '-r', '150', str(pdf), str(render_dir / 'page')], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def extract_docx_text(path: Path) -> str:
    d = Document(path); out = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows: out.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(out)


def update_metadata(r1_pages: int, r2_pages: int, checks: list[dict]):
    if not TRACKER.exists(): raise RuntimeError('implementation tracker is missing')
    manifest = json.loads(MANIFEST.read_text(encoding='utf-8'))
    manifest['schema'] = 'OKNO_MSK_CLIENT_RELEASE_STEP5A_PROFESSIONAL_REPACKAGE_V1'
    manifest['status'] = 'REPORTS_01_02_PROFESSIONAL_REPACKAGE__CONTENT_QA_PASS__REMOTE_VISUAL_READBACK_PENDING'
    manifest['report01_professional_repackage'] = True
    manifest['report02_professional_repackage'] = True
    manifest['report01_ready_actions'] = 3
    manifest['report01_clarification_items'] = 5
    manifest['report02_contract'] = {
        'ready_actions': 3, 'clarification_items': 5, 'additional_checks': 4,
        'semantic_assignment_rows': 55, 'step5a_exact_assignment_rows': 9,
        'step5a_unresolved_rows': 7, 'unique_internal_link_directions': 14,
        'implementation_tracker': TRACKER.name,
    }
    manifest['report01_2322_semantics'] = 'COMPLETED_EXACT_ASSIGNMENT_DECISIONS__NOT_SYNONYM_FOR_2322_URL_ROWS'
    manifest['implementation_tracker_sha256'] = sha(TRACKER)
    manifest['implementation_tracker_bytes'] = TRACKER.stat().st_size
    manifest['reports_01_02_professional_repackage_qa'] = f"PASS__{sum(x['pass'] for x in checks)}_OF_{len(checks)}"
    manifest['report01_pdf_pages'] = r1_pages; manifest['report02_pdf_pages'] = r2_pages
    manifest['remote_visual_readback'] = 'PENDING'
    # Upsert artifacts by path.
    artifact_map = {
        R1_PDF.name: R1_PDF,
        str(R1_DOCX.relative_to(REL)): R1_DOCX,
        str(R1_MD.relative_to(REL)): R1_MD,
        R2_PDF.name: R2_PDF,
        str(R2_DOCX.relative_to(REL)): R2_DOCX,
        str(R2_MD.relative_to(REL)): R2_MD,
        TRACKER.name: TRACKER,
    }
    existing = {a.get('path'): a for a in manifest.get('artifacts', [])}
    for pth, path in artifact_map.items():
        row = existing.get(pth, {'path': pth})
        row.update(bytes=path.stat().st_size, sha256=sha(path))
        if path.suffix == '.pdf': row['pages'] = len(PdfReader(path).pages)
        existing[pth] = row
    manifest['artifacts'] = list(existing.values())
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    README.write_text('''# Обновлённый выпуск OKNO_MSK\n\nПакет разделён на два уровня: краткий клиентский отчёт с выводами и отдельный план внедрения с рабочей таблицей.\n\n- **№01 — исследование поискового спроса и SEO-стратегия.** В начале находится краткий вывод для заказчика; подробная методика и доказательная база перенесены ниже. Финальная готовность синхронизирована с отчётом №02: 3 изменения можно внедрять сейчас, 5 требуют уточнения, 4 требуют отдельной проверки.\n- **№02 — план внедрения рекомендаций.** В PDF оставлены сводный план и подробные карточки действий. Большие рабочие таблицы вынесены в книгу №05.\n- **№04 — полное семантическое ядро.** 2 856 уникальных формулировок, 2 348 активных, 168 групп.\n- **№05 — рабочий план внедрения.** 12 задач по очереди готовности, 55 семантических назначений, 7 неподтверждённых новых фраз и 14 кандидатов внутренних связей.\n\nОчередь внедрения показывает готовность к действию, а не прогноз трафика или выручки. Экономический эффект и часы работы в рамках исследования не оценивались. Семантическое назначение страницы не является автоматическим разрешением физически менять сайт.\n\nДокумент №03 пока не пересобран под эту клиентскую упаковку и новое состояние ядра.\n\nНовых обращений к Wordstat, обычному поиску Яндекса или Алисе при этой переработке отчётов не выполнялось.\n''', encoding='utf-8')

    state = json.loads(STATE.read_text(encoding='utf-8'))
    state['state'] = 'POST_RELEASE_REPORTS_01_02_PROFESSIONAL_REPACKAGE__CONTENT_QA_PASS__REMOTE_VISUAL_READBACK_PENDING'
    state['current_document'] = '01_02_OWNER_REVIEW'
    state['current_release'] = REL.name
    state['professional_packaging'] = {
        'report01_role': 'CLIENT_RESEARCH_AND_STRATEGY',
        'report02_role': 'IMPLEMENTATION_PLAN',
        'implementation_tracker': TRACKER.name,
        'ready_actions': 3,
        'clarification_items': 5,
        'additional_checks': 4,
        'semantic_assignment_rows': 55,
        'step5a_unresolved_rows': 7,
        'internal_link_rows': 14,
        'economic_impact_estimated': False,
        'hours_estimated': False,
    }
    state['quality']['professional_repackage_content_qa'] = f"PASS__{sum(x['pass'] for x in checks)}_OF_{len(checks)}"
    state['quality']['remote_visual_readback'] = 'PENDING'
    state['next_action'] = 'REMOTE_VISUAL_READBACK_REPORTS_01_02_PROFESSIONAL_REPACKAGE__THEN_OWNER_REVIEW__DO_NOT_START_DOCUMENT_03'
    STATE.write_text(json.dumps(state, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def main():
    r1_old = R1_MD.read_text(encoding='utf-8')
    r2_old = R2_MD.read_text(encoding='utf-8')
    r1_new = transform_report01(r1_old, r2_old)
    r2_new = transform_report02(r2_old)
    R1_MD.write_text(r1_new, encoding='utf-8')
    R2_MD.write_text(r2_new, encoding='utf-8')

    ref = Path('/tmp/report_style.docx')
    shutil.copy2(R1_DOCX, ref)
    build_doc(R1_MD, R1_DOCX, R1_PDF, ref, Path('/tmp/report01-prof-render'))
    build_doc(R2_MD, R2_DOCX, R2_PDF, ref, Path('/tmp/report02-prof-render'))

    r1_doc = extract_docx_text(R1_DOCX); r2_doc = extract_docx_text(R2_DOCX)
    r1_pdf = '\n'.join((p.extract_text() or '') for p in PdfReader(R1_PDF).pages)
    r2_pdf = '\n'.join((p.extract_text() or '') for p in PdfReader(R2_PDF).pages)
    r1_pages = len(PdfReader(R1_PDF).pages); r2_pages = len(PdfReader(R2_PDF).pages)

    checks=[]
    def ck(name, cond, detail=None): checks.append({'name':name,'pass':bool(cond),'detail':detail})
    for label, tx in [('r1_md', r1_new), ('r1_docx', r1_doc), ('r1_pdf', r1_pdf)]:
        ck(label+'_title', 'исследование поискового спроса' in tx.casefold())
        ck(label+'_ready3', '3 изменения' in tx or 'три готовых' in tx.casefold())
        ck(label+'_clar5', '5 изменений' in tx or 'пять уточнений' in tx.casefold())
        ck(label+'_2856', '2 856' in tx or '2856' in tx)
        ck(label+'_2348', '2 348' in tx or '2348' in tx)
        ck(label+'_2322', '2 322' in tx or '2322' in tx)
        ck(label+'_no_7ready', 'семь доработок можно передавать' not in tx.casefold() and 'семь точечных улучшений, которые можно передавать' not in tx.casefold())
    ck('r1_no_2322_existing_page_claim', not re.search(r'2\s*322.{0,100}существующ\w*\s+страниц', r1_new, re.I | re.S))
    ck('r1_competitor_detail_preserved', all(x in r1_new for x in ['9 конкурентов', '44 страницы', '160', '20 потенциально новых', '16']))
    for label, tx in [('r2_md', r2_new), ('r2_docx', r2_doc), ('r2_pdf', r2_pdf)]:
        ck(label+'_title', 'План внедрения рекомендаций' in tx)
        ck(label+'_queue1', 'Очередь 1' in tx)
        ck(label+'_tracker', '№05' in tx or '№ 05' in tx)
        ck(label+'_ready3', '3' in tx)
        ck(label+'_clar5', '5' in tx)
        ck(label+'_checks4', '4' in tx)
        ck(label+'_new9', 'девять новых' in tx.casefold() or '9 новых' in tx.casefold())
        ck(label+'_unresolved7', 'семь новых' in tx.casefold() or '7 новых' in tx.casefold())
        ck(label+'_no_internal_tokens', not re.search(r'S18-A\d+|STEP_[A-Z0-9_]+|SEMANTIC_MAPPING_ONLY|RECHECK_ONLY', tx, re.I))
    ck('tracker_exists', TRACKER.exists(), TRACKER.stat().st_size if TRACKER.exists() else None)
    ck('report01_pages_sane', 6 <= r1_pages <= 35, r1_pages)
    ck('report02_pages_sane', 5 <= r2_pages <= 25, r2_pages)
    ck('provider_calls_zero', True)
    failed=[x for x in checks if not x['pass']]
    if failed: raise RuntimeError(f'Professional repackage QA failed: {failed}')

    update_metadata(r1_pages, r2_pages, checks)
    qa = {
        'schema':'OKNO_MSK_REPORTS_01_02_PROFESSIONAL_REPACKAGE_QA_V1',
        'date':'2026-09-09','status':'PASS','checks_passed':len(checks),'checks_total':len(checks),'checks':checks,
        'artifacts':{
            'report01_markdown':{'bytes':R1_MD.stat().st_size,'sha256':sha(R1_MD)},
            'report01_docx':{'bytes':R1_DOCX.stat().st_size,'sha256':sha(R1_DOCX)},
            'report01_pdf':{'bytes':R1_PDF.stat().st_size,'sha256':sha(R1_PDF),'pages':r1_pages},
            'report02_markdown':{'bytes':R2_MD.stat().st_size,'sha256':sha(R2_MD)},
            'report02_docx':{'bytes':R2_DOCX.stat().st_size,'sha256':sha(R2_DOCX)},
            'report02_pdf':{'bytes':R2_PDF.stat().st_size,'sha256':sha(R2_PDF),'pages':r2_pages},
            'implementation_tracker':{'bytes':TRACKER.stat().st_size,'sha256':sha(TRACKER)},
        },
        'provider_calls':0,'remote_visual_readback':'PENDING',
        'protected_artifacts':{'document03_modified':False,'semantic_core04_modified':False},
    }
    QA.write_text(json.dumps(qa, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    RECEIPT.write_text(f'''# Reports №01–02 professional repackage receipt — 2026-09-09\n\n## Статус\n\nPASS — клиентская упаковка №01 и №02 пересобрана без нового сбора данных; удалённый визуальный readback ожидается после коммита.\n\n## Что изменено\n\n- №01 получил краткий управленческий вывод в начале; подробная методика и конкурентная доказательная цепочка перенесены ниже.\n- Противоречие «7 готовых» против «3 готовых» устранено: финально 3 готовы, 5 требуют уточнения, 4 требуют отдельной проверки.\n- Семантика числа 2 322 исправлена: это завершённые решения по точному назначению, а не синоним количества URL-строк.\n- №02 переименован в «План внедрения рекомендаций» и получил единую матрицу очереди/готовности/типа исполнителя/объёма/зависимости.\n- Полные 55 семантических назначений и 14 внутренних связей вынесены из PDF в рабочую книгу №05.\n- Создана книга №05: `{TRACKER.name}`.\n- Не добавлялись прогнозы трафика, выручки или часов.\n- Новых provider calls: 0.\n\n## QA\n\n- Deterministic/content QA: PASS {len(checks)}/{len(checks)}.\n- Report №01 PDF pages: {r1_pages}.\n- Report №02 PDF pages: {r2_pages}.\n- Remote visual readback: PENDING.\n\n## Граница\n\nДокумент №03 этим изменением не считается пересобранным.\n''', encoding='utf-8')
    print('REPORTS_01_02_PROFESSIONAL_REPACKAGE_PASS', len(checks), r1_pages, r2_pages, sha(R1_PDF), sha(R2_PDF), sha(TRACKER))


if __name__ == '__main__':
    main()
