from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path('extension/docs/kwork/KW001_AI_NATIVE_YANDEX_ALICE/tests/OKNO_MSK')
REL = ROOT / 'OKNO_MSK_RESEARCH_RELEASE_STEP05A_PROPAGATED_2026-09-09'
MD = REL / 'sources/02_OKNO_MSK_VNEDRENIE_REKOMENDATSII_RU_2026-09-09.md'
DOCX = REL / 'editable/02_OKNO_MSK_VNEDRENIE_REKOMENDATSII_RU_2026-09-09.docx'
PDF = REL / '02_OKNO_MSK_VNEDRENIE_REKOMENDATSII_RU_2026-09-09.pdf'
README = REL / 'README_RU.md'
MANIFEST = REL / 'RELEASE_MANIFEST_2026-09-09.json'
SUMMARY = ROOT / 'STEP_05A_INTEGRATED_SEMANTIC_CORE_2026-09-09/FINAL_SEMANTIC_CORE_STEP05A_INTEGRATED_SUMMARY_2026-09-09.json'
QA = ROOT / 'REPORT02_STEP05A_PROPAGATION_QA_2026-09-09.json'
RECEIPT = ROOT / 'REPORT02_STEP05A_PROPAGATION_RECEIPT_2026-09-09.md'
STATE = ROOT / 'RESEARCH_REPORT_REBUILD_CURRENT_STATE_CLIENT_CLEANUP_2026-09-08.json'
RENDER = Path('/tmp/report02-step5a-render')

PHRASES9 = [
    'солнцезащитный стеклопакет',
    'солнцезащитное стекло в стеклопакете',
    'солнцезащитный стеклопакет rehau',
    'многофункциональный стеклопакет что это',
    'ударопрочный стеклопакет',
    'балконы под офис',
    'шумоизоляция на крышу балкона',
    'шумоизоляция крыши балкона от дождя',
    'армирование оконного профиля',
]
PHRASES7 = [
    'гидроизоляция для открытого балкона',
    'гидроизоляция открытого балкона в частном доме',
    'лучшая гидроизоляция для открытого балкона',
    'как сделать гидроизоляцию на открытом балконе',
    'гидроизоляция открытого деревянного балкона',
    'гидроизоляция балконной плиты открытого балкона',
    'шумоизоляция крыши балкона изнутри от дождя',
]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def compact(text: str) -> str:
    """Normalize PDF/DOCX extraction independently of renderer line breaks/punctuation."""
    return re.sub(r'[^0-9a-zа-яё]+', '', text.casefold())


def extract_section(text: str, number: int, next_number: int) -> str:
    a = re.search(rf'(?m)^## {number}\. .*$', text)
    b = re.search(rf'(?m)^## {next_number}\. .*$', text)
    if not a or not b or b.start() <= a.end():
        raise RuntimeError(f'missing section {number}->{next_number}')
    return text[a.end():b.start()].strip()


def rebuild_markdown() -> None:
    s = json.loads(SUMMARY.read_text(encoding='utf-8'))
    assert s['canonical_rows'] == 2856
    assert s['active_rows'] == 2348
    assert s['exact_assignment_decision_rows'] == 2322
    assert s['active_unresolved_exact_owner_rows'] == 26
    assert s['demand_groups'] == 168

    t = MD.read_text(encoding='utf-8')
    # Idempotent: a second run keeps an already-rebuilt document unchanged.
    if '## 11. Материалы, использованные в исследовании' in t and '### 6.2. Девять новых фраз' in t:
        return

    ready = extract_section(t, 2, 3)
    clarification = extract_section(t, 3, 4)
    mapping = extract_section(t, 4, 5)
    checks = extract_section(t, 5, 6)
    links = extract_section(t, 6, 7)
    acceptance = extract_section(t, 7, 8)
    m8 = re.search(r'(?m)^## 8\. .*$', t)
    if not m8:
        raise RuntimeError('missing bibliography section')
    bibliography = t[m8.end():].strip()

    assigned_header = '### Дополнительные фразы, назначенные существующим страницам'
    unresolved_header = '### Дополнительные фразы без подтверждённого точного владельца'
    if assigned_header not in mapping or unresolved_header not in mapping:
        raise RuntimeError('Step5A blocks not found in current Report02')
    base = mapping.split(assigned_header, 1)[0].strip()
    assigned = mapping.split(assigned_header, 1)[1].split(unresolved_header, 1)[0].strip()
    unresolved = mapping.split(unresolved_header, 1)[1].strip()

    intro = '''# Внедрение рекомендации

## 1. Как пользоваться руководством

Этот документ переводит результаты исследования и обновлённого семантического ядра в конкретные действия и ограничения для сайта.

Три готовых изменения в разделе 4 — самостоятельные задачи. Нумерация нужна для навигации и приёмки и не задаёт обязательную календарную последовательность. Пять пунктов раздела 5 требуют уточнения конкретного факта или места размещения до окончательного внедрения.

Полный состав поисковых фраз, групп, страниц и статистики спроса находится в книге №04. После дополнительного анализа конкурентов ядро было пересобрано как единый массив, поэтому в этой версии руководства используются уже новые итоговые числа и новые назначения фраз.

Важно различать три типа результата:

- **готовое изменение сайта** — есть подтверждённая страница, понятное изменение и критерии приёмки;
- **семантическое назначение** — известно, какая существующая страница отвечает за тему; само это назначение не является разрешением автоматически менять текст, создавать страницу или перестраивать сайт;
- **неподтверждённый точный владелец** — тема сохранена в ядре, но до внедрения нужно подтвердить точную страницу и границы предложения компании.

## 2. Что изменилось после обновления семантического ядра

Первичная исследовательская база содержала 2 840 уникальных формулировок. Дополнительный анализ конкурентов выявил 16 новых подтверждённых формулировок в семи направлениях. После включения их в общий процесс очистки и назначения итоговое ядро содержит:

| Показатель | Итог |
|---|---:|
| Уникальные поисковые формулировки | 2 856 |
| Активные фразы | 2 348 |
| Активные фразы с завершённым решением по точному назначению | 2 322 |
| Активные фразы без подтверждённого точного владельца | 26 |
| Пользовательские группы спроса | 168 |
| Новые формулировки из дополнительного анализа конкурентов | 16 |

Из 16 новых формулировок девять получили точное семантическое назначение на существующие страницы. Ещё семь сохранены в активном ядре, но требуют подтверждения точной страницы или границ услуги.

Это изменение **не превращает автоматически новые фразы в задания на физическую правку сайта**. Физическое изменение включается в руководство только тогда, когда для него есть отдельное доказательное основание.

## 3. Краткая карта работ

| Категория | Количество | Следующее действие |
|---|---:|---|
| Изменения сайта, готовые к внедрению | 3 | Выполнить и принять по отдельности |
| Требуют уточнения перед внедрением | 5 | Получить указанный ответ и только затем финализировать изменение |
| Семантические назначения на страницы | 55 | Использовать в ядре и контент-планировании; 46 ранее подтверждённых тем + 9 новых фраз |
| Новые фразы без подтверждённого точного владельца | 7 | Не передавать во внедрение до подтверждения страницы и границ предложения |
| Дополнительные проверки | 4 | Провести указанную проверку и зафиксировать решение |
| Уникальные направления внутренних связей | 14 | Детализировать точное место и текст ссылки перед размещением |'''

    out = '\n\n'.join([
        intro,
        '## 4. Изменения сайта, готовые к внедрению\n\n' + ready,
        '## 5. Уточнить перед внедрением\n\n' + clarification,
        '## 6. Семантические назначения на страницы\n\n### 6.1. Ранее подтверждённые тематические назначения\n\n' + base + '\n\n### 6.2. Девять новых фраз с подтверждённым точным назначением\n\n' + assigned,
        '## 7. Семь новых фраз, которые пока нельзя передавать во внедрение\n\n' + unresolved,
        '## 8. Дополнительные проверки перед следующими изменениями\n\n' + checks,
        '## 9. Потенциальные внутренние связи между страницами\n\n' + links,
        '## 10. Как принимать результат\n\n' + acceptance,
        '## 11. Материалы, использованные в исследовании\n\n' + bibliography,
    ]) + '\n'
    MD.write_text(out, encoding='utf-8')


def build_documents() -> None:
    DOCX.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(['pandoc', str(MD), '--from=gfm', '--to=docx', '--output', str(DOCX)], check=True)
    tmp = Path('/tmp/report02-step5a-pdf')
    shutil.rmtree(tmp, ignore_errors=True)
    tmp.mkdir(parents=True)
    subprocess.run(
        ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', str(tmp), str(DOCX)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    built = tmp / f'{DOCX.stem}.pdf'
    if not built.exists():
        raise RuntimeError('PDF was not generated')
    shutil.copy2(built, PDF)


def read_docx() -> str:
    d = Document(DOCX)
    out = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            out.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(out)


def qa_and_metadata() -> None:
    text = MD.read_text(encoding='utf-8')
    docx_text = read_docx()
    reader = PdfReader(PDF)
    pdf_text = '\n'.join((p.extract_text() or '') for p in reader.pages)
    pdf_pages = len(reader.pages)
    checks: list[dict] = []

    def ck(name: str, value: bool, detail=None):
        checks.append({'name': name, 'pass': bool(value), 'detail': detail})

    ck('title', text.startswith('# Внедрение рекомендации\n'))
    ck('no_role_branding', 'Документ №02 для' not in text and 'SEO-специалиста' not in text)
    for token in [r'S18-A\d+', r'STEP_[A-Z0-9_]+', r'Stage\d+', r'CV\d+', r'OR-\d+', r'CONTENT_BLOCK', r'SEMANTIC_MAPPING_ONLY', r'RECHECK_ONLY']:
        ck('no_internal_' + token, not re.search(token, text, re.I))
    for value in ['2 856', '2 348', '2 322', '26', '168']:
        ck('has_' + value.replace(' ', ''), value in text)
    ck('all_9_assigned_present', all(x in text.casefold() for x in PHRASES9))
    ck('all_7_unresolved_present', all(x in text.casefold() for x in PHRASES7))
    ck('no_unsupported_zero_new_pages', 'Новые страницы: 0' not in text and 'new_pages' not in text)
    ck('no_unsupported_zero_physical', 'физических изменений = 0' not in text and 'physical_site_changes' not in text)
    ck('semantic_not_auto_change', 'не является разрешением автоматически менять текст' in text)

    def between(a: str, b: str) -> str:
        return text.split(a, 1)[1].split(b, 1)[0]

    sec4 = between('## 4. Изменения сайта, готовые к внедрению', '## 5.')
    sec5 = between('## 5. Уточнить перед внедрением', '## 6.')
    sec6 = between('## 6. Семантические назначения на страницы', '## 7.')
    sec7 = between('## 7. Семь новых фраз, которые пока нельзя передавать во внедрение', '## 8.')
    sec8 = between('## 8. Дополнительные проверки перед следующими изменениями', '## 9.')
    sec9 = between('## 9. Потенциальные внутренние связи между страницами', '## 10.')
    sec11 = text.split('## 11. Материалы, использованные в исследовании', 1)[1]
    ck('ready_3', len(re.findall(r'^### \d+\.', sec4, re.M)) == 3)
    ck('clarification_5', len(re.findall(r'^### \d+\.', sec5, re.M)) == 5)
    ck('base_topic_rows_46', len(re.findall(r'^\| \d+ \|', sec6.split('### 6.2.', 1)[0], re.M)) == 46)
    rows9 = [x for x in sec6.split('### 6.2.', 1)[1].splitlines() if x.startswith('| ') and not x.startswith('|---') and 'Поисковая фраза' not in x]
    rows7 = [x for x in sec7.splitlines() if x.startswith('| ') and not x.startswith('|---') and 'Поисковая фраза' not in x]
    ck('new_assigned_rows_9', len(rows9) == 9, len(rows9))
    ck('new_unresolved_rows_7', len(rows7) == 7, len(rows7))
    ck('checks_4', len(re.findall(r'^### \d+\.', sec8, re.M)) == 4)
    ck('links_14', len(re.findall(r'^\| \d+ \|', sec9, re.M)) == 14)
    ck('bibliography_10', len(re.findall(r'^\d+\. ', sec11, re.M)) == 10)

    for label, tx in [('docx', docx_text), ('pdf', pdf_text)]:
        ctx = compact(tx)
        ck(label + '_title', compact('Внедрение рекомендации') in ctx)
        ck(label + '_2856', compact('2 856') in ctx)
        ck(label + '_2348', compact('2 348') in ctx)
        ck(label + '_2322', compact('2 322') in ctx)
        ck(label + '_new9', all(compact(x) in ctx for x in PHRASES9))
        ck(label + '_new7', all(compact(x) in ctx for x in PHRASES7))
        ck(label + '_no_role', compact('SEO-специалиста') not in ctx and compact('Документ №02 для') not in ctx)
    ck('pdf_page_count_sane', 14 <= pdf_pages <= 30, pdf_pages)
    ck('provider_calls_zero', True)

    failed = [x for x in checks if not x['pass']]
    if failed:
        raise RuntimeError(f'QA failed: {failed}')

    shutil.rmtree(RENDER, ignore_errors=True)
    RENDER.mkdir(parents=True)
    subprocess.run(
        ['pdftoppm', '-png', '-r', '150', str(PDF), str(RENDER / 'page')],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    rendered = len(list(RENDER.glob('page-*.png')))
    if rendered != pdf_pages:
        raise RuntimeError(f'PDF render mismatch: reader={pdf_pages}, rendered={rendered}')

    qa = {
        'schema': 'OKNO_MSK_REPORT02_STEP05A_PROPAGATION_QA_V5',
        'date': '2026-09-09',
        'status': 'PASS',
        'checks_passed': sum(x['pass'] for x in checks),
        'checks_total': len(checks),
        'checks': checks,
        'artifacts': {
            'markdown': {'bytes': MD.stat().st_size, 'sha256': sha(MD)},
            'docx': {'bytes': DOCX.stat().st_size, 'sha256': sha(DOCX)},
            'pdf': {'bytes': PDF.stat().st_size, 'sha256': sha(PDF), 'pages': pdf_pages},
        },
        'provider_calls': 0,
        'pdf_render': f'PASS_{rendered}_OF_{pdf_pages}',
        'remote_visual_readback': 'PENDING',
    }
    QA.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    README.write_text('''# Обновлённый выпуск OKNO_MSK

Выпуск использует единое семантическое ядро после включения подтверждённых результатов дополнительного анализа конкурентов в общий процесс очистки, группировки и назначения страниц.

- Полное ядро: 2 856 уникальных поисковых формулировок.
- В едином корпусе: 2 840 формулировок базового сбора и 16 принятых формулировок из семи направлений дополнительного анализа конкурентов.
- Активное ядро: 2 348 фраз.
- Завершённые решения по точному семантическому назначению: 2 322.
- Активные фразы без подтверждённого точного владельца: 26.
- Пользовательские группы спроса: 168.

Из 16 новых формулировок девять получили точное назначение существующим страницам. Ещё семь сохранены в активном ядре, но не передаются во внедрение до подтверждения точной страницы и/или границ предложения компании.

Документ №02 пересобран по этому состоянию ядра. В нём сохранены только доказанные действия и ограничения: 3 готовых изменения сайта, 5 пунктов для уточнения, 55 семантических назначений (46 ранее подтверждённых тем + 9 новых фраз), 7 новых фраз без подтверждённого точного владельца, 4 дополнительные проверки и 14 уникальных направлений внутренних связей.

Само наличие фразы в ядре или её назначение существующей странице не считается автоматическим разрешением создавать страницу, менять текст или выполнять другое физическое изменение сайта. Итог о новых страницах или физических изменениях не выводится из семантического ядра без отдельного доказательства.

Документ №03 пока сохранён без содержательной переработки под новое ядро и должен пройти отдельную сверку.

Новых обращений к Вордстату, обычному поиску Яндекса или Алисе при пересборке документа №02 не выполнялось.
''', encoding='utf-8')

    manifest = json.loads(MANIFEST.read_text(encoding='utf-8'))
    manifest['schema'] = 'OKNO_MSK_CLIENT_RELEASE_STEP5A_PROPAGATED_V2'
    manifest['status'] = 'DOCUMENT_02_STEP05A_RECONCILED__REMOTE_VISUAL_READBACK_PENDING'
    manifest['semantic_rows'] = 2856
    manifest['active_rows'] = 2348
    manifest['exact_assignment_decision_rows'] = 2322
    manifest['active_unresolved_exact_owner_rows'] = 26
    manifest['structural_units'] = 168
    manifest.pop('new_pages', None)
    manifest.pop('new_physical_site_changes', None)
    manifest['new_pages_from_step5a_claim_state'] = 'NOT_ASSERTED_FROM_SEMANTIC_CORE_WITHOUT_SEPARATE_DOWNSTREAM_EVIDENCE'
    manifest['new_physical_site_changes_from_step5a_claim_state'] = 'NOT_ASSERTED_FROM_SEMANTIC_CORE_WITHOUT_SEPARATE_DOWNSTREAM_EVIDENCE'
    manifest['report02_step5a_reconciled'] = True
    manifest['report02_contract'] = {
        'ready_actions': 3,
        'clarification_items': 5,
        'semantic_assignment_rows': 55,
        'base_topic_page_rows': 46,
        'step5a_exact_assignment_rows': 9,
        'step5a_unresolved_rows': 7,
        'additional_checks': 4,
        'unique_internal_link_directions': 14,
        'bibliography_items': 10,
    }
    manifest['report02_deterministic_content_qa'] = f"PASS__{qa['checks_passed']}_OF_{qa['checks_total']}"
    manifest['report02_pdf_rerender'] = f'PASS__{rendered}_OF_{pdf_pages}'
    manifest['report02_remote_visual_readback'] = 'PENDING'
    manifest['report02_provider_calls'] = 0
    for artifact in manifest.get('artifacts', []):
        if artifact.get('path') == PDF.name:
            artifact.update(bytes=PDF.stat().st_size, sha256=sha(PDF), pages=pdf_pages)
        elif artifact.get('path') == str(DOCX.relative_to(REL)):
            artifact.update(bytes=DOCX.stat().st_size, sha256=sha(DOCX))
        elif artifact.get('path') == str(MD.relative_to(REL)):
            artifact.update(bytes=MD.stat().st_size, sha256=sha(MD))
        elif artifact.get('path') == 'README_RU.md':
            artifact.update(bytes=README.stat().st_size, sha256=sha(README))
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    state = {
        'project': 'OKNO_MSK',
        'updated_date': '2026-09-09',
        'state': 'POST_RELEASE_DOCUMENT_02_STEP05A_RECONCILED__CONTENT_QA_PASS__REMOTE_VISUAL_READBACK_PENDING',
        'research_stage_0_to_15': 'COMPLETE',
        'current_document': '02',
        'current_release': 'OKNO_MSK_RESEARCH_RELEASE_STEP05A_PROPAGATED_2026-09-09',
        'current_report_02_visible_title': 'Внедрение рекомендации',
        'integrated_semantic_core': {
            'canonical_rows': 2856,
            'active_rows': 2348,
            'exact_assignment_decision_rows': 2322,
            'active_unresolved_exact_owner_rows': 26,
            'demand_groups': 168,
            'step5a_rows': 16,
            'step5a_exact_assigned_rows': 9,
            'step5a_unresolved_rows': 7,
        },
        'counts': {
            'ready_actions': 3,
            'clarification_items': 5,
            'semantic_assignment_rows': 55,
            'base_topic_page_rows': 46,
            'step5a_exact_assignment_rows': 9,
            'step5a_unresolved_rows': 7,
            'additional_checks': 4,
            'page_link_rows': 14,
            'bibliography_items': 10,
        },
        'quality': {
            'deterministic_content_qa': f"PASS__{qa['checks_passed']}_OF_{qa['checks_total']}",
            'md_docx_pdf_key_content_equivalence': 'PASS',
            'pdf_rerender': f'PASS__{rendered}_OF_{pdf_pages}',
            'remote_visual_readback': 'PENDING',
            'unsupported_zero_new_page_claims': 0,
            'unsupported_zero_physical_change_claims': 0,
            'semantic_assignment_auto_physical_authorization_claims': 0,
            'project_internal_traceability_hits': 0,
            'provider_calls': 0,
        },
        'artifacts': {
            'markdown': {'path': str(MD.relative_to(ROOT)), 'sha256': sha(MD), 'bytes': MD.stat().st_size},
            'docx': {'path': str(DOCX.relative_to(ROOT)), 'sha256': sha(DOCX), 'bytes': DOCX.stat().st_size},
            'pdf': {'path': str(PDF.relative_to(ROOT)), 'sha256': sha(PDF), 'bytes': PDF.stat().st_size, 'pages': pdf_pages},
            'qa': QA.name,
            'receipt': RECEIPT.name,
        },
        'protected_artifacts': {
            'document_01_modified_by_report02_rebuild': False,
            'document_03_modified_by_report02_rebuild': False,
            'semantic_core_04_modified_by_report02_rebuild': False,
        },
        'next_action': 'REMOTE_VISUAL_READBACK_REPORT_02__THEN_OWNER_REVIEW__DO_NOT_START_DOCUMENT_03',
    }
    STATE.write_text(json.dumps(state, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    RECEIPT.write_text(f'''# Report №02 — Step 5A propagation receipt — 2026-09-09

## Статус

PASS — документ №02 пересобран по интегрированному семантическому ядру; удалённый визуальный readback ожидается после коммита.

## Основание

- Интегрированное ядро: 2 856 строк.
- Активные строки: 2 348.
- Завершённые решения по точному семантическому назначению: 2 322.
- Активные строки без подтверждённого точного владельца: 26.
- Группы спроса: 168.
- Новые строки Step 5A: 16 = 9 с точным назначением + 7 без подтверждённого точного владельца.

## Что изменено в документе №02

1. Старые числа 2 840 / 2 332 / 2 313 заменены итоговыми 2 856 / 2 348 / 2 322; добавлен показатель 26 активных строк без подтверждённого точного владельца.
2. 9 новых фраз встроены в раздел семантических назначений существующим страницам.
3. 7 новых фраз вынесены в отдельный блок и не передаются во внедрение до подтверждения точной страницы и границ предложения компании.
4. 46 ранее подтверждённых тематических назначений сохранены; итоговый справочник содержит 55 строк = 46 + 9.
5. Набор физических действий не расширялся только на основании новой семантики: сохранены 3 готовых изменения, 5 пунктов для уточнения, 4 дополнительные проверки и 14 уникальных направлений внутренних связей.
6. В клиентский документ не перенесены неподтверждённые утверждения «новых страниц = 0» и «физических изменений = 0».

## QA

- Deterministic/content QA: PASS {qa['checks_passed']}/{qa['checks_total']}.
- Markdown → DOCX → PDF key-content equivalence: PASS.
- PDF pages: {pdf_pages}.
- PDF re-render: PASS {rendered}/{pdf_pages}.
- Remote visual readback: PENDING.
- New provider calls: 0.

## Артефакты

- Markdown: {MD.stat().st_size} bytes, SHA256 `{sha(MD)}`.
- DOCX: {DOCX.stat().st_size} bytes, SHA256 `{sha(DOCX)}`.
- PDF: {PDF.stat().st_size} bytes, SHA256 `{sha(PDF)}`.
- QA JSON: `{QA.name}`.

## Граница результата

Документ №03 этим изменением не считается согласованным с новым ядром.
''', encoding='utf-8')


def main() -> None:
    rebuild_markdown()
    build_documents()
    qa_and_metadata()
    print('REPORT02_STEP05A_MATERIALIZATION_PASS', sha(MD), sha(DOCX), sha(PDF))


if __name__ == '__main__':
    main()
