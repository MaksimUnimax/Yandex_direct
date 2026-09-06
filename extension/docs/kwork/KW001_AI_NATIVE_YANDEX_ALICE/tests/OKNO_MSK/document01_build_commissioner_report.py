from pathlib import Path
import re, hashlib
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parent
REL = ROOT / 'OKNO_MSK_RESEARCH_RELEASE_CORRECTED_2026-09-05'
SRC = REL / 'sources/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.md'
DOCX = REL / 'editable/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.docx'
TITLE = 'ОКНО МОСКВА — поисковое ядро сайта: спрос, обычный Яндекс и ответы Алисы'


def sha256(path):
    h = hashlib.sha256(); h.update(path.read_bytes()); return h.hexdigest()


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tcPr = cell._tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); x = OxmlElement('w:tblHeader'); x.set(qn('w:val'), 'true'); trPr.append(x)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))


def set_cell_width(cell, cm):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(cm * 567))); tcW.set(qn('w:type'), 'dxa')


def add_page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('Страница '); r.font.name = 'Liberation Sans'; r.font.size = Pt(8)
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
    b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = ' PAGE '
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
    r._r.append(a); r._r.append(b); r._r.append(c)


def add_runs(p, s):
    for part in re.split(r'(\*\*.*?\*\*)', s):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part)
        r.font.name = 'Liberation Sans'


def build_docx():
    text = SRC.read_text(encoding='utf-8'); lines = text.splitlines()
    assert lines[0] == '# ' + TITLE

    # Report-01 non-repeat gate: semantic core is the study object; intent is only a grouping aid.
    forbidden = [
        'точных поисковых формулировок', '2 840 точных', 'пользовательских задач', 'пользовательская задача',
        'спорные семьи', 'спорные семейства', 'изменения запрещены', 'частично готов', 'связанные страницы',
        'семантический маппинг', 'family owner', 'structural unit', 'causal delta', 'proxy',
        'READY', 'HOLD', 'RECHECK', 'SEARCH_REQUIRED', 'PENDING_BUSINESS_DETAIL',
        'Главное за одну минуту', 'Как использовать результаты', 'для владельца сайта', 'владельцу сайта',
        'владелец бизнеса', 'цель — не расширять сайт', 'цель не расширять сайт любой ценой'
    ]
    low = text.lower()
    for bad in forbidden:
        assert bad.lower() not in low, bad

    required = [
        'Задача работы — пересобрать поисковое ядро сайта для Москвы',
        '18 широких стартовых формулировок',
        'Первый проход Wordstat дал 2 415 исходных строк',
        'Это расширение добавило ещё 550 строк',
        'Всего перед очисткой было 2 965 исходных строк',
        '2 840 после очистки — это число уникальных поисковых формулировок',
        'а не 2 840 отдельных замеров точной частотности',
        '2 332 активные формулировки', '2 313', '19', '168 групп поискового спроса',
        'Это было критерием для правильного объединения запросов и распределения их по страницам, а не отдельным предметом исследования',
        '75 проверок — не весь объём исследования',
        'Восемь выбрали для полного сравнения', 'шесть', 'две', '16', 'одну тему',
        '34 значимых вывода',
        'Под ответами Алисы здесь имеется в виду генеративный ответ в поиске Яндекса, работающий на технологиях Алисы',
        'Сайт → Wordstat → очистка и объединение повторов → 168 групп поискового спроса',
        'Семь доработок можно передавать в работу сейчас'
    ]
    for marker in required:
        assert marker in text, marker
    assert 'без операторов точной частотности' in text
    assert 'показатель Wordstat `count`' in text

    assert text.count('**Почему это важно:**') == 7
    assert text.count('**Что рекомендуется сделать:**') == 7
    assert len(re.findall(r'^\|\s*\d+\s*\|', text, re.M)) == 75
    for q in [
        'панорамные алюминиевые окна', 'алюминиевые окна для веранды', 'панорамное остекление балкона',
        'установка подоконника на пластиковые окна', 'французские панорамные окна',
        'замена окна на пластиковое цена москва', 'как открыть пластиковое окно', 'лучшие пластиковые окна'
    ]:
        assert f'| {q} |' in text, q
    for native in ['REHAU', 'Provedal', 'KBE', 'Accado', 'Vorne', 'Futurus', 'fapim', 'rehau thermo']:
        assert native in text, native

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65); sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(.7); sec.footer_distance = Cm(.7)

    styles = doc.styles
    normal = styles['Normal']; normal.font.name = 'Liberation Sans'; normal.font.size = Pt(10.4)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in [('Title', 20, 0, 12), ('Heading 1', 15.5, 14, 7), ('Heading 2', 12.5, 11, 5), ('Heading 3', 11.2, 9, 4)]:
        st = styles[name]; st.font.name = 'Liberation Sans'; st.font.size = Pt(size); st.font.bold = True
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after); st.paragraph_format.keep_with_next = True
    if 'Callout' not in styles:
        s = styles.add_style('Callout', WD_STYLE_TYPE.PARAGRAPH); s.font.name = 'Liberation Sans'; s.font.size = Pt(10.2)
        s.paragraph_format.left_indent = Cm(.35); s.paragraph_format.right_indent = Cm(.35); s.paragraph_format.space_before = Pt(5); s.paragraph_format.space_after = Pt(7)
    if 'SmallTable' not in styles:
        s = styles.add_style('SmallTable', WD_STYLE_TYPE.PARAGRAPH); s.font.name = 'Liberation Sans'; s.font.size = Pt(8.2)
        s.paragraph_format.space_after = Pt(0); s.paragraph_format.line_spacing = 1.0

    hp = sec.header.paragraphs[0]; hp.text = 'ОКНО МОСКВА · поисковое ядро · обычный Яндекс и ответы Алисы'; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.name = 'Liberation Sans'; r.font.size = Pt(7.5)
    add_page_number(sec.footer.paragraphs[0])

    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1; continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title'); p.add_run(line[2:]); i += 1; continue
        if line.startswith('## '):
            doc.add_paragraph(line[3:], style='Heading 1'); i += 1; continue
        if line.startswith('### '):
            doc.add_paragraph(line[4:], style='Heading 2'); i += 1; continue
        if line.startswith('|'):
            tls = []
            while i < len(lines) and (lines[i].startswith('|') or not lines[i].strip()):
                if lines[i].startswith('|'):
                    tls.append(lines[i])
                i += 1
            rows = []
            for tl in tls:
                cols = [c.strip() for c in tl.strip().strip('|').split('|')]
                if all(set(c) <= set('-: ') for c in cols):
                    continue
                rows.append(cols)
            if not rows:
                continue
            if len(rows[0]) == 4 and rows[0][0] == 'Запрос':
                for row in rows[1:]:
                    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; set_cant_split(t.rows[0])
                    c = t.cell(0, 0); set_cell_shading(c, 'F2F4F6'); set_cell_margins(c, 140, 150, 140, 150)
                    p = c.paragraphs[0]; rr = p.add_run(row[0]); rr.bold = True; rr.font.name = 'Liberation Sans'; rr.font.size = Pt(10.6)
                    for label, val in [('Почему выбрали', row[1]), ('Что показал ответ Алисы', row[2]), ('Что это значит для сайта', row[3])]:
                        pp = c.add_paragraph(); pp.paragraph_format.space_after = Pt(2); pp.paragraph_format.keep_together = True
                        a = pp.add_run(label + ': '); a.bold = True; a.font.name = 'Liberation Sans'; a.font.size = Pt(9.2)
                        b = pp.add_run(val); b.font.name = 'Liberation Sans'; b.font.size = Pt(9.2)
                    doc.add_paragraph().paragraph_format.space_after = Pt(1)
            else:
                t = doc.add_table(rows=1, cols=len(rows[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'; set_repeat_header(t.rows[0])
                for j, val in enumerate(rows[0]):
                    c = t.rows[0].cells[j]; c.text = val; set_cell_shading(c, 'E9EDF2'); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c, 70, 70, 70, 70)
                    for p in c.paragraphs:
                        p.style = 'SmallTable'
                        for run in p.runs:
                            run.bold = True
                for row in rows[1:]:
                    rr = t.add_row(); set_cant_split(rr)
                    for j, val in enumerate(row):
                        c = rr.cells[j]; c.text = val; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP; set_cell_margins(c, 55, 60, 55, 60)
                        for p in c.paragraphs:
                            p.style = 'SmallTable'
                if len(rows[0]) == 3:
                    for row in t.rows:
                        for j, w in enumerate([.8, 5.9, 9.2]):
                            set_cell_width(row.cells[j], w)
            continue
        if line.startswith('- '):
            p = doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent = Cm(.45); p.paragraph_format.first_line_indent = Cm(-.25); p.paragraph_format.space_after = Pt(3)
            p.add_run('• '); add_runs(p, line[2:]); i += 1; continue
        if re.match(r'^\d+\. ', line):
            m = re.match(r'^(\d+)\. (.*)', line); p = doc.add_paragraph(style='Normal')
            p.paragraph_format.left_indent = Cm(.5); p.paragraph_format.first_line_indent = Cm(-.32); p.paragraph_format.space_after = Pt(3)
            r = p.add_run(m.group(1) + '. '); r.bold = True; r.font.name = 'Liberation Sans'; add_runs(p, m.group(2)); i += 1; continue

        para = [line]; i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#', '|', '- ')) and not re.match(r'^\d+\. ', lines[i]):
            para.append(lines[i]); i += 1
        s = ' '.join(x.strip() for x in para); p = doc.add_paragraph(style='Normal')
        if s.startswith('**Сайт →'):
            p.style = 'Callout'
            pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F2F4F6'); pPr.append(shd)
        add_runs(p, s)
        if s.startswith('**Что сохранить:**') or s.startswith('**Как должен выглядеть результат:**'):
            p.paragraph_format.keep_with_next = True

    props = doc.core_properties
    props.title = TITLE; props.author = ''; props.last_modified_by = ''; props.subject = 'Исследовательский отчёт для заказчика'; props.comments = ''
    doc.save(DOCX)


if __name__ == '__main__':
    build_docx()
    print('DOCUMENT_01_SEMANTIC_CORE_DOCX_BUILD_PASS', DOCX, sha256(DOCX))
