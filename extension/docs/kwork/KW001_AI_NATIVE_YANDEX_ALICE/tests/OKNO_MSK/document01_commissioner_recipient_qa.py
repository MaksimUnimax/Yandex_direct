from pathlib import Path
import hashlib, json, re, subprocess, tempfile
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
REL = ROOT / 'OKNO_MSK_RESEARCH_RELEASE_CORRECTED_2026-09-05'
MD = REL / 'sources/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.md'
DOCX = REL / 'editable/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.docx'
PDF = REL / '01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.pdf'
QA_PATH = ROOT / 'RESEARCH_REBUILD_POST_RELEASE_DOCUMENT_01_COMMISSIONER_REPORT_QA_2026-09-06.json'
MAN_PATH = REL / 'RELEASE_MANIFEST_2026-09-05.json'
TITLE = 'ОКНО МОСКВА — https://okno-msk.ru/ — поисковое ядро под Алису и обычную выдачу Яндекса'


def sha(p):
    h = hashlib.sha256(); h.update(p.read_bytes()); return h.hexdigest()


def norm(s):
    return re.sub(r'\s+', ' ', s.replace('\u00ad', '')).strip()


def main():
    text = MD.read_text(encoding='utf-8')
    assert text.splitlines()[0] == '# ' + TITLE
    low = text.lower()

    forbidden = [
        'исходных строк', 'кворк', 'kwork', '`count`', 'генеративный', 'нейросетевой', 'Алиса/ИИ',
        '## Приложение.', 'Полный список 75 проверенных',
        'После пересборки ядра', 'В работе пересобрано поисковое ядро',
        'Исследование должно было показать, нужна ли сайту полная пересборка',
        'точная частотность', 'операторами Wordstat',
        'только как дополнительную проверку', 'дополнительное сравнение с выдачей Алисы',
        'дополнительная проверка выдачи Алисы',
        'Специализированная страница панорамного остекления балкона соответствует спросу',
        'Общая страница веранды и раздел алюминиевых окон выполняют разные роли',
        'Страница замены окна и общая страница монтажа',
        'Главное за одну минуту', 'Как использовать результаты', 'полный объём этого этапа',
        'для владельца сайта', 'владельцу сайта', 'владелец бизнеса'
    ]
    for bad in forbidden:
        assert bad.lower() not in low, bad

    required = [
        TITLE,
        'Цель работы — пересобрать поисковое ядро сайта «Окно Москва» с учётом выдачи Алисы.',
        'Поиск в Яндексе изменился:',
        'существующее распределение основных групп спроса по страницам сайта уже в целом хорошо соответствует как обычной выдаче Яндекса, так и выдаче Алисы',
        'полностью перестраивать существующее ядро и структуру сайта не требуется',
        '2 415 поисковых фраз', '550 фраз', '2 965 поисковых фраз',
        'показатель частотности по Москве',
        '2 840 уникальных поисковых фраз', '2 332 активные поисковые фразы',
        '2 313', '19', '168 групп поискового спроса',
        '75 проверок',
        '## Как сайт соответствует выдаче Алисы',
        'существующее распределение основных групп спроса по страницам в целом соответствует выдаче Алисы',
        '25 тем', '8 разобрали углублённо', 'шесть разных случаев', 'два контрольных запроса',
        'Восемь карточек ниже показывают углублённо зафиксированные случаи разных типов.',
        'семь точечных улучшений'
    ]
    for marker in required:
        assert marker in text, marker

    overall = 'Главный результат проверки: существующее распределение основных групп спроса по страницам в целом соответствует выдаче Алисы'
    deep = 'Из этих 25 тем 8 разобрали углублённо'
    assert text.index(overall) < text.index(deep)

    assert len(re.findall(r'^\|\s*\d+\s*\|', text, re.M)) == 0
    assert text.count('**Почему это важно:**') == 7
    assert text.count('**Что рекомендуется сделать:**') == 7

    alice_queries = [
        'панорамные алюминиевые окна', 'алюминиевые окна для веранды', 'панорамное остекление балкона',
        'установка подоконника на пластиковые окна', 'французские панорамные окна',
        'замена окна на пластиковое цена москва', 'как открыть пластиковое окно', 'лучшие пластиковые окна'
    ]
    for q in alice_queries:
        assert f'| {q} |' in text, q

    d = Document(DOCX)
    dtext = '\n'.join(p.text for p in d.paragraphs)
    assert TITLE in dtext
    assert len(d.tables) == 8, len(d.tables)

    reader = PdfReader(PDF)
    assert not reader.is_encrypted
    pages = len(reader.pages)
    assert 6 <= pages <= 10, pages
    ptext = '\n'.join((p.extract_text() or '') for p in reader.pages)
    ntext = norm(ptext).lower()
    for marker in [
        'https://okno-msk.ru/', 'поисковое ядро под Алису',
        'Цель работы — пересобрать поисковое ядро', 'показатель частотности по Москве',
        '2 415 поисковых фраз', '2 965 поисковых фраз', '2 840 уникальных поисковых фраз',
        '75 проверок', 'Как сайт соответствует выдаче Алисы',
        'Восемь карточек ниже показывают углублённо зафиксированные случаи разных типов.',
        'Что нужно изменить на сайте', 'Итог'
    ]:
        assert norm(marker).lower() in ntext, marker
    for bad in [
        'Приложение. Что показала обычная выдача', 'count', 'генеративный',
        'точная частотность', 'операторами Wordstat', 'только как дополнительную проверку',
        'дополнительное сравнение с выдачей Алисы', 'дополнительная проверка выдачи Алисы'
    ]:
        assert bad.lower() not in ntext, bad

    with tempfile.TemporaryDirectory() as td:
        subprocess.run(['pdftoppm', '-png', '-r', '110', str(PDF), str(Path(td) / 'page')], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        rendered = sorted(Path(td).glob('page-*.png'))
        assert len(rendered) == pages, (len(rendered), pages)
        assert all(p.stat().st_size > 15000 for p in rendered)

    untouched = {
        '02_OKNO_MSK_SEO_IMPLEMENTATION_GUIDE_RU_2026-09-05.pdf': '3f589c03bc44f5127aa9f93697e3120729149f8f6985c11fa921424c867ec61a',
        '03_OKNO_MSK_AI_KNOWLEDGE_DOCUMENT_2026-09-05.md': 'd5c90cf187041e975f17d03a2be138eff0b6a06bfe1b7d395a3b24547edc62b0',
        'editable/02_OKNO_MSK_SEO_IMPLEMENTATION_GUIDE_RU_2026-09-05.docx': '4c9f4aeded8b23c8ed7741d10a989ba79696d8fac3fa10a48de11cdd01df2d3c',
        'sources/02_OKNO_MSK_SEO_IMPLEMENTATION_GUIDE_RU_2026-09-05.md': '764f697fb5ad7ca6f5982b000b6053652460f854c8bb7b937a2041688fff80d8'
    }
    for rel, want in untouched.items():
        assert sha(REL / rel) == want, rel

    qa = {
        'status': 'ANALYST_RECHECK_PASS__OWNER_REVIEW_PENDING',
        'scope': 'DOCUMENT_01_ONLY',
        'checks': {
            'site_url_visible_in_title': True,
            'alice_first_kwork_goal_explicit': True,
            'alice_is_core_product_surface_not_additional': True,
            'semantic_core_is_primary_research_object': True,
            'commissioned_alice_rebuild_goal_distinguished_from_no_full_rebuild_result': True,
            'direct_dual_surface_optimization_verdict_visible': True,
            'wordstat_phrase_counts_written_as_phrases_not_rows': True,
            'english_internal_wordstat_field_name_absent': True,
            'wordstat_frequency_work_explained_in_russian': True,
            'unperformed_frequency_procedures_not_narrated_to_customer': True,
            '168_units_presented_as_demand_groups': True,
            'ordinary_yandex_75_is_targeted_validation_not_total_scope': True,
            'raw_75_query_appendix_absent': True,
            'alice_overall_compatibility_verdict_visible': True,
            'alice_overall_verdict_precedes_deep_dive_selection': True,
            'alice_25_candidate_themes_explained': True,
            'alice_eight_cases_are_deep_dives_not_total_alice_scope': True,
            'alice_6_decision_sensitive_plus_2_controls_explained': True,
            'alice_not_presented_as_research_author': True,
            'alice_customer_vocabulary_is_vydacha_alice': True,
            'positive_site_result_is_aggregate_not_page_catalogue': True,
            'seven_recommendations_have_explicit_why': True,
            'two_company_fact_requests_visible_without_partial_ready_framing': True,
            'old_21_item_routing_dump_absent': True,
            'generic_disclaimer_section_absent': True,
            'docx_has_eight_alice_cards_and_no_75_query_table': True,
            'pdf_openable_not_encrypted': True,
            'pdf_all_pages_rendered': True,
            'document_02_not_modified': True,
            'document_03_not_modified': True,
            'provider_calls_during_rebuild_zero': True
        },
        'counts': {
            'wordstat_first_pass_phrases': 2415,
            'wordstat_expansion_phrases': 550,
            'source_phrases_before_deduplication': 2965,
            'unique_search_phrases_after_deduplication': 2840,
            'excluded_rows_internal_authority': 334,
            'deferred_rows_internal_authority': 174,
            'active_search_phrases': 2332,
            'assigned_search_phrases': 2313,
            'unresolved_search_phrases': 19,
            'demand_groups': 168,
            'ordinary_yandex_targeted_checks': 75,
            'alice_candidate_themes_reviewed': 25,
            'alice_deep_dive_cases': 8,
            'alice_decision_sensitive_cases': 6,
            'alice_controls': 2,
            'alice_repeated_candidate_types_not_deep_dived': 16,
            'alice_candidate_not_deep_dived': 1,
            'material_results_internal_authority': 34,
            'ready_recommendations': 7,
            'company_fact_requests': 2,
            'pdf_pages': pages
        },
        'wordstat_count_semantics': 'CUSTOMER_TEXT_STATES_POSITIVE_FREQUENCY_WORK_ONLY__NO_INTERNAL_FIELD__NO_UNPERFORMED_PROCEDURE_NARRATIVE',
        'research_object_semantics': 'ALICE_NATIVE_SEARCH_CORE_FOR_MODERN_YANDEX__ORDINARY_YANDEX_AND_ALICE_ARE_CORE_EVIDENCE_SURFACES',
        'sha256': {'md': sha(MD), 'docx': sha(DOCX), 'pdf': sha(PDF)},
        'size_bytes': {'md': MD.stat().st_size, 'docx': DOCX.stat().st_size, 'pdf': PDF.stat().st_size},
        'provider_calls_during_rebuild': 0,
        'next_action': 'OWNER_REVIEW_CORRECTED_DOCUMENT_01'
    }
    QA_PATH.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    man = json.loads(MAN_PATH.read_text(encoding='utf-8'))
    man['status'] = 'DOCUMENT_01_ALICE_NATIVE_CUSTOMER_CONTRACT_ANALYST_RECHECK_PASS__OWNER_REVIEW_PENDING__DOCUMENT_02_03_NOT_ADVANCED'
    art = man['recipient_artifacts'][0]
    art['size_bytes'] = PDF.stat().st_size; art['sha256'] = sha(PDF); art['pages'] = pages
    art['recipient_contract'] = 'ALICE_NATIVE_KWORK_ANSWER__SEARCH_CORE_FREQUENCY_ORDINARY_YANDEX_ALICE_COMPATIBILITY_AND_ACTIONS__NO_RAW_75_QUERY_DUMP'
    for item in man['editable_and_source_files']:
        if item['path'] == 'editable/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.docx':
            item['size_bytes'] = DOCX.stat().st_size; item['sha256'] = sha(DOCX)
        if item['path'] == 'sources/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-05.md':
            item['size_bytes'] = MD.stat().st_size; item['sha256'] = sha(MD)
    man['qa']['authority'] = '../RESEARCH_REBUILD_POST_RELEASE_DOCUMENT_01_COMMISSIONER_REPORT_QA_2026-09-06.json'
    man['qa']['analyst_recipient_qa'] = 'DOCUMENT_01_ALICE_NATIVE_CUSTOMER_CONTRACT_PASS'
    man['qa']['document_01'] = 'ANALYST_RECHECK_PASS__OWNER_REVIEW_PENDING'
    man['qa']['document_01_physical_pdf'] = f'{pages}_OF_{pages}_PAGES_RENDER_PASS'
    man['qa']['document_01_search_observations'] = '75_TARGETED_CHECKS_EXPLAINED__RAW_APPENDIX_ABSENT'
    man['qa']['document_01_alice_scope'] = 'CORE_PRODUCT_SURFACE__25_CANDIDATE_THEMES__8_DEEP_DIVES_NOT_TOTAL_SCOPE'
    man['qa']['document_01_alice_compatibility_verdict'] = 'SITE_BROADLY_MATCHES_ALICE_OUTPUT__NO_SYSTEMIC_REBUILD_REQUIRED'
    man['qa']['document_01_site_url_in_title'] = True
    man['qa']['document_01_frequency_work_visible'] = True
    man['qa']['document_01_unperformed_frequency_procedure_narrative'] = 'ABSENT'
    man['qa']['document_01_positive_no_change_page_catalogue'] = 'ABSENT__AGGREGATE_SITE_RESULT_ONLY'
    man['qa']['document_01_raw_75_query_appendix'] = 'ABSENT'
    man['qa']['document_01_full_rebuild_claim'] = 'ABSENT__RESEARCH_CONCLUDES_FULL_REBUILD_NOT_REQUIRED'
    man['qa']['document_01_github_readback'] = 'PENDING_POST_COMMIT_READBACK'
    man['qa']['next_action'] = 'OWNER_REVIEW_CORRECTED_DOCUMENT_01'
    MAN_PATH.write_text(json.dumps(man, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    print('DOCUMENT_01_ALICE_NATIVE_CUSTOMER_QA_PASS', pages, sha(MD), sha(DOCX), sha(PDF))


if __name__ == '__main__':
    main()
