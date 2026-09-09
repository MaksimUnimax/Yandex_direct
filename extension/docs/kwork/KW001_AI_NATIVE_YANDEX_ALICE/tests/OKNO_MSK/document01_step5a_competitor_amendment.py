from pathlib import Path
import hashlib, json, re, subprocess, shutil
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
REL = ROOT / 'OKNO_MSK_RESEARCH_RELEASE_STEP05A_PROPAGATED_2026-09-09'
SRC = REL / 'sources/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.md'
DOCX = REL / 'editable/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.docx'
PDF = REL / '01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.pdf'
MANIFEST = REL / 'RELEASE_MANIFEST_2026-09-09.json'
QA = ROOT / 'DOCUMENT_01_STEP05A_COMPETITOR_AMENDMENT_QA_2026-09-09.json'
LOG = ROOT / 'DOCUMENT_01_STEP05A_COMPETITOR_AMENDMENT_2026-09-09.md'
TITLE = 'ОКНО МОСКВА — https://okno-msk.ru/ — поисковое ядро под Алису и обычную выдачу Яндекса'

COMPETITORS = [
    ('mosokna.ru', 5), ('i-okna.ru', 5), ('msk.okna-servise.com', 5),
    ('okna-moskva.ru', 5), ('oknafactoria.ru', 5), ('okna-germany.ru', 5),
    ('fabrikaokon.ru', 5), ('aluminarium.ru', 4), ('elit-balkon.ru', 5),
]
PHRASES = [
    'гидроизоляция для открытого балкона',
    'гидроизоляция открытого балкона в частном доме',
    'лучшая гидроизоляция для открытого балкона',
    'как сделать гидроизоляцию на открытом балконе',
    'гидроизоляция открытого деревянного балкона',
    'гидроизоляция балконной плиты открытого балкона',
    'солнцезащитный стеклопакет',
    'солнцезащитное стекло в стеклопакете',
    'солнцезащитный стеклопакет rehau',
    'многофункциональный стеклопакет что это',
    'ударопрочный стеклопакет',
    'балконы под офис',
    'шумоизоляция на крышу балкона',
    'шумоизоляция крыши балкона от дождя',
    'шумоизоляция крыши балкона изнутри от дождя',
    'армирование оконного профиля',
]

NEW_BLOCK = r'''### Дополнительный анализ конкурентов и расширение ядра

После основного исследования отдельно проверили, какие темы у конкурентов заметны в обычной выдаче Яндекса, но могли быть пропущены при первоначальном сборе ядра. Цель этого этапа — не копировать структуру чужих сайтов и не создавать страницы по одному факту наличия темы у конкурента, а найти возможные пробелы, затем проверить реальный спрос и только после этого решить, нужна ли тема сайту «Окно Москва».

**Какие конкуренты были проанализированы.** Для проверки использовались девять сайтов, страницы которых ранее реально встречались в сохранённой обычной выдаче Яндекса по исследуемым запросам. Случайный обход сайтов и расширение списка конкурентов по доменам не выполнялись.

| Конкурент | Проверено страниц |
|---|---:|
| mosokna.ru | 5 |
| i-okna.ru | 5 |
| msk.okna-servise.com | 5 |
| okna-moskva.ru | 5 |
| oknafactoria.ru | 5 |
| okna-germany.ru | 5 |
| fabrikaokon.ru | 5 |
| aluminarium.ru | 4 |
| elit-balkon.ru | 5 |
| **Итого** | **44** |

На 44 страницах зафиксировали 92 тематических упоминания и свели их в 43 направления, чтобы одинаковые или близкие темы не считать несколько раз. Дальше каждое направление сравнили с уже собранным ядром и границами бизнеса.

Результат отбора был таким:

- 22 направления уже были покрыты существующим ядром и не требовали расширения;
- 3 направления выходили за подтверждённые границы бизнеса;
- по 4 направлениям до Wordstat не было достаточных оснований продолжать расширение;
- 14 направлений передали в Wordstat для проверки реального спроса;
- по этим 14 направлениям получили 160 проверяемых формулировок;
- из них выделили 20 формулировок, где требовалась точная проверка смысла в обычной выдаче Яндекса;
- итогом стали 16 новых поисковых фраз в семи подтверждённых направлениях.

Таким образом, цепочка дополнительной работы выглядела так: **9 конкурентов → 44 страницы → 92 тематических упоминания → 43 сводных направления → 14 направлений с проверкой спроса → 160 формулировок → 20 кандидатов на точную проверку → 16 новых фраз в 7 направлениях.**

Две отдельные темы — «окна для старого фонда» и «кладовая на балконе» — не были превращены в новые страницы или готовые задания: имеющихся данных оказалось недостаточно для такого решения.

### Какие 16 фраз добавлены и чем закончилась их проверка

Все 16 новых фраз включены в итоговый сохранённый корпус. После дополнительной проверки девять получили точное назначение на существующую страницу. Семь фраз оставлены активными без принудительного назначения точной страницы: по ним тематическое направление понятно, но для внедрения не доказана точная страница или граница услуги. Это конечный результат текущего исследования, а не пропущенный этап работы.

#### Гидроизоляция открытого балкона — 6 фраз

- гидроизоляция для открытого балкона
- гидроизоляция открытого балкона в частном доме
- лучшая гидроизоляция для открытого балкона
- как сделать гидроизоляцию на открытом балконе
- гидроизоляция открытого деревянного балкона
- гидроизоляция балконной плиты открытого балкона

**Итог:** все шесть фраз сохранены в ядре, но точная существующая страница и соответствие конкретной услуги компании не доказаны. Новую страницу по ним не создаём и готовое задание на сайт не формируем до появления такого подтверждения.

#### Солнцезащитные стеклопакеты — 3 фразы

- солнцезащитный стеклопакет
- солнцезащитное стекло в стеклопакете
- солнцезащитный стеклопакет rehau

**Итог:** все три фразы относятся к существующей странице https://okno-msk.ru/okna-rehau/steklopakety-dlya-plastikovykh-okon. Отдельная новая страница не требуется.

#### Многофункциональный стеклопакет — 1 фраза

- многофункциональный стеклопакет что это

**Итог:** фраза относится к существующей статье https://okno-msk.ru/stati/kak-vybrat-steklopaket-dlya-plastikovogo-okna/. Отдельная новая страница не требуется.

#### Ударопрочный стеклопакет — 1 фраза

- ударопрочный стеклопакет

**Итог:** фраза относится к существующей странице https://okno-msk.ru/okna-rehau/steklopakety-dlya-plastikovykh-okon. Отдельная новая страница не требуется.

#### Балкон под рабочее место — 1 фраза

- балконы под офис

**Итог:** фраза относится к существующему разделу https://okno-msk.ru/balkony-i-lodzhii/. Отдельная новая страница не требуется.

#### Шумоизоляция крыши балкона — 3 фразы

- шумоизоляция на крышу балкона
- шумоизоляция крыши балкона от дождя
- шумоизоляция крыши балкона изнутри от дождя

**Итог:** первые две фразы относятся к существующей странице https://okno-msk.ru/balkony-i-lodzhii/balkon-s-kryshej/. Для варианта «изнутри» точный способ выполнения услуги компанией не подтверждён, поэтому эта фраза сохранена без принудительного назначения точной страницы.

#### Армирование оконного профиля — 1 фраза

- армирование оконного профиля

**Итог:** фраза относится к существующей статье https://okno-msk.ru/stati/kak-vybrat-plastikovye-okna/. Отдельная новая страница не требуется.

Дополнительная проверка не дала оснований создавать новые страницы: **новых страниц — 0, новых готовых физических изменений сайта — 0**. Практический результат этого этапа — расширение семантического покрытия и более точное распределение новых формулировок между уже существующими материалами без искусственного раздувания структуры.

После добавления этих фраз полный сохранённый корпус содержит 2 856 уникальных фраз. Активное ядро содержит 2 348 фраз; 2 322 активные фразы имеют точное назначение на существующую страницу, 26 оставлены без принудительного точного назначения, включая семь новых формулировок из проверки конкурентов.

Проверка конкурентов использовалась только как способ найти возможные пробелы. Сам факт наличия темы у другого сайта не считался доказательством спроса, необходимости новой страницы или будущего трафика. Решения принимались после сверки с существующим ядром, Wordstat, обычной выдачей Яндекса и границами сайта «Окно Москва».
'''


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def patch_source():
    text = SRC.read_text(encoding='utf-8')
    start = '### Дополнение после проверки конкурентов'
    end = '\n## Как проводилась работа'
    assert text.count(start) == 1, text.count(start)
    a = text.index(start); b = text.index(end, a)
    text = text[:a] + NEW_BLOCK.rstrip() + '\n' + text[b:]

    text = text.replace(
        'По 75 таким запросам сохранена выдача и зафиксирован преобладающий тип ответа.',
        'По первоначальному ядру сохранена выдача по 75 таким запросам. После анализа конкурентов ещё 6 новых спорных формулировок были отдельно проверены в обычной выдаче Яндекса. Всего в исследовании сохранена точная выдача по 81 запросу.'
    )
    text = text.replace(
        'По 75 конкретным запросам сохранена выдача и зафиксировано, какой тип ответа преобладал. Эти 75 проверок служили точечной проверкой решений внутри уже собранного ядра.',
        'По первоначальному ядру по 75 конкретным запросам сохранена выдача и зафиксировано, какой тип ответа преобладал. После анализа конкурентов ещё 6 новых формулировок проверили тем же способом. Итого сохранена точная обычная выдача по 81 запросу: 75 проверок исходного ядра и 6 проверок новых фраз.'
    )
    text = text.replace(
        '**Сайт → Wordstat по Москве и частотность → очистка и объединение повторов → 168 групп поискового спроса → распределение запросов по страницам → 75 выборочных проверок в обычном Яндексе → проверка ядра по выдаче Алисы → 8 углублённых разборов → 7 готовых улучшений и 2 уточнения у компании.**',
        '**Сайт → Wordstat по Москве и частотность → очистка и объединение повторов → 168 групп поискового спроса → распределение запросов по страницам → 75 выборочных проверок исходного ядра в обычном Яндексе → проверка ядра по выдаче Алисы → 8 углублённых разборов → анализ 9 конкурентов и 44 их страниц → проверка 14 новых направлений в Wordstat → 6 дополнительных точных проверок в обычном Яндексе → 16 новых фраз в 7 направлениях → 7 готовых улучшений и 2 уточнения у компании.**'
    )
    text = text.replace(
        '75 проверок конкретных формулировок понадобились для тех случаев, где от выдачи зависело правильное распределение спроса или понимание самой формулировки.',
        'Всего сохранена обычная выдача по 81 конкретному запросу: 75 проверок исходного ядра и 6 дополнительных проверок после анализа конкурентов. Они понадобились для тех случаев, где от выдачи зависело правильное распределение спроса или понимание самой формулировки.'
    )
    text = text.replace(
        '75 проверок обычной выдачи использовались для неоднозначных и важных для распределения запросов случаев. Восемь случаев по выдаче Алисы — это углублённо зафиксированные проверки разных типов решений внутри общей работы по ядру под Алису.',
        'В обычной выдаче сохранена точная проверка 81 запроса: 75 неоднозначных случаев исходного ядра и ещё 6 формулировок, появившихся после анализа конкурентов. Восемь случаев по выдаче Алисы — это углублённо зафиксированные проверки разных типов решений внутри общей работы по ядру под Алису.'
    )
    SRC.write_text(text, encoding='utf-8')
    return text


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))


def add_page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('Страница '); r.font.name = 'Liberation Sans'; r.font.size = Pt(8)
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
    b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = ' PAGE '
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
    r._r.extend([a,b,c])


def add_runs(p, s):
    parts = re.split(r'(\*\*.*?\*\*)', s)
    for part in parts:
        if not part: continue
        bold = part.startswith('**') and part.endswith('**')
        text = part[2:-2] if bold else part
        r = p.add_run(text); r.bold = bold; r.font.name = 'Liberation Sans'


def parse_table(lines, i):
    raw=[]
    while i < len(lines) and lines[i].startswith('|'):
        raw.append(lines[i]); i += 1
    rows=[]
    for ln in raw:
        cols=[c.strip() for c in ln.strip().strip('|').split('|')]
        if all(set(c) <= set('-: ') for c in cols): continue
        rows.append(cols)
    return rows, i


def build_docx(text):
    lines=text.splitlines(); assert lines[0] == '# ' + TITLE
    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.65); sec.bottom_margin=Cm(1.55); sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
    sec.header_distance=Cm(.7); sec.footer_distance=Cm(.7)
    normal=doc.styles['Normal']; normal.font.name='Liberation Sans'; normal.font.size=Pt(10.2)
    normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.06
    for name,size,before,after in [('Title',18.5,0,12),('Heading 1',15.5,14,7),('Heading 2',12.5,11,5),('Heading 3',11.2,9,4)]:
        st=doc.styles[name]; st.font.name='Liberation Sans'; st.font.size=Pt(size); st.font.bold=True
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    hp=sec.header.paragraphs[0]; hp.text='ОКНО МОСКВА · ядро под Алису · обычная выдача Яндекса'; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs: r.font.name='Liberation Sans'; r.font.size=Pt(7.5)
    add_page_number(sec.footer.paragraphs[0])

    i=0
    while i < len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        if line.startswith('# '): p=doc.add_paragraph(style='Title'); add_runs(p,line[2:]); i+=1; continue
        if line.startswith('## '): doc.add_paragraph(line[3:],style='Heading 1'); i+=1; continue
        if line.startswith('### '): doc.add_paragraph(line[4:],style='Heading 2'); i+=1; continue
        if line.startswith('#### '): doc.add_paragraph(line[5:],style='Heading 3'); i+=1; continue
        if line.startswith('|'):
            rows,i=parse_table(lines,i)
            if not rows: continue
            ncols=max(len(r) for r in rows)
            table=doc.add_table(rows=len(rows), cols=ncols); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
            for ri,row in enumerate(rows):
                set_cant_split(table.rows[ri])
                for ci in range(ncols):
                    cell=table.cell(ri,ci); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                    val=row[ci] if ci < len(row) else ''
                    p=cell.paragraphs[0]; add_runs(p,val)
                    for rr in p.runs: rr.font.size=Pt(8.4 if ncols>=4 else 9.2)
                    if ri==0: set_cell_shading(cell,'E8EDF2');
                    if ri==0:
                        for rr in p.runs: rr.bold=True
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent=Cm(.45); p.paragraph_format.first_line_indent=Cm(-.25); p.paragraph_format.space_after=Pt(2.5)
            p.add_run('• '); add_runs(p,line[2:]); i+=1; continue
        if re.match(r'^\d+\. ',line):
            m=re.match(r'^(\d+)\. (.*)',line); p=doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent=Cm(.5); p.paragraph_format.first_line_indent=Cm(-.32)
            r=p.add_run(m.group(1)+'. '); r.bold=True; r.font.name='Liberation Sans'; add_runs(p,m.group(2)); i+=1; continue
        para=[line]; i+=1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','- ')) and not re.match(r'^\d+\. ',lines[i]):
            para.append(lines[i]); i+=1
        p=doc.add_paragraph(style='Normal'); add_runs(p,' '.join(x.strip() for x in para))
    props=doc.core_properties; props.title=TITLE; props.author=''; props.last_modified_by=''; props.subject='Исследовательский отчёт для заказчика'; props.comments=''
    doc.save(DOCX)


def build_pdf():
    if PDF.exists(): PDF.unlink()
    subprocess.run(['libreoffice','--headless','--convert-to','pdf','--outdir',str(REL),str(DOCX)],check=True)
    assert PDF.exists() and PDF.stat().st_size > 10000


def docx_text():
    d=Document(DOCX); out=[]
    for p in d.paragraphs: out.append(p.text)
    for t in d.tables:
        for row in t.rows:
            out.extend(c.text for c in row.cells)
    return '\n'.join(out)


def pdf_text():
    return '\n'.join((p.extract_text() or '') for p in PdfReader(str(PDF)).pages)


def qa_all(text):
    checks={}
    for domain,_ in COMPETITORS: checks['source_competitor_'+domain]=domain in text
    checks['competitor_count_9']=len(COMPETITORS)==9
    checks['competitor_pages_44']=sum(x[1] for x in COMPETITORS)==44 and '9 конкурентов → 44 страницы' in text
    for marker in ['92 тематических упоминания','43 сводных направления','22 направления уже были покрыты','3 направления выходили','4 направлениям','14 направлений','160 формулировок','20 кандидатов','16 новых фраз','7 направлениях']:
        checks['source_'+marker]=marker in text
    for p in PHRASES: checks['source_phrase_'+p]=p in text
    checks['phrase_count_16']=len(PHRASES)==16
    checks['full_2856']='2 856 уникальных фраз' in text
    checks['active_2348']='2 348' in text
    checks['exact_2322']='2 322' in text
    checks['unassigned_26']='26 оставлены' in text
    checks['search_81']='81 запрос' in text and '75' in text and '6' in text
    checks['new_pages_zero']='новых страниц — 0' in text
    checks['physical_zero']='новых готовых физических изменений сайта — 0' in text
    forbidden=['STEP_','S5A-DELTA','REVIEW_SEARCH','CORE_CANDIDATE','SEMANTIC_MAPPING_ONLY','S18-A','CV00','OR-0']
    for token in forbidden: checks['no_internal_'+token]=token not in text
    dt=docx_text(); pt=pdf_text()
    for domain,_ in COMPETITORS:
        checks['docx_'+domain]=domain in dt; checks['pdf_'+domain]=domain in pt
    for p in PHRASES:
        checks['docx_phrase_'+p]=p in dt; checks['pdf_phrase_'+p]=p in pt
    pages=len(PdfReader(str(PDF)).pages); checks['pdf_pages_nonzero']=pages>0
    failed=[k for k,v in checks.items() if not v]
    result={'schema':'OKNO_MSK_DOCUMENT01_STEP05A_COMPETITOR_AMENDMENT_QA_V1','date':'2026-09-09','status':'PASS' if not failed else 'FAIL','checks_total':len(checks),'checks_passed':len(checks)-len(failed),'checks_failed':failed,'pdf_pages':pages,'source_sha256':sha256(SRC),'docx_sha256':sha256(DOCX),'pdf_sha256':sha256(PDF)}
    QA.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    if failed: raise AssertionError(failed)
    return result


def refresh_manifest():
    data=json.loads(MANIFEST.read_text(encoding='utf-8'))
    targets={
        '01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.pdf':PDF,
        'editable/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.docx':DOCX,
        'sources/01_OKNO_MSK_CLIENT_RESEARCH_REPORT_RU_2026-09-09.md':SRC,
    }
    seen=set()
    for item in data['artifacts']:
        if item['path'] in targets:
            p=targets[item['path']]; item['bytes']=p.stat().st_size; item['sha256']=sha256(p); seen.add(item['path'])
    assert seen==set(targets)
    data['report01_competitor_analysis_amended']=True
    data['report01_competitors']=9
    data['report01_competitor_pages']=44
    data['report01_new_phrases_explicit']=16
    data['report01_search_checks_total']=81
    MANIFEST.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def write_log(result):
    LOG.write_text(f'''# Document 01 — competitor-analysis amendment\n\nStatus: **{result['status']}**\n\nReport 01 now explicitly contains:\n\n- all 9 analyzed competitor domains;\n- 44 inspected competitor pages;\n- 92 topic mentions consolidated into 43 directions;\n- the 22 / 3 / 4 / 14 direction screening outcome;\n- 160 Wordstat formulations and 20 exact-Search candidates;\n- all 16 accepted phrases grouped under seven demand directions;\n- final exact-page vs unresolved-owner result for every new phrase;\n- 75 original + 6 post-competitor Search checks = 81 exact Search checks;\n- zero justified new pages and zero new physical site changes from the 16-phrase delta.\n\nPDF pages: {result['pdf_pages']}\n\nSource SHA-256: `{result['source_sha256']}`\nDOCX SHA-256: `{result['docx_sha256']}`\nPDF SHA-256: `{result['pdf_sha256']}`\n''',encoding='utf-8')


def main():
    text=patch_source(); build_docx(text); build_pdf(); result=qa_all(text); refresh_manifest(); result=qa_all(text); write_log(result)
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
