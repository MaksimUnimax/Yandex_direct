#!/usr/bin/env python3
"""Independent, read-only QA for the MK01-only OKNO_MSK rehearsal."""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl


EXPECTED_SHEETS = [
    "Как пользоваться",
    "Все запросы",
    "Рабочее ядро",
    "Группы запросов",
    "На проверку",
    "Исключено",
    "Методика",
]
EXPECTED_ROWS = {
    "Все запросы": 2845,
    "Рабочее ядро": 2190,
    "Группы запросов": 64,
    "На проверку": 192,
    "Исключено": 473,
    "Методика": 22,
}
EXPECTED_FREEZE = {
    "Все запросы": "C6",
    "Рабочее ядро": "C6",
    "Группы запросов": "B6",
    "На проверку": "B6",
    "Исключено": "B6",
    "Методика": "A6",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def read_tsv_gzip(path: Path) -> list[dict[str, str]]:
    with gzip.open(path, "rt", encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def norm(value: str) -> str:
    return " ".join((value or "").strip().lower().split())


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--remote-materialization-commit", default="")
    args = parser.parse_args()

    repo = Path(args.repo_root).resolve()
    out = Path(args.output).resolve()
    work = out.parent
    source = repo / "extension/docs/kwork/KW001_AI_NATIVE_YANDEX_ALICE/tests/OKNO_MSK"

    source_universe = read_tsv(source / "STEP_08_SEARCH_STAGE_SEMANTIC_SET.tsv")
    occurrences = read_tsv(source / "STEP_07C_SEMANTIC_CORRECTION_OCCURRENCES.tsv")
    assignments = read_tsv(source / "STEP_10_FRESH_R1_ASSIGNMENTS_FINAL.tsv")
    taxonomy = read_tsv(source / "STEP_10_FRESH_R1_TAXONOMY_FINAL.tsv")
    search_decisions = read_tsv(source / "STEP_09_EVIDENCE_QUESTION_DECISIONS.tsv")
    step5a_delta = read_tsv(source / "STEP_05A_FIRST_EXECUTION_2026-09-08/STEP_05A_ACCEPTED_SEMANTIC_PIPELINE_DELTA.tsv")

    universe_gzip = work / "MK01_SEMANTIC_UNIVERSE_2026-09-09.tsv.gz"
    groups_tsv = work / "MK01_CLUSTER_SUMMARY_2026-09-09.tsv"
    workbook_path = work / "MK01_OKNO_MSK_SEMANTIC_CORE_2026-09-09.xlsx"
    manifest_path = work / "MK01_MATERIALIZATION_MANIFEST_2026-09-09.json"
    build_report_path = work / "MK01_WORKBOOK_BUILD_REPORT_2026-09-09.json"

    universe = read_tsv_gzip(universe_gzip)
    groups = read_tsv(groups_tsv)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    build_report = json.loads(build_report_path.read_text(encoding="utf-8"))

    checks: dict[str, dict[str, object]] = {}

    def check(name: str, condition: bool, evidence: object) -> None:
        require(condition, f"{name}: {evidence}")
        checks[name] = {"status": "PASS", "evidence": evidence}

    source_keys = [norm(row["phrase"]) for row in source_universe]
    universe_keys = [norm(row["Поисковая фраза"]) for row in universe]
    assignment_keys = [norm(row["phrase"]) for row in assignments]
    delta_keys = {norm(row["phrase"]) for row in step5a_delta}

    check("source_universe_rows", len(source_universe) == 2840, len(source_universe))
    check("source_universe_unique", len(set(source_keys)) == 2840, len(set(source_keys)))
    check("occurrence_rows", len(occurrences) == 2965, len(occurrences))
    check("output_universe_rows", len(universe) == 2840, len(universe))
    check("output_universe_unique", len(set(universe_keys)) == 2840, len(set(universe_keys)))
    check("source_output_phrase_set", set(source_keys) == set(universe_keys), {
        "missing": len(set(source_keys) - set(universe_keys)),
        "added": len(set(universe_keys) - set(source_keys)),
    })
    check("source_assignment_phrase_set", set(source_keys) == set(assignment_keys), {
        "missing": len(set(source_keys) - set(assignment_keys)),
        "added": len(set(assignment_keys) - set(source_keys)),
    })
    check("step5a_delta_rows", len(step5a_delta) == 16, len(step5a_delta))
    check("step5a_contamination", not (set(universe_keys) & delta_keys), len(set(universe_keys) & delta_keys))

    statuses = Counter(row["Итоговый статус"] for row in universe)
    expected_statuses = {
        "В рабочем ядре": 2185,
        "Нужна проверка в обычной выдаче Яндекса": 13,
        "Проверка отложена": 174,
        "Исключено после построчной очистки": 334,
        "Исключено после смысловой группировки": 134,
    }
    check("final_status_partition", statuses == expected_statuses, dict(statuses))
    working = [row for row in universe if row["В рабочем ядре"] == "Да"]
    review = [row for row in universe if row["Итоговый статус"] in {"Нужна проверка в обычной выдаче Яндекса", "Проверка отложена"}]
    excluded = [row for row in universe if row["Итоговый статус"].startswith("Исключено")]
    check("row_reconciliation", len(working) + len(review) + len(excluded) == 2840, {
        "working": len(working), "review": len(review), "excluded": len(excluded), "total": len(universe)
    })
    check("no_silent_default_keep", all(row["Основание решения"].strip() for row in universe), "2840/2840 rows have an explicit reason")
    check("uncertainty_preserved", len(review) == 187 and all(row["Следующий шаг"].strip() for row in review), {
        "review": len(review), "with_next_step": sum(bool(row["Следующий шаг"].strip()) for row in review)
    })
    check("excluded_provenance_preserved", len(excluded) == 468 and all(row["Происхождение данных"].strip() for row in excluded), {
        "excluded": len(excluded), "with_provenance": sum(bool(row["Происхождение данных"].strip()) for row in excluded)
    })
    check("occurrence_reconciliation", sum(int(row["Исходных наблюдений"]) for row in universe) == 2965, sum(int(row["Исходных наблюдений"]) for row in universe))
    check("full_provenance_preserved", all(row["Полная техническая provenance"].strip() for row in universe), "2840/2840 non-empty")
    search_checked = [row for row in universe if row["Проверка в Яндексе"] != "Отдельная проверка этой фразы не проводилась"]
    search_queries = {norm(row["query"]) for row in search_decisions}
    search_in_universe = search_queries & set(universe_keys)
    search_controls = search_queries - set(universe_keys)
    check("saved_search_decisions_reconciled", len(search_decisions) == 75 and len(search_checked) == 66 and len(search_controls) == 9, {
        "decision_rows": len(search_decisions), "output_rows_with_direct_search": len(search_checked), "control_anchors_outside_universe": len(search_controls)
    })

    outside_working = [row for row in working if row["Код группы (для аудита)"].startswith("OUTSIDE_")]
    active_without_group = [row for row in working if row["Код группы (для аудита)"] in {"", "Не назначен"}]
    check("outside_not_in_working_core", not outside_working, len(outside_working))
    check("working_rows_have_group_task_intent", not active_without_group and all(
        row["Группа запросов"].strip() and row["Задача пользователя"].strip() and row["Интент"].strip()
        for row in working
    ), {"working": len(working), "without_group": len(active_without_group)})

    group_ids = [row["Код группы (для аудита)"] for row in groups]
    assignment_counts = Counter(row["cluster_id"] for row in assignments if row["assignment_status"] == "ASSIGNED")
    group_total = sum(int(row["Всего фраз"]) for row in groups)
    group_working = sum(int(row["Фраз в рабочем ядре"]) for row in groups)
    check("cluster_count_and_uniqueness", len(groups) == 59 and len(set(group_ids)) == 59, {"rows": len(groups), "unique": len(set(group_ids))})
    check("cluster_members_reconcile", group_total == 2319 and group_working == 2185, {"all_assigned": group_total, "working": group_working})
    check("cluster_assignment_counts_match", all(int(row["Всего фраз"]) == assignment_counts[row["Код группы (для аудита)"]] for row in groups), "59/59 group counts match assignments")
    representative_membership = []
    assignment_phrases_by_group: dict[str, set[str]] = {}
    for row in assignments:
        if row["assignment_status"] == "ASSIGNED":
            assignment_phrases_by_group.setdefault(row["cluster_id"], set()).add(norm(row["phrase"]))
    for row in groups:
        representative_membership.append(norm(row["Основная формулировка"]) in assignment_phrases_by_group[row["Код группы (для аудита)"]])
    check("cluster_representatives_are_members", all(representative_membership), f"{sum(representative_membership)}/59")
    check("cluster_task_first_fields", all(row["Задача пользователя"].strip() and row["Интент"].strip() and row["Граница группы"].strip() for row in groups), "59/59 groups have task, intent and boundary")
    check("outside_cluster_partition", sum(row["Код группы (для аудита)"].startswith("OUTSIDE_") for row in groups) == 5 and sum(int(row["Всего фраз"]) for row in groups if row["Код группы (для аудита)"].startswith("OUTSIDE_")) == 134, {"groups": 5, "rows": 134})
    roles = [row["Роль в результате"] for row in groups]
    check("cluster_working_first_order", roles[:54] == ["Рабочая смысловая группа"] * 54 and roles[54:] == ["Исключённая смысловая группа"] * 5, {"first": roles[0], "last": roles[-1]})
    check("taxonomy_coverage", set(group_ids) == {row["cluster_id"] for row in taxonomy}, {"groups": len(group_ids), "taxonomy": len(taxonomy)})

    check("manifest_status", manifest["status"] == "PASS", manifest["status"])
    check("manifest_counts", manifest["final_counts"] == {
        "universe": 2840,
        "working_core": 2185,
        "review_or_uncertain": 187,
        "excluded_total": 468,
        "excluded_during_cleanup": 334,
        "excluded_after_task_clustering": 134,
        "semantic_groups_total": 59,
        "working_groups": 54,
        "outside_groups": 5,
        "exact_search_decisions_available": 75,
        "exact_search_decisions_joined_to_universe": 66,
        "search_control_anchors_outside_universe": 9,
        "step5a_contamination": 0,
    }, manifest["final_counts"])
    check("manifest_no_data_loss", manifest["invariants"]["silent_row_loss"] == 0 and manifest["invariants"]["duplicate_phrase_keys"] == 0, manifest["invariants"])
    check("provider_calls_rehearsal", manifest["provider_calls_during_rehearsal"] == 0, manifest["provider_calls_during_rehearsal"])
    artifact_hashes = {
        "semantic_universe_tsv_gzip": sha256(universe_gzip),
        "cluster_summary_tsv": sha256(groups_tsv),
        "client_xlsx": sha256(workbook_path),
        "workbook_build_report": sha256(build_report_path),
    }
    check("artifact_hashes_match_manifest", all(
        artifact_hashes[key] == manifest["artifacts"][key]["sha256"] for key in artifact_hashes
    ), artifact_hashes)

    with zipfile.ZipFile(workbook_path) as archive:
        bad_member = archive.testzip()
    check("xlsx_zip_integrity", bad_member is None, bad_member or "all members readable")
    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=False)
    check("xlsx_sheet_order", workbook.sheetnames == EXPECTED_SHEETS, workbook.sheetnames)
    check("xlsx_sheet_visibility", all(sheet.sheet_state == "visible" for sheet in workbook.worksheets), {sheet.title: sheet.sheet_state for sheet in workbook.worksheets})
    check("xlsx_row_counts", all(workbook[name].max_row == expected for name, expected in EXPECTED_ROWS.items()), {name: workbook[name].max_row for name in EXPECTED_ROWS})
    check("xlsx_freeze_panes", all(str(workbook[name].freeze_panes) == expected for name, expected in EXPECTED_FREEZE.items()), {name: str(workbook[name].freeze_panes) for name in EXPECTED_FREEZE})
    table_info = {}
    for name in EXPECTED_SHEETS[1:]:
        sheet = workbook[name]
        table_info[name] = {table.name: {"ref": table.ref, "filter": table.autoFilter.ref if table.autoFilter else None} for table in sheet.tables.values()}
    check("xlsx_tables_and_filters", all(len(workbook[name].tables) == 1 and next(iter(workbook[name].tables.values())).autoFilter is not None for name in EXPECTED_SHEETS[1:]), table_info)
    check("xlsx_readable_widths", all(
        all((dimension.width or 0) >= 8 for dimension in workbook[name].column_dimensions.values())
        for name in EXPECTED_SHEETS
    ), "all explicitly sized columns >= 8 characters")

    formulas = []
    errors = []
    urls = []
    forbidden_headers = []
    internal_terms = []
    forbidden_header_tokens = re.compile(r"(^|_)(target_url|url_owner|page_owner|structural_action|create_page|merge_page|split_page)($|_)", re.I)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    if value.startswith("="):
                        formulas.append(f"{sheet.title}!{cell.coordinate}")
                    if value.startswith("#") and value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}:
                        errors.append(f"{sheet.title}!{cell.coordinate}:{value}")
                    urls.extend(re.findall(r"https?://[^\s]+", value))
                    if cell.row == 5 and forbidden_header_tokens.search(value):
                        forbidden_headers.append(f"{sheet.title}!{cell.coordinate}:{value}")
                    if re.search(r"\b(provenance|CORE_CANDIDATE|REVIEW_SEARCH|ASSIGNED|PRESERVED_DEFERRED|PRESERVED_EXCLUDED)\b", value):
                        internal_terms.append(f"{sheet.title}!{cell.coordinate}:{value}")
    check("xlsx_no_formula_errors", not errors, errors)
    check("xlsx_no_unexpected_formulas", not formulas, formulas)
    check("xlsx_no_downstream_headers", not forbidden_headers, forbidden_headers)
    check("xlsx_no_unexplained_internal_terms", not internal_terms, internal_terms)
    check("xlsx_no_external_urls", all(url.rstrip(".,)") == "https://okno-msk.ru/" for url in urls), urls)
    check("xlsx_client_language", workbook["Рабочее ядро"]["A5"].value == "Группа запросов" and workbook["Все запросы"]["B5"].value == "Итоговый статус", {
        "primary_sheet_header": workbook["Рабочее ядро"]["A5"].value,
        "status_header": workbook["Все запросы"]["B5"].value,
    })
    check("xlsx_primary_sheet_counts", {
        "Рабочее ядро": workbook["Рабочее ядро"].max_row - 5,
        "На проверку": workbook["На проверку"].max_row - 5,
        "Исключено": workbook["Исключено"].max_row - 5,
    } == {"Рабочее ядро": 2185, "На проверку": 187, "Исключено": 468}, {
        "Рабочее ядро": workbook["Рабочее ядро"].max_row - 5,
        "На проверку": workbook["На проверку"].max_row - 5,
        "Исключено": workbook["Исключено"].max_row - 5,
    })
    check("workbook_build_report", build_report["status"] == "PASS" and build_report["rendered_sheet_count"] == 7 and build_report["workbook_sheet_order"] == EXPECTED_SHEETS, {
        "status": build_report["status"], "rendered_sheet_count": build_report["rendered_sheet_count"], "sheet_order": build_report["workbook_sheet_order"]
    })
    formula_scans = [
        build_report["formula_error_scan_before_export"],
        build_report["formula_error_scan_after_import"],
    ]
    check("workbook_formula_scans", all(
        "matched 0 entries" in scan or '"matches":[]' in scan.replace(" ", "")
        for scan in formula_scans
    ), "before/after import scans contain no matches")

    check("required_files_exist", all(path.exists() and path.stat().st_size > 0 for path in [
        universe_gzip, groups_tsv, workbook_path, manifest_path, build_report_path,
        work / "DELIVERY_SUMMARY.md", work / "REHEARSAL_EXECUTION_LOG.md",
    ]), "all materialization deliverables exist and are non-empty")
    check("remote_materialization_commit_recorded", bool(args.remote_materialization_commit), args.remote_materialization_commit)

    report = {
        "schema": "MK01_OKNO_MSK_INDEPENDENT_QA_V1",
        "date": "2026-09-09",
        "status": "PASS",
        "remote_materialization_commit": args.remote_materialization_commit,
        "counts": {
            "source_occurrences": len(occurrences),
            "source_universe": len(source_universe),
            "working_core": len(working),
            "review_or_uncertain": len(review),
            "excluded": len(excluded),
            "clusters": len(groups),
            "step5a_contamination": len(set(universe_keys) & delta_keys),
            "provider_calls_during_rehearsal": manifest["provider_calls_during_rehearsal"],
        },
        "checks_total": len(checks),
        "checks_passed": len(checks),
        "checks_failed": 0,
        "checks": checks,
    }
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": report["status"],
        "checks_total": report["checks_total"],
        "counts": report["counts"],
        "output": str(out),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
