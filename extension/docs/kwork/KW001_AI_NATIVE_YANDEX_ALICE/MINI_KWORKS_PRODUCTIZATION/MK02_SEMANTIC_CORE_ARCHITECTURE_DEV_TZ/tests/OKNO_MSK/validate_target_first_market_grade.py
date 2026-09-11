#!/usr/bin/env python3
"""Semantic and final-byte validator for the MK02 market-grade Phase-7 package."""

from __future__ import annotations

import csv
import gzip
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
DATE = "2026-09-10"
DELIVERY = HERE / f"CLIENT_DELIVERY_PHASE_7_TARGET_FIRST_MARKET_GRADE_{DATE}"
OUTPUT = HERE / f"TARGET_FIRST_MARKET_GRADE_MACHINE_QA_{DATE}.json"

FOUNDATION = HERE / f"MK02_SEMANTIC_FOUNDATION_{DATE}.tsv.gz"
PHRASES = HERE / f"TARGET_FIRST_PHRASE_LANDING_MAP_MARKET_GRADE_{DATE}.tsv.gz"
CLUSTERS = HERE / f"TARGET_FIRST_CLUSTER_LANDING_MAP_{DATE}.tsv"
REGISTRY = HERE / f"TARGET_PAGE_REGISTRY_{DATE}.tsv"
HIERARCHY = HERE / f"TARGET_ARCHITECTURE_HIERARCHY_{DATE}.tsv"
RECONCILIATION = HERE / f"CURRENT_TARGET_RECONCILIATION_{DATE}.tsv"
SPECS = HERE / f"TARGET_PAGE_SPEC_REGISTER_MARKET_GRADE_{DATE}.tsv"
DELTA = HERE / f"CURRENT_TARGET_CHANGE_DELTA_{DATE}.tsv"

WORKBOOK_REPORT = HERE / f"TARGET_FIRST_MARKET_GRADE_CLIENT_WORKBOOK_BUILD_REPORT_{DATE}.json"
XLSX_PHYSICAL = HERE / f"TARGET_FIRST_MARKET_GRADE_XLSX_PHYSICAL_QA_{DATE}.json"
PDF_REPORT = HERE / f"TARGET_FIRST_MARKET_GRADE_CLIENT_PDF_BUILD_REPORT_{DATE}.json"
PDF_PHYSICAL = HERE / f"TARGET_FIRST_MARKET_GRADE_PDF_PHYSICAL_QA_{DATE}.json"
AUTHORITY_REPORT = HERE / f"TARGET_FIRST_MARKET_GRADE_AUTHORITY_QA_{DATE}.json"
RECIPIENT_REPORT = HERE / f"TARGET_FIRST_MARKET_GRADE_PRODUCT_RECIPIENT_QA_{DATE}.md"

XLSX_NAME = f"SEMANTIC_CORE_AND_TARGET_SEO_STRUCTURE_OKNO_MSK_{DATE}.xlsx"
ANALYTICAL_NAME = f"TARGET_SEO_ARCHITECTURE_REPORT_OKNO_MSK_{DATE}.pdf"
TZ_NAME = f"TARGET_PAGE_SPECIFICATION_TZ_OKNO_MSK_{DATE}.pdf"

REQUIRED_SHEETS = {
    "Начните здесь",
    "Все запросы",
    "Кластеры и задачи",
    "Рассадка запросов",
    "Посадочные страницы",
    "Целевая структура",
    "Реестр страниц",
    "ТЗ по страницам",
    "Изменения сайта",
    "Проверить и отложено",
    "Связи страниц",
    "Текущий сайт",
    "Как проверить",
}

LEVEL_1_FILES = [
    ROOT / "GENERAL_RULES.md",
    ROOT / "DELIVERABLE_SPEC.md",
    ROOT / "QA_AND_RELEASE.md",
    ROOT / "PRODUCT_SCOPE.md",
    ROOT / "PRODUCT_ROADMAP.md",
    ROOT / "ERRORS_AND_LESSONS.md",
    ROOT / "STEP_RULES_INDEX.md",
    ROOT / "steps" / "STEP_14_IMPLEMENTATION_SPECIFICATION.md",
    ROOT / "steps" / "STEP_15_CLIENT_MATERIALIZATION_QA.md",
]

CLIENT_FORBIDDEN = [
    "TARGET_FIRST",
    "MARKET_GRADE",
    "PHASE_7",
    "KEEP_LOCK_AS_TARGET_OWNER",
    "OPTIMIZE_STRENGTHEN",
    "ROUTE_INTERNAL_LINK_CHANGE",
    "RECHECK_NEEDS_EVIDENCE",
    "READY_IMPLEMENTATION_SPEC",
    "PENDING_PLACEMENT_OR_CONTEXT",
    "provider_calls",
    "source_authority",
    "repository",
    "commit",
    "DIY-задача",
    "�",
]

checks: dict[str, dict[str, object]] = {}


def check(name: str, passed: bool, evidence: object) -> None:
    checks[name] = {"status": "PASS" if passed else "FAIL", "evidence": evidence}


def read_tsv(path: Path) -> list[dict[str, str]]:
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_json(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def command_text(arguments: list[str]) -> str:
    completed = subprocess.run(arguments, check=True, capture_output=True, text=True, encoding="utf-8")
    return completed.stdout


def pdf_text(path: Path) -> str:
    return command_text(["pdftotext", "-layout", str(path), "-"])


def page_names(rows: list[dict[str, str]]) -> set[str]:
    return {row["target_page_name_ru"] for row in rows}


def split_topics(value: str) -> list[str]:
    return [item.strip() for item in value.split(" | ") if item.strip()]


def normalized_topic(value: str) -> str:
    value = re.sub(r"^(Раскрыть на этой странице|Встроить без отдельного URL|Кратко упомянуть или использовать как переход):\s*", "", value)
    value = re.sub(r"\s+", " ", value.casefold()).strip(" .«»")
    return value


def workbook_snapshot(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet_rows: dict[str, list[list[object]]] = {}
    visible = []
    frozen = []
    text_values: list[str] = []
    formula_errors: list[str] = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state == "visible":
            visible.append(sheet.title)
        if sheet.freeze_panes:
            frozen.append(sheet.title)
        rows: list[list[object]] = []
        for source_row in sheet.iter_rows(values_only=True):
            values = list(source_row)
            rows.append(values)
            for value in values:
                if isinstance(value, str):
                    text_values.append(value)
                    if any(token in value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                        formula_errors.append(f"{sheet.title}: {value}")
        sheet_rows[sheet.title] = rows
    return {
        "sheet_names": workbook.sheetnames,
        "visible_sheets": visible,
        "frozen_sheets": frozen,
        "rows": sheet_rows,
        "texts": text_values,
        "formula_errors": formula_errors,
    }


def find_header(rows: list[list[object]], required: set[str]) -> list[str]:
    for row in rows:
        values = [str(value) for value in row if value is not None]
        if required.issubset(values):
            return values
    return []


def main() -> int:
    xlsx = DELIVERY / XLSX_NAME
    analytical = DELIVERY / ANALYTICAL_NAME
    tz = DELIVERY / TZ_NAME
    generators = [
        HERE / "build_target_first_market_grade_page_specs.py",
        HERE / "build_target_first_client_workbook.mjs",
        HERE / "build_target_first_client_pdfs.py",
        HERE / "validate_target_first_corrective.py",
    ]
    historical = [
        HERE / f"TARGET_FIRST_PHRASE_LANDING_MAP_{DATE}.tsv.gz",
        HERE / f"TARGET_PAGE_SPEC_REGISTER_{DATE}.tsv",
        HERE / f"CLIENT_DELIVERY_PHASE_7_TARGET_FIRST_CORRECTED_{DATE}",
    ]
    required = [
        FOUNDATION, PHRASES, CLUSTERS, REGISTRY, HIERARCHY, RECONCILIATION, SPECS, DELTA,
        WORKBOOK_REPORT, XLSX_PHYSICAL, PDF_REPORT, PDF_PHYSICAL, AUTHORITY_REPORT,
        RECIPIENT_REPORT, xlsx, analytical, tz, *generators, *LEVEL_1_FILES,
    ]
    missing = [str(path.relative_to(ROOT)) for path in required if not path.exists()]
    check("required_market_grade_inputs_and_outputs_exist", not missing, {"missing": missing, "required": len(required)})
    check("historical_authorities_and_package_preserved", all(path.exists() for path in historical), [str(path.relative_to(ROOT)) for path in historical])
    package_files = sorted(path.name for path in DELIVERY.iterdir() if path.is_file()) if DELIVERY.is_dir() else []
    check("client_delivery_folder_exactly_three_files", package_files == sorted([XLSX_NAME, ANALYTICAL_NAME, TZ_NAME]), package_files)
    if missing:
        return finalize({}, {})

    foundation = read_tsv(FOUNDATION)
    phrases = read_tsv(PHRASES)
    clusters = read_tsv(CLUSTERS)
    registry = read_tsv(REGISTRY)
    hierarchy = read_tsv(HIERARCHY)
    reconciliation = read_tsv(RECONCILIATION)
    specs = read_tsv(SPECS)
    delta = read_tsv(DELTA)

    authority_report = load_json(AUTHORITY_REPORT)
    workbook_report = load_json(WORKBOOK_REPORT)
    xlsx_physical = load_json(XLSX_PHYSICAL)
    pdf_report = load_json(PDF_REPORT)
    pdf_physical = load_json(PDF_PHYSICAL)

    working_ids = {row["phrase_id"] for row in foundation if row["in_working_core"] == "Да"}
    review_ids = {
        row["phrase_id"] for row in foundation
        if row["product_status"] in {"Проверка отложена", "Нужна проверка в обычной выдаче Яндекса"}
    }
    excluded_ids = {row["phrase_id"] for row in foundation if row["product_status"].startswith("Исключено")}
    all_ids = [row["phrase_id"] for row in foundation]
    check(
        "source_accounting_2840_2185_187_468",
        len(foundation) == 2840 and len(working_ids) == 2185 and len(review_ids) == 187 and len(excluded_ids) == 468
        and len(set(all_ids)) == 2840 and not (working_ids & review_ids or working_ids & excluded_ids or review_ids & excluded_ids),
        {"universe": len(foundation), "working": len(working_ids), "review": len(review_ids), "excluded": len(excluded_ids), "unique_ids": len(set(all_ids))},
    )
    phrase_ids = [row["phrase_key"] for row in phrases]
    check("working_phrase_routes_2185_unique", len(phrases) == len(set(phrase_ids)) == 2185, {"rows": len(phrases), "unique": len(set(phrase_ids))})
    check("phrase_routes_exactly_match_working_authority", set(phrase_ids) == working_ids, {"missing": len(working_ids - set(phrase_ids)), "extra": len(set(phrase_ids) - working_ids)})
    check("cluster_task_routes_161", len(clusters) == 161 and len({row["cluster_task_key"] for row in clusters}) == 161, len(clusters))

    target_sets = [
        {row["target_page_key"] for row in registry},
        {row["target_page_key"] for row in hierarchy},
        {row["target_page_key"] for row in reconciliation},
        {row["target_page_key"] for row in specs},
    ]
    check("target_role_alignment_60", all(len(items) == 60 for items in target_sets) and all(items == target_sets[0] for items in target_sets[1:]), [len(items) for items in target_sets])
    check("full_market_grade_page_specs_60", len(specs) == 60 and len({row["page_spec_key"] for row in specs}) == 60, len(specs))

    action_counts = Counter(row["target_action"] for row in specs)
    expected_actions = Counter({"KEEP_LOCK_AS_TARGET_OWNER": 48, "OPTIMIZE_STRENGTHEN": 7, "ROUTE_INTERNAL_LINK_CHANGE": 4, "RECHECK_NEEDS_EVIDENCE": 1})
    check("reconciliation_counts_48_7_4_1", action_counts == expected_actions, dict(action_counts))
    check("fake_create_zero", not any("CREATE" in row["target_action"] for row in specs + reconciliation + delta), {"create": 0})
    spec_by_key = {row["target_page_key"]: row for row in specs}
    delta_failures = [row["change_delta_key"] for row in delta if row["target_page_key"] not in spec_by_key or spec_by_key[row["target_page_key"]]["target_action"] != row["target_action"]]
    check("change_delta_14_is_page_spec_subset", len(delta) == 14 and not delta_failures, {"rows": len(delta), "failures": delta_failures, "unique_pages": len({row["target_page_key"] for row in delta})})

    missing_phrase_wordstat = [row["phrase_key"] for row in phrases if not row["wordstat_popular_count"].strip() or not row["wordstat_popular_count"].isdigit()]
    check("primary_phrase_map_individual_wordstat_2185", not missing_phrase_wordstat, {"present": len(phrases) - len(missing_phrase_wordstat), "missing": missing_phrase_wordstat[:10]})
    metric_note_failures = [row["phrase_key"] for row in phrases if "не сумм" not in row["wordstat_metric_note"].casefold()]
    check("wordstat_is_individual_not_fake_page_sum", not metric_note_failures, {"metric_note_failures": metric_note_failures[:10]})

    phrases_by_page: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in phrases:
        if row["target_landing_page_key"]:
            phrases_by_page[row["target_landing_page_key"]].append(row)
    primary_failures = []
    for row in specs:
        matches = [phrase for phrase in phrases_by_page[row["target_page_key"]] if phrase["phrase"] == row["primary_representative_query"]]
        if not row["primary_query_wordstat"].isdigit() or not matches or all(phrase["wordstat_popular_count"] != row["primary_query_wordstat"] for phrase in matches):
            primary_failures.append(row["target_page_key"])
    check("primary_query_wordstat_matches_phrase_evidence_60", not primary_failures, {"matched": 60 - len(primary_failures), "failures": primary_failures})

    secondary_contract_failures = []
    secondary_value_failures = []
    pages_with_secondaries = 0
    for row in specs:
        count = int(row["secondary_query_count"])
        value = row["secondary_queries_with_wordstat"].strip()
        if count == 0:
            if not value.startswith("Нет:"):
                secondary_contract_failures.append(row["target_page_key"])
            continue
        pages_with_secondaries += 1
        items = split_topics(value)
        if len(items) != count or len(set(items)) != count or row["primary_representative_query"].casefold() in {re.sub(r"\s+—\s+\d+$", "", item).casefold() for item in items}:
            secondary_contract_failures.append(row["target_page_key"])
        if any(not re.fullmatch(r".+\s+—\s+\d+", item) for item in items):
            secondary_value_failures.append(row["target_page_key"])
    check("secondary_query_set_or_explicit_none_60", not secondary_contract_failures, {"pages_with_secondaries": pages_with_secondaries, "explicit_none": 60 - pages_with_secondaries, "failures": secondary_contract_failures})
    check("secondary_queries_have_individual_demand", not secondary_value_failures, {"failures": secondary_value_failures})

    job_failures = [row["target_page_key"] for row in specs if not row["primary_page_job_ru"].strip() or ";" in row["primary_page_job_ru"] or " | " in row["primary_page_job_ru"] or len(row["primary_page_job_ru"]) > 240]
    check("one_clear_primary_page_job_60", not job_failures, {"failures": job_failures})
    page_type_values = Counter(row["page_type"] for row in specs)
    unresolved_types = [row["target_page_key"] for row in specs if row["page_type"] == "Требует определения роли"]
    check("page_type_and_unresolved_role_classification", unresolved_types == ["TP-UNRESOLVED-DIY-WINDOW-TASK"] and sum(page_type_values.values()) == 60, dict(page_type_values))

    boundary_fields = ["own_coverage_clean", "embedded_no_standalone_topics", "support_mention_link_topics", "elsewhere_named_pages"]
    boundary_missing = [row["target_page_key"] for row in specs if any(not row[field].strip() for field in boundary_fields)]
    check("four_page_boundary_fields_present_60", not boundary_missing, {"failures": boundary_missing})
    overlaps: dict[str, list[str]] = {}
    explanation_failures = []
    for row in specs:
        own = {normalized_topic(item) for item in split_topics(row["own_coverage_clean"])}
        elsewhere = {normalized_topic(item) for item in split_topics(row["elsewhere_named_pages"])}
        exact = sorted(own & elsewhere)
        if exact:
            overlaps[row["target_page_key"]] = exact
        named_elsewhere = not row["elsewhere_named_pages"].startswith("Отдельные соседние владельцы")
        explanation = row["boundary_overlap_explanations"].casefold()
        if named_elsewhere and (not explanation.strip() or not ("полное раскрытие" in explanation or "не входит в собственное покрытие" in explanation)):
            explanation_failures.append(row["target_page_key"])
    check("unexplained_own_coverage_elsewhere_overlap_zero", not overlaps, overlaps)
    check("human_readable_boundary_explanations_present", not explanation_failures, {"failures": explanation_failures})

    h1_failures = []
    title_failures = []
    for row in specs:
        if not row["recommended_h1_or_blocker"].strip() or (row["target_action"] == "RECHECK_NEEDS_EVIDENCE" and "Блокер" not in row["recommended_h1_or_blocker"]):
            h1_failures.append(row["target_page_key"])
        state = row["title_requirement_state"]
        value = row["recommended_title_direction_or_blocker"]
        if row["target_action"] == "OPTIMIZE_STRENGTHEN":
            valid = state == "TITLE_DIRECTION_REQUIRED_AND_PRESENT" and value.startswith("Направление Title:")
        elif row["target_action"] == "RECHECK_NEEDS_EVIDENCE":
            valid = state == "BLOCKED_UNRESOLVED_ROLE" and value.startswith("Блокер:")
        else:
            valid = state == "TITLE_CHANGE_NOT_REQUIRED" and "не требуется" in value
        if not valid:
            title_failures.append(row["target_page_key"])
    check("recommended_h1_or_role_blocker_60", not h1_failures, {"failures": h1_failures})
    check("title_direction_status_matches_action_60", not title_failures, {"failures": title_failures})

    priority_failures = [row["target_page_key"] for row in specs if row["analytical_seo_priority"] not in {"Высокий", "Средний", "Низкий"} or not row["analytical_seo_priority_basis"].strip()]
    schedule_directives = ["сделать первым", "в первую очередь внедр", "этап 1", "дедлайн", "за неделю", "за месяц"]
    schedule_failures = [row["target_page_key"] for row in specs if any(token in row["analytical_seo_priority_basis"].casefold() for token in schedule_directives)]
    check("analytical_seo_priority_and_basis_60", not priority_failures, {"counts": dict(Counter(row["analytical_seo_priority"] for row in specs)), "failures": priority_failures})
    check("analytical_priority_not_implementation_schedule", not schedule_failures, {"failures": schedule_failures})

    routed_counts = Counter(row["target_landing_page_key"] for row in phrases if row["target_landing_page_key"])
    routed_failures = [row["target_page_key"] for row in specs if int(row["total_routed_phrase_count"]) != routed_counts[row["target_page_key"]]]
    check("page_spec_total_routed_phrase_counts_match_map", not routed_failures, {"failures": routed_failures})
    non_page_route_sentinels = {"NO_TARGET_OUTSIDE_SCOPE"}
    phrase_page_failures = [row["phrase_key"] for row in phrases if row["target_landing_page_key"] and row["target_landing_page_key"] not in spec_by_key and row["target_landing_page_key"] not in non_page_route_sentinels]
    check("phrase_target_page_keys_valid", not phrase_page_failures, {"failures": phrase_page_failures[:10]})

    workbook = workbook_snapshot(xlsx)
    check("xlsx_hash_matches_build_and_physical_receipts", sha256(xlsx) == workbook_report["output"]["sha256"] == xlsx_physical["final_bytes"]["sha256"], sha256(xlsx))
    check("xlsx_required_sheets_present_without_fixed_count_gate", REQUIRED_SHEETS.issubset(set(workbook["sheet_names"])), {"actual_count": len(workbook["sheet_names"]), "missing": sorted(REQUIRED_SHEETS - set(workbook["sheet_names"]))})
    check("xlsx_all_visible_sheets_rendered_and_frozen", len(workbook["visible_sheets"]) == xlsx_physical["visual_checks"]["rendered_sheets"] == xlsx_physical["visual_checks"]["inspected_sheets"] and set(workbook["visible_sheets"]) == set(workbook["frozen_sheets"]), {"visible": len(workbook["visible_sheets"]), "frozen": len(workbook["frozen_sheets"])})
    check("xlsx_formula_error_scan_zero", not workbook["formula_errors"] and xlsx_physical["package_checks"]["formula_error_tokens"] == 0, workbook["formula_errors"])
    phrase_header = find_header(workbook["rows"]["Рассадка запросов"], {"Фраза", "Вордстат", "Кластер / задача", "Целевая посадочная", "Действие", "Неопределённость"})
    register_header = find_header(workbook["rows"]["Реестр страниц"], {"Основной запрос", "Вордстат основного", "Дополнительные запросы + Вордстат", "Всего распределённых фраз", "Рекомендуемый H1 / блокер", "SEO-приоритет"})
    specs_header = find_header(workbook["rows"]["ТЗ по страницам"], {"Главная задача страницы", "Собственное покрытие", "Встроить без отдельного URL", "Только упомянуть / связать", "Отдать другой названной странице", "Пояснение границы"})
    check("xlsx_phrase_target_view_has_individual_wordstat", bool(phrase_header), phrase_header)
    check("xlsx_page_register_has_primary_secondary_demand_h1_priority", bool(register_header), register_header)
    check("xlsx_page_specs_expose_clean_boundary_concepts", bool(specs_header), specs_header)
    workbook_text_set = {value for value in workbook["texts"]}
    missing_keep_names = [row["target_page_name_ru"] for row in specs if row["target_action"] == "KEEP_LOCK_AS_TARGET_OWNER" and row["target_page_name_ru"] not in workbook_text_set]
    check("all_48_keep_roles_visible_in_xlsx_complete_register", not missing_keep_names, {"visible": 48 - len(missing_keep_names), "missing": missing_keep_names})
    workbook_text = "\n".join(workbook["texts"])
    workbook_leaks = [token for token in CLIENT_FORBIDDEN if token.casefold() in workbook_text.casefold()]
    check("xlsx_client_visible_internal_token_leak_zero", not workbook_leaks, workbook_leaks)
    check("xlsx_physical_package_qa_pass", xlsx_physical["status"] == "PASS" and xlsx_physical["package_checks"]["zip_integrity"] == "PASS" and xlsx_physical["visual_checks"]["clipping"] == xlsx_physical["visual_checks"]["overlap"] == xlsx_physical["visual_checks"]["broken_glyphs"] == 0, xlsx_physical["visual_checks"])

    analytical_text = pdf_text(analytical)
    tz_text = pdf_text(tz)
    actual_pdf = {
        ANALYTICAL_NAME: {"sha256": sha256(analytical), "pages": len(PdfReader(analytical).pages)},
        TZ_NAME: {"sha256": sha256(tz), "pages": len(PdfReader(tz).pages)},
    }
    report_pdf = {row["file"]: row for row in pdf_report["outputs"]}
    physical_pdf = {Path(row["file"]).name: row for row in pdf_physical["files"]}
    check("analytical_pdf_hash_matches_build_and_physical_receipts", actual_pdf[ANALYTICAL_NAME]["sha256"] == report_pdf[ANALYTICAL_NAME]["sha256"] == physical_pdf[ANALYTICAL_NAME]["sha256"], actual_pdf[ANALYTICAL_NAME])
    check("tz_pdf_hash_matches_build_and_physical_receipts", actual_pdf[TZ_NAME]["sha256"] == report_pdf[TZ_NAME]["sha256"] == physical_pdf[TZ_NAME]["sha256"], actual_pdf[TZ_NAME])
    page_match = all(actual_pdf[name]["pages"] == report_pdf[name]["pages"] == physical_pdf[name]["pages"] == physical_pdf[name]["rendered_pages"] == physical_pdf[name]["visually_inspected_pages"] for name in (ANALYTICAL_NAME, TZ_NAME))
    check("both_pdfs_exact_final_pages_parsed_rendered_inspected", page_match, actual_pdf)
    pdf_leaks = {name: [token for token in CLIENT_FORBIDDEN if token.casefold() in text.casefold()] for name, text in [(ANALYTICAL_NAME, analytical_text), (TZ_NAME, tz_text)]}
    check("both_pdfs_client_visible_internal_token_leak_zero", not any(pdf_leaks.values()), pdf_leaks)
    analytical_contract = pdf_report["analytical_contract"]
    check("analytical_pdf_scannable_real_60_role_tree", analytical_contract["tree_role_count"] == 60 and analytical_contract["tree_section_count"] == 9 and analytical_contract["scannable_tree_marker"] in re.sub(r"\s+", " ", analytical_text), analytical_contract)
    check("analytical_pdf_complete_page_model_and_keyword_examples", analytical_contract["complete_page_model_rows"] == 60 and analytical_contract["primary_secondary_wordstat_visible"] is True, analytical_contract)
    tz_contract = pdf_report["tz_contract"]
    check("tz_complete_compact_all_60_page_register", tz_contract["register_rows"] == 60 and "Часть B. Полный компактный реестр 60 целевых страниц" in re.sub(r"\s+", " ", tz_text), tz_contract)
    check("tz_selective_detailed_cards_7_4_1", tz_contract["detailed_cards"] == 12 and tz_contract["detailed_optimize_cards"] == 7 and tz_contract["detailed_route_cards"] == 4 and tz_contract["detailed_recheck_cards"] == 1, tz_contract)
    check("tz_no_mechanical_detailed_keep_cards", tz_contract["detailed_keep_cards"] == 0, tz_contract["detailed_keep_cards"])
    check("both_pdf_physical_qa_pass", pdf_physical["status"] == "PASS" and all(value == 0 for value in pdf_physical["visual_checks"].values()), pdf_physical["visual_checks"])

    sources = {path.name: path.read_text(encoding="utf-8") for path in generators[:3]}
    check("pdf_generator_deterministic_and_no_fixed_page_contract", "rl_config.invariant = 1" in sources["build_target_first_client_pdfs.py"] and pdf_report["no_fixed_pdf_page_count_requirement"] is True, pdf_report["no_fixed_pdf_page_count_requirement"])
    fixed_sheet_patterns = [r"EXPECTED_SHEET_COUNT", r"sheet_count\s*==\s*13", r"len\([^\n]*sheets[^\n]*\)\s*==\s*13"]
    fixed_sheet_hits = [pattern for pattern in fixed_sheet_patterns if re.search(pattern, sources["build_target_first_client_workbook.mjs"], re.IGNORECASE)]
    check("workbook_qa_has_no_fixed_sheet_count_contract", not fixed_sheet_hits, fixed_sheet_hits)
    e6_hits = [name for name, source in sources.items() if re.search(r"[\"']E6[\"']", source)]
    check("no_hard_coded_e6_semantic_authority", not e6_hits, e6_hits)
    description_fields = [field for field in specs[0] if field.casefold() in {"description", "meta_description", "recommended_description"}]
    check("description_not_silently_added_to_base_scope", not description_fields, description_fields)
    check("old_target_first_validator_preserved_not_weakened", (HERE / "validate_target_first_corrective.py").is_file(), "validate_target_first_corrective.py")

    level_1_text = "\n".join(path.read_text(encoding="utf-8") for path in LEVEL_1_FILES)
    level_1_requirements = {
        "phrase_wordstat": "PRIMARY PHRASE MAP WITHOUT INDIVIDUAL DEMAND = FAIL",
        "secondary_queries": "PAGE SPEC WITH ONLY ONE QUERY WHEN USEFUL SECONDARIES EXIST = FAIL",
        "heterogeneous_job": "HETEROGENEOUS PRIMARY PAGE JOB = FAIL",
        "boundary_overlap": "UNEXPLAINED OWN COVERAGE INTERSECT ELSEWHERE != EMPTY = FAIL",
        "h1": "MATERIAL PAGE WITHOUT H1 OR BLOCKER = FAIL",
        "title": "CREATE/OPTIMIZE WITHOUT TITLE DIRECTION OR EXPLICIT BLOCKER = FAIL",
        "priority": "ANALYTICAL PRIORITY PRESENTED AS SCHEDULE = FAIL",
        "tree": "ANALYTICAL PDF WITHOUT SCANNABLE TREE = FAIL",
        "keep_register": "KEEP DROPPED FROM COMPLETE REGISTER = FAIL",
        "keep_compression": "EVERY KEEP FORCED INTO FULL DETAILED CARD = QUALITY FAIL",
        "fake_create": "FAKE CREATE USED FOR PORTFOLIO = FAIL",
    }
    missing_level_1 = [name for name, marker in level_1_requirements.items() if marker not in level_1_text]
    check("level_1_qg01_qg08_hard_fail_markers_permanent", not missing_level_1, {"present": len(level_1_requirements) - len(missing_level_1), "required": len(level_1_requirements), "missing": missing_level_1})
    check("level_1_no_fake_create_portfolio_rule", "clearly labelled demo" in level_1_text and "fake CREATE" in level_1_text, "truthful case or clearly labelled demo")

    recipient_text = RECIPIENT_REPORT.read_text(encoding="utf-8")
    check("recipient_qa_ten_scenarios_pass", "RECIPIENT QA = PASS (10/10)" in recipient_text and "only the final three client files" in recipient_text.casefold(), RECIPIENT_REPORT.name)
    provider_counts = [authority_report["counts"]["provider_calls"], workbook_report.get("provider_calls", 0), pdf_report["provider_calls"]]
    check("new_provider_calls_zero", all(value == 0 for value in provider_counts), provider_counts)
    check("authority_generator_qa_18_of_18_pass", authority_report["status"] == "PASS" and authority_report["checks_total"] == 18 and all(value == "PASS" for value in authority_report["checks"].values()), {"status": authority_report["status"], "checks": authority_report["checks_total"]})

    file_hashes = {path.name: sha256(path) for path in (xlsx, analytical, tz)}
    counts = {
        "semantic_universe": len(foundation),
        "working_phrase_routes": len(phrases),
        "review": len(review_ids),
        "excluded": len(excluded_ids),
        "cluster_tasks": len(clusters),
        "target_pages": len(registry),
        "page_specs": len(specs),
        "keep": action_counts["KEEP_LOCK_AS_TARGET_OWNER"],
        "optimize": action_counts["OPTIMIZE_STRENGTHEN"],
        "route": action_counts["ROUTE_INTERNAL_LINK_CHANGE"],
        "recheck": action_counts["RECHECK_NEEDS_EVIDENCE"],
        "change_delta": len(delta),
        "create": 0,
        "provider_calls": 0,
        "xlsx_sheets": len(workbook["sheet_names"]),
        "analytical_pdf_pages": actual_pdf[ANALYTICAL_NAME]["pages"],
        "tz_pdf_pages": actual_pdf[TZ_NAME]["pages"],
    }
    return finalize(counts, file_hashes)


def finalize(counts: dict[str, object], file_hashes: dict[str, str]) -> int:
    failed = [name for name, result in checks.items() if result["status"] != "PASS"]
    report = {
        "schema": "MK02_TARGET_FIRST_MARKET_GRADE_MACHINE_QA_V1",
        "date": DATE,
        "status": "PASS" if not failed else "FAIL",
        "checks_total": len(checks),
        "checks_passed": len(checks) - len(failed),
        "checks_failed": len(failed),
        "failed_checks": failed,
        "counts": counts,
        "client_file_sha256": file_hashes,
        "checks": checks,
    }
    OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not failed else 1


if __name__ == "__main__":
    sys.exit(main())
